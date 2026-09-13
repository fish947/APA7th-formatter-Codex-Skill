"""End-to-end preservation and safety regression tests for the APA formatter.

Run: python -m unittest discover -s apa7_formatter -p 'test_*.py' -v
Fixtures and results exist only inside an automatically removed temp directory.
These tests verify engineering behavior, not full APA or visual compliance.
"""
from __future__ import annotations

import copy
import hashlib
import io
from pathlib import Path
import struct
import tempfile
import unittest
from zipfile import ZipFile
import zlib

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from lxml import etree

import apa7_format as apa


def tiny_png():
    """A valid small image generated in memory, with no imaging dependency."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00\x00\xff\x00" * 2))
            + chunk(b"IEND", b""))


def field(paragraph, instruction, result):
    for kind, value in (("begin", None), ("instruction", instruction),
                        ("separate", None), ("result", result), ("end", None)):
        run = paragraph.add_run()
        if kind == "result":
            run.text = value
        elif kind == "instruction":
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = value
            run._r.append(node)
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            run._r.append(node)


def hyperlink(paragraph, text, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run, content = OxmlElement("w:r"), OxmlElement("w:t")
    content.text = text
    run.append(content)
    link.append(run)
    paragraph._p.append(link)


def c14n(element):
    return etree.tostring(element, method="c14n")


def body_content_without_formatting(doc):
    """Independent oracle: retain content tree order, excluding format properties."""
    body = copy.deepcopy(doc._element.body)
    for tag in ("w:pPr", "w:rPr", "w:sectPr", "w:tblPr", "w:trPr", "w:tcPr"):
        for element in list(body.iter(qn(tag))):
            element.getparent().remove(element)
    return c14n(body)


def page_fields(path):
    with ZipFile(path) as package:
        counts = {}
        for name in package.namelist():
            if name.startswith("word/header") and name.endswith(".xml"):
                root = etree.fromstring(package.read(name))
                counts[name] = sum((node.text or "").strip() == "PAGE" for node in root.iter(qn("w:instrText")))
        return counts


class FormatterRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="apa7_tests_")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def save(self, doc, name="source.docx"):
        path = self.root / name
        doc.save(path)
        return path

    def format(self, source, name="formatted.docx", **kwargs):
        kwargs.setdefault("profile", "student")
        return apa.format_file(source, output=self.root / name, **kwargs)

    def assert_no_output(self, name="formatted.docx"):
        out = self.root / name
        for path in (out, out.with_suffix(".feedback.html")):
            self.assertFalse(path.exists(), str(path))

    def test_fields_hyperlinks_emphasis_and_media_preserved(self):
        doc = Document()
        p = doc.add_paragraph("A scientific finding: ")
        p.add_run("p").italic = True
        p.add_run("2").font.superscript = True
        p.add_run(" substantive emphasis ").bold = True
        field(p, ' ADDIN ZOTERO_ITEM CSL_CITATION {"citationID":"test-α"} ', "(Smith, 2024)")
        hyperlink(p, " DOI", "https://doi.org/10.1234/example")
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), "42")
        start.set(qn("w:name"), "ResultAnchor")
        end.set(qn("w:id"), "42")
        p._p.insert(0, start)
        p._p.append(end)
        doc.add_paragraph().add_run().add_picture(io.BytesIO(tiny_png()), width=Inches(1))
        source = self.save(doc)
        source_bytes = source.read_bytes()
        before = Document(source)
        expected_content = body_content_without_formatting(before)
        before_media = apa.package_payloads(source)
        out, report = self.format(source)
        after = Document(out)
        self.assertEqual(source.read_bytes(), source_bytes)
        self.assertEqual(body_content_without_formatting(after), expected_content)
        self.assertEqual(apa.package_payloads(out), before_media)
        self.assertEqual(apa.visible_text(after.paragraphs[0]), apa.visible_text(before.paragraphs[0]))
        by_text = {run.text: run for run in after.paragraphs[0].runs if run.text}
        self.assertTrue(by_text["p"].italic)
        self.assertTrue(by_text["2"].font.superscript)
        self.assertTrue(by_text[" substantive emphasis "].bold)
        links = [rel.target_ref for rel in after.part.rels.values() if rel.reltype == RT.HYPERLINK]
        self.assertEqual(links, ["https://doi.org/10.1234/example"])
        self.assertEqual(report["preservation"], "passed")
        self.assertEqual(report["source_sha256"], hashlib.sha256(source_bytes).hexdigest())
        self.assertEqual(set(report["feedback"]), {"summary", "changed", "apa_sources", "locations", "needs_review"})
        self.assertEqual({path.name for path in self.root.iterdir()}, {"source.docx", "formatted.docx"})

    def test_second_run_does_not_duplicate_page_fields_or_body_content(self):
        doc = Document()
        doc.add_paragraph("Ordinary text with no special formatting.")
        source = self.save(doc)
        first, _ = self.format(source, name="first.docx")
        second, _ = self.format(first, name="second.docx")
        counts = page_fields(first)
        self.assertEqual(len(counts), 3, "All first/even/default header variants should have a PAGE field")
        self.assertTrue(all(count == 1 for count in counts.values()), counts)
        self.assertEqual(page_fields(second), counts)
        self.assertEqual(c14n(Document(first)._element.body), c14n(Document(second)._element.body))

    def test_unaccepted_body_revision_rejected_without_outputs(self):
        doc = Document()
        p = doc.add_paragraph("Original ")
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), "1")
        insertion.set(qn("w:author"), "Reviewer")
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "unaccepted insertion"
        run.append(text)
        insertion.append(run)
        p._p.append(insertion)
        source = self.save(doc)
        unchanged = source.read_bytes()
        with self.assertRaisesRegex(ValueError, "修订"):
            self.format(source)
        self.assertEqual(source.read_bytes(), unchanged)
        self.assert_no_output()

    def test_unaccepted_header_revision_also_rejected(self):
        doc = Document()
        doc.add_paragraph("Body without revisions.")
        p = doc.sections[0].header.paragraphs[0]
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), "2")
        insertion.set(qn("w:author"), "Reviewer")
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "unaccepted header revision"
        run.append(text)
        insertion.append(run)
        p._p.append(insertion)
        source = self.save(doc)
        with self.assertRaisesRegex(ValueError, "修订"):
            self.format(source)
        self.assert_no_output()

    def test_tracked_table_property_change_rejected(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(1, 0).text = "Value"
        change = OxmlElement("w:tblPrChange")
        change.set(qn("w:id"), "3")
        change.set(qn("w:author"), "Reviewer")
        change.append(OxmlElement("w:tblPr"))
        table._tbl.tblPr.append(change)
        source = self.save(doc)
        with self.assertRaisesRegex(ValueError, "修订"):
            self.format(source)
        self.assert_no_output()

    def test_existing_source_output_and_optional_feedback_are_never_overwritten(self):
        doc = Document()
        doc.add_paragraph("Keep this original.")
        source = self.save(doc)
        original_bytes = source.read_bytes()
        with self.assertRaises(ValueError):
            apa.format_file(source, output=source, profile="student")
        self.assertEqual(source.read_bytes(), original_bytes)
        existing = self.save(Document(), "existing.docx")
        existing_bytes = existing.read_bytes()
        with self.assertRaises(ValueError):
            apa.format_file(source, output=existing, profile="student")
        self.assertEqual(existing.read_bytes(), existing_bytes)
        out = self.root / "collision.docx"
        feedback = out.with_suffix(".feedback.html")
        feedback.write_text("existing feedback", encoding="utf-8")
        with self.assertRaises(ValueError):
            apa.format_file(source, output=out, profile="student", save_feedback=True)
        self.assertFalse(out.exists())
        self.assertEqual(feedback.read_text(encoding="utf-8"), "existing feedback")

    def test_profile_is_required_and_optional_feedback_is_short(self):
        doc = Document()
        doc.add_paragraph("A short paper.")
        source = self.save(doc)
        with self.assertRaisesRegex(ValueError, "请先选择论文类型"):
            apa.format_file(source, output=self.root / "missing-profile.docx")
        out, report = apa.format_file(source, output=self.root / "with-feedback.docx",
                                      profile="student", save_feedback=True)
        feedback = out.with_suffix(".feedback.html")
        self.assertTrue(feedback.is_file())
        rendered = feedback.read_text(encoding="utf-8")
        for heading in ("改了什么", "APA 来源", "改了原稿哪里", "还要审核"):
            self.assertIn(heading, rendered)
        self.assertNotIn("source_sha256", rendered)
        self.assertIn("原稿和论文文字没有改动", report["feedback"]["summary"])

    def test_heading4_without_boundary_configuration_remains_unchanged(self):
        doc = Document()
        doc.add_paragraph("Participants. Forty adults participated.", style="Heading 4")
        source = self.save(doc)
        before = c14n(Document(source).paragraphs[0]._p)
        out, report = self.format(source)
        self.assertEqual(c14n(Document(out).paragraphs[0]._p), before)
        self.assertTrue(any(e["location"] == "p1" and e["status"] == "review" and "run_in_headings" in e["message"] for e in report["events"]))

    def test_heading4_explicit_prefix_splits_run_without_changing_text(self):
        doc = Document()
        text = "Participants. Forty adults participated."
        doc.add_paragraph(text, style="Heading 4")
        source = self.save(doc)
        out, _ = self.format(source, config={"run_in_headings": {"1": "Participants."}})
        p = Document(out).paragraphs[0]
        self.assertEqual(p.text, text)
        self.assertEqual(p.runs[0].text, "Participants.")
        self.assertTrue(p.runs[0].bold)
        self.assertFalse(p.runs[0].italic)
        self.assertEqual(p.runs[1].text, " Forty adults participated.")
        self.assertFalse(p.runs[1].bold)
        self.assertEqual(p.style.name, "Normal")

    def test_heading4_preserves_deliberate_italic_and_bold_in_following_prose(self):
        doc = Document()
        p = doc.add_paragraph("Results. The statistic ", style="Heading 4")
        p.add_run("p").italic = True
        p.add_run(" was ")
        p.add_run("significant").bold = True
        p.add_run(".")
        source = self.save(doc)
        out, _ = self.format(source, config={"run_in_headings": {"1": "Results."}})
        runs = {run.text: run for run in Document(out).paragraphs[0].runs}
        self.assertTrue(runs["p"].italic, "Formatting a run-in heading must preserve scientific italic in subsequent prose")
        self.assertTrue(runs["significant"].bold, "Formatting a run-in heading must preserve deliberate prose emphasis")

    def test_merged_and_nested_table_structure_and_format_are_preserved(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2)).text = "Merged group heading"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "Vertical group"
        table.cell(1, 1).text = "A"
        nested = table.cell(2, 2).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested content"
        source = self.save(doc)
        before = c14n(Document(source).tables[0]._tbl)
        out, report = self.format(source)
        self.assertEqual(c14n(Document(out).tables[0]._tbl), before)
        self.assertTrue(any(e["location"] == "table1" and e["status"] == "review" and "合并" in e["message"] for e in report["events"]))

    def test_image_width_uses_actual_section_and_custom_headers_survive(self):
        doc = Document()
        first = doc.sections[0]
        first.page_width, first.page_height = Inches(8.5), Inches(11)
        first.header.paragraphs[0].text = "Research group header"
        first.footer.paragraphs[0].text = "Confidential review copy"
        doc.add_paragraph().add_run().add_picture(io.BytesIO(tiny_png()), width=Inches(12))
        second = doc.add_section(WD_SECTION_START.NEW_PAGE)
        second.orientation = WD_ORIENT.LANDSCAPE
        second.page_width, second.page_height = Inches(11), Inches(8.5)
        doc.add_paragraph().add_run().add_picture(io.BytesIO(tiny_png()), width=Inches(12))
        source = self.save(doc)
        before = Document(source)
        header_before = c14n(before.sections[0].header._element)
        footer_before = c14n(before.sections[0].footer._element)
        media_before = apa.package_payloads(source)
        out, _ = self.format(source)
        after = Document(out)
        self.assertEqual(len(after.sections), 2)
        self.assertEqual([(s.page_width, s.page_height, s.orientation) for s in after.sections],
                         [(s.page_width, s.page_height, s.orientation) for s in before.sections])
        self.assertEqual([s.width for s in after.inline_shapes], [Inches(6.5), Inches(9)])
        self.assertEqual([s.height for s in after.inline_shapes], [Inches(6.5), Inches(9)])
        self.assertEqual(c14n(after.sections[0].header._element), header_before)
        self.assertEqual(c14n(after.sections[0].footer._element), footer_before)
        self.assertTrue(after.sections[1].header.is_linked_to_previous)
        self.assertEqual(apa.package_payloads(out), media_before)

    def test_references_label_without_page_boundary_does_not_turn_body_into_title_metadata(self):
        doc = Document()
        doc.add_paragraph("A Study of Language", style="Title")
        doc.add_paragraph("Jane Researcher")
        doc.add_paragraph("This is ordinary research prose that must remain a body paragraph.")
        doc.add_paragraph("References")
        doc.add_paragraph("Researcher, J. (2024). Example publication.")
        source = self.save(doc)
        out, report = self.format(source)
        self.assertEqual(report["paragraph_roles"]["3"], "body")
        self.assertNotIn("title_meta", report["paragraph_roles"].values())
        body = Document(out).paragraphs[2]
        self.assertEqual(body.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(body.paragraph_format.first_line_indent, Inches(0.5))
        self.assertEqual(report["paragraph_roles"]["5"], "reference")

    def test_abstract_state_ends_at_explicit_page_boundary(self):
        for boundary in ("paragraph_before", "run_break", "separate_break", "section_break"):
            with self.subTest(boundary=boundary):
                doc = Document()
                doc.add_paragraph("Abstract")
                abstract = doc.add_paragraph("This paragraph summarizes the study.")
                if boundary == "run_break":
                    abstract.add_run().add_break(WD_BREAK.PAGE)
                elif boundary == "separate_break":
                    doc.add_page_break()
                elif boundary == "section_break":
                    doc.add_section(WD_SECTION_START.NEW_PAGE)
                body = doc.add_paragraph("This is the first ordinary paragraph of the paper.")
                if boundary == "paragraph_before":
                    body.paragraph_format.page_break_before = True
                body_number = str(len(doc.paragraphs))
                source = self.save(doc, name=boundary + "_source.docx")
                out, report = self.format(source, name=boundary + "_formatted.docx")
                self.assertEqual(report["paragraph_roles"]["2"], "abstract")
                self.assertEqual(report["paragraph_roles"][body_number], "body")
                after = Document(out)
                self.assertEqual(after.paragraphs[1].paragraph_format.first_line_indent, 0)
                self.assertEqual(after.paragraphs[-1].paragraph_format.first_line_indent, Inches(0.5))

    def test_final_note_after_table_is_classified_and_formatted_as_note(self):
        doc = Document()
        doc.add_paragraph("Table 1")
        doc.add_paragraph("Descriptive Statistics")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Measure"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "N"
        table.cell(1, 1).text = "40"
        doc.add_paragraph("Note. Values refer to complete cases.")
        source = self.save(doc)
        out, report = self.format(source)
        self.assertEqual(report["paragraph_roles"]["3"], "note")
        note = Document(out).paragraphs[-1]
        self.assertEqual(note.text, "Note. Values refer to complete cases.")
        self.assertEqual(note.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(note.paragraph_format.first_line_indent, 0)
        self.assertEqual(note.paragraph_format.line_spacing, 2)

    def test_professional_linked_header_right_tab_adapts_to_landscape_section(self):
        doc = Document()
        first = doc.sections[0]
        first.page_width, first.page_height = Inches(8.5), Inches(11)
        doc.add_paragraph("Portrait section.")
        second = doc.add_section(WD_SECTION_START.NEW_PAGE)
        second.orientation = WD_ORIENT.LANDSCAPE
        second.page_width, second.page_height = Inches(11), Inches(8.5)
        doc.add_paragraph("Landscape section.")
        self.assertTrue(second.header.is_linked_to_previous)
        source = self.save(doc)
        out, _ = self.format(source, profile="professional", running_head="Section widths")
        after = Document(out)
        for index, expected_width in enumerate((Inches(6.5), Inches(9))):
            section = after.sections[index]
            for header in (section.header, section.first_page_header, section.even_page_header):
                with self.subTest(section=index, header=type(header).__name__):
                    self.assertEqual(len(header.paragraphs), 1)
                    p = header.paragraphs[0]
                    self.assertIn("SECTION WIDTHS", p.text)
                    tabs = list(p.paragraph_format.tab_stops)
                    self.assertEqual(len(tabs), 1)
                    self.assertEqual(tabs[0].alignment, WD_TAB_ALIGNMENT.RIGHT)
                    self.assertEqual(tabs[0].position, expected_width)
        self.assertFalse(after.sections[1].header.is_linked_to_previous)
        counts = page_fields(out)
        self.assertEqual(len(counts), 6)
        self.assertTrue(all(count == 1 for count in counts.values()), counts)
        again, _ = self.format(out, name="rerun.docx", profile="professional", running_head="Section widths")
        self.assertEqual(page_fields(again), counts)


if __name__ == "__main__":
    unittest.main()

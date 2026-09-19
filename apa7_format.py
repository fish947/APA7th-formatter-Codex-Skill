#!/usr/bin/env python3
"""APA 7 Word formatter with concise, evidence-linked feedback, v0.14.0.

Python >= 3.10; pip install 'python-docx>=1.2,<2'
Run without arguments for a local file-picker GUI, or:
  python apa7_format.py paper.docx --profile student
  python apa7_format.py paper.docx --profile professional --running-head 'SHORT TITLE'
  python apa7_format.py paper.docx --inspect

This is a formatting assistant, not a certification of full APA compliance.
Official rules were read on 2026-09-12. No document is uploaded. Source is never
overwritten. Legacy .doc requires a separately installed LibreOffice converter.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import html
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from urllib.parse import urlsplit
from zipfile import ZipFile

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Inches, Pt, RGBColor
    from docx.text.run import Run
    from lxml import etree
except ImportError:
    raise SystemExit("缺少依赖。请先运行：python3 -m pip install 'python-docx>=1.2,<2'")

VERSION = "0.14.0"
BASE = "https://apastyle.apa.org/style-grammar-guidelines/"
# Short verbatim excerpts, each <=25 words per source. The linked page carries
# the full rule and its exceptions; implementation summaries are our paraphrases.
SOURCES = {
    "scope": {"title": "Paper Format", "path": "paper-format",
              "quote": "Follow the guidelines of your institution or publisher to adapt APA Style formatting guidelines as needed.",
              "rule": "学校、课程和期刊的明确要求可要求调整 APA 格式。"},
    "font": {"title": "Font", "path": "paper-format/font",
             "quote": "A variety of fonts are permitted in APA Style papers.",
             "rule": "正文允许多种易读字体；本工具提供官网列出的字体组合。图内字体、代码及脚注有例外。"},
    "margins": {"title": "Margins", "path": "paper-format/margins",
                "quote": "Use 1-inch margins on every side of the page for an APA Style paper.",
                "rule": "四边 1 英寸；毕业论文可能另有装订要求。"},
    "spacing": {"title": "Line Spacing", "path": "paper-format/line-spacing",
                "quote": "Do not add extra space before or after paragraphs.",
                "rule": "正文、摘要、块引用、参考文献及图表编号标题注释双倍行距。表内允许 1、1.5 或 2 倍；脚注、公式有例外。"},
    "paragraph": {"title": "Paragraph Alignment and Indentation", "path": "paper-format/paragraph-format",
                  "quote": "Indent the first line of each paragraph of text 0.5 in. from the left margin.",
                  "rule": "正文左对齐、首行缩进 0.5 英寸；摘要首行不缩进，参考文献悬挂缩进，图表说明左齐。"},
    "header": {"title": "Page Header", "path": "paper-format/page-header",
               "quote": "Insert page numbers in the top right corner. The page number should show on all pages.",
               "rule": "标题页为第 1 页。学生通常仅页码；投稿另需左齐、全大写且含空格不超过 50 字符的短标题。"},
    "title": {"title": "Title Page Setup", "path": "paper-format/title-page",
              "quote": "Place the title three to four lines down from the top of the title page.",
              "rule": "标题居中加粗，与作者间空一双倍行。学生需作者、单位、课程、教师、日期；投稿需对应单位及作者注。"},
    "student_title_elements": {
        "title": "A Step-by-Step Guide for APA Style Student Papers",
        "url": "https://www.apa.org/ed/precollege/psn/2020/09/apa-style-student-papers",
        "quote": "Unless instructed otherwise, students should use the student title page format and include the following elements, in the order listed, on the title page:",
        "rule": "学生标题页依次核对论文标题、作者、单位、课程、教师、截止日期和页码；学校要求优先。",
        "checked_on": "2026-09-18",
    },
    "headings": {"title": "Headings", "path": "paper-format/headings",
                 "quote": "Do not label headings with numbers or letters.",
                 "rule": "一级居中粗体；二级左齐粗体；三级左齐粗斜体；四级缩进粗体、五级缩进粗斜体，后两级句号后同段续正文。"},
    "case": {"title": "Title Case Capitalization", "path": "capitalization/title-case",
             "quote": "In title case, major words are capitalized, and most minor words are lowercase.",
             "rule": "论文标题、各级标题、图表标题使用 title case；词性、专名、符号及缩写需结合内容判断。"},
    "references": {"title": "Reference List Setup", "path": "paper-format/reference-list",
                   "quote": "Type each reference as a single paragraph, justified to the left margin.",
                   "rule": "新页居中粗体 References；条目双倍行距、0.5 英寸悬挂缩进。排序及正文引用对应须核对，个人通信等有例外。"},
    "reference_order": {"title": "Reference List Setup: Alphabetical Order",
                        "url": "https://apastyle.apa.org/style-grammar-guidelines/paper-format/reference-list",
                        "quote": "Alphabetize references according to the first word of the reference (usually the last name of the first author).",
                        "rule": "参考文献通常按第一个词（多为第一作者姓氏）的字母顺序排列；本工具只对可靠解析的拉丁字母作者条目做保守提示。",
                        "checked_on": "2026-09-18"},
    "same_author_date": {"title": "Citing Works With the Same Author and Date",
                         "path": "citations/basic-principles/same-year-author",
                         "quote": "When multiple references have an identical author (or authors) and publication year, include a lowercase letter after the year.",
                         "rule": "完全相同的作者组合和出版年份需使用 a、b 等年份后缀，并在正文引用中保持一致。",
                         "checked_on": "2026-09-18"},
    "citation_match": {"title": "Author-Date Citation System", "path": "citations/basic-principles/author-date",
                       "quote": "Each work cited must appear in the reference list, and each work in the reference list must be cited in the text.",
                       "rule": "正文引文与参考文献表应相互对应；本工具只做作者—年份的轻量提示，不自动增删文献。"},
    "reference_links": {"title": "DOIs and URLs", "path": "references/dois-urls",
                        "quote": "Present both DOIs and URLs as hyperlinks (i.e., beginning with “http://” or “https://”).",
                        "rule": "参考文献中已有的完整 DOI／URL 在电子文档中保留为可点击链接；显示可使用普通黑色文字。",
                        "checked_on": "2026-09-14"},
    "tables": {"title": "Table Setup", "path": "tables-figures/tables",
               "quote": "Do not use vertical borders to separate data, and do not use borders around every cell in a table.",
               "rule": "编号粗体、标题另行斜体，均位于表上方。表头居中；首列正文左齐。通常保留顶线、底线及表头下线；复杂表可有必要横线。"},
    "figures": {"title": "Figure Setup", "path": "tables-figures/figures",
                "quote": "Number figures in the order in which they are mentioned in your paper.",
                "rule": "编号粗体、标题另行斜体且在图上方，必要注释在下方。应检查清晰度、坐标、图例和正文首次提及位置。"},
    "appendices": {"title": "Appendices Setup", "path": "paper-format/appendices",
                   "quote": "Write the appendix label at the top of the page in bold and centered.",
                   "rule": "每个附录另起一页，标签和下一行标题均居中粗体；附录图表有独立编号规则及单个图表例外。"},
    "quotations": {"title": "Quotations", "path": "citations/quotations",
                   "quote": "Format quotations of 40 words or more as block quotations:",
                   "rule": "40 词及以上原文引用用块引用，整体左缩进 0.5 英寸、双倍行距；同一引文后续段首行再缩进 0.5 英寸。"},
    "statistics": {"title": "Numbers and Statistics Guide",
                   "url": "https://apastyle.apa.org/instructional-aids/numbers-statistics-guide.pdf",
                   "quote": "Report exact p values to two or three decimals (e.g., p = .006, p = .03).",
                   "rule": "统计符号、运算符空格、前导零和 p 值按 APA 呈现；不自动舍入、重算或改变任何研究结果。",
                   "checked_on": "2026-09-13"},
    "equations": {"title": "Publication Manual: Presentation of Equations",
                  "url": "https://www.apa.org/pubs/books/publication-manual-7th-edition-spiral",
                  "quote": "Number all displayed equations consecutively, with the number in parentheses near the right margin of the page.",
                  "rule": "识别原生 Word 公式和纯文本公式；安全处理公式段落版式，编号、标点、变量与公式含义仍需审核。",
                  "checked_on": "2026-09-13"},
    "jars": {"title": "APA Research Transparency Standards",
             "url": "https://www.apa.org/pubs/journals/resources/standards-disclosures",
             "quote": "Exact p values, effect sizes, and 95% confidence intervals, or an explanation of why this was not possible.",
             "rule": "专业／投稿论文应结合具体期刊要求核对精确 p 值、效应量和置信区间。",
             "checked_on": "2026-09-13"},
}
for _source in SOURCES.values():
    if "url" not in _source:
        _source["url"] = BASE + _source.pop("path")
    _source.setdefault("checked_on", "2026-09-12")

FONTS = {"Times New Roman": 12, "Arial": 11, "Calibri": 11,
         "Georgia": 11, "Lucida Sans Unicode": 10, "Aptos": 12}
CAPTION = re.compile(r"^(Table|Figure)\s+([A-Z]?\d+)\s*$", re.I)
COMBINED_CAPTION = re.compile(r"^(Table|Figure)\s+[A-Z]?\d+[.:：\s]+\S", re.I)
NUMBERED_CALLOUT = re.compile(
    r"\b(?P<kind>Tables?|Figures?|Equations?)\s+"
    r"(?P<identifiers>\(?[A-Z]?\d+\)?(?:\s*(?:,\s*(?:and\s+)?|and\s+|&\s+|to\s+|through\s+|[-–—]\s*)\(?[A-Z]?\d+\)?)*)",
    re.I,
)
NUMBERED_IDENTIFIER = re.compile(r"[A-Z]?\d+", re.I)
SECTION_LABELS = {"abstract", "references", "reference", "author note", "footnotes"}
ROLES = {"body", "title", "title_meta", "section", "abstract", "reference",
         "heading1", "heading2", "heading3", "heading4", "heading5",
         "caption_number", "caption_title", "note", "quote", "quote_continuation",
         "appendix", "appendix_title", "keywords", "equation", "preserve"}

APA7_PARAGRAPH_STYLES = {
    "APA7 Body": {"indent": 0.5, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": False, "keep": False},
    "APA7 Title": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.CENTER, "bold": True, "italic": False, "keep": True},
    "APA7 Title Metadata": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.CENTER, "bold": False, "italic": False, "keep": False},
    "APA7 Abstract": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": False, "keep": False},
    "APA7 Keywords": {"indent": 0.5, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": False, "keep": False},
    "APA7 Heading 1": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.CENTER, "bold": True, "italic": False, "keep": True},
    "APA7 Heading 2": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": True, "italic": False, "keep": True},
    "APA7 Heading 3": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": True, "italic": True, "keep": True},
    "APA7 Reference": {"indent": -0.5, "left": 0.5, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": False, "keep": False},
    "APA7 Caption Number": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": True, "italic": False, "keep": True},
    "APA7 Caption Title": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": True, "keep": True},
    "APA7 Note": {"indent": 0, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": False, "keep": False},
    "APA7 Block Quote": {"indent": 0, "left": 0.5, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": False, "keep": False},
    "APA7 Run-in Heading": {"indent": 0.5, "left": 0, "align": WD_ALIGN_PARAGRAPH.LEFT, "bold": False, "italic": False, "keep": False},
}
APA7_CHARACTER_STYLES = {
    "APA7 Level 4 Prefix": {"bold": True, "italic": False},
    "APA7 Level 5 Prefix": {"bold": True, "italic": True},
}
ROLE_STYLE_NAMES = {
    "body": "APA7 Body", "title": "APA7 Title", "title_meta": "APA7 Title Metadata",
    "abstract": "APA7 Abstract", "keywords": "APA7 Keywords", "section": "APA7 Heading 1",
    "appendix": "APA7 Heading 1", "appendix_title": "APA7 Heading 1",
    "heading1": "APA7 Heading 1", "heading2": "APA7 Heading 2", "heading3": "APA7 Heading 3",
    "reference": "APA7 Reference", "caption_number": "APA7 Caption Number",
    "caption_title": "APA7 Caption Title", "note": "APA7 Note",
    "quote": "APA7 Block Quote", "quote_continuation": "APA7 Block Quote",
}
STYLE_ROLE_HINTS = {
    "APA7 Body": "body", "APA7 Title": "title", "APA7 Title Metadata": "title_meta",
    "APA7 Abstract": "abstract", "APA7 Keywords": "keywords",
    "APA7 Heading 1": "heading1", "APA7 Heading 2": "heading2", "APA7 Heading 3": "heading3",
    "APA7 Reference": "reference", "APA7 Caption Number": "caption_number",
    "APA7 Caption Title": "caption_title", "APA7 Note": "note", "APA7 Block Quote": "quote",
}

_YEAR = r"(?:1[5-9]\d{2}|20\d{2})[a-z]?|n\.d\."
_PARENTHETICAL_CITATION = re.compile(r"(?P<author>[^,;()]{1,120}?),\s*(?P<year>" + _YEAR + r")", re.I)
_NARRATIVE_CITATION = re.compile(
    r"\b(?P<author>[A-Z][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+(?:\s+(?:et\s+al\.|and\s+[A-Z][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+|&\s*[A-Z][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+|[A-Z][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+)){0,4})"
    r"\s*\((?P<year>" + _YEAR + r")\)")
_REFERENCE_YEAR = re.compile(r"\((?P<year>" + _YEAR + r")\)", re.I)
_REFERENCE_DATE = re.compile(
    r"\((?P<year>(?P<base>(?:1[5-9]\d{2}|20\d{2}|n\.d\.))(?P<suffix>-?[a-z])?)\)", re.I
)
_ANY_DOI = re.compile(
    r"(?<![\w/])(?:https?://(?:dx\.)?doi\.org/|doi\s*:\s*)?(?P<doi>10\.\d{4,9}/[-._;()/:A-Z0-9]+)",
    re.I,
)
_WEB_URL = re.compile(r"https?://[^\s<>\"“”'{}]+", re.I)
_BARE_DOI = re.compile(r"(?<!doi\.org/)(?<![\w/])(?:doi\s*:\s*)?(10\.\d{4,9}/[-._;()/:A-Z0-9]+)", re.I)


def digest(path: Path) -> str:
    with path.open("rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest() if hasattr(hashlib, "file_digest") else hashlib.sha256(f.read()).hexdigest()


def visible_text(p) -> str:
    # Includes hyperlink/field results without invoking the destructive .text setter.
    return "".join(p._p.xpath(".//w:t/text()"))


def descendant_count(element, *names) -> int:
    """Count qualified descendants without relying on an element's XPath class.

    Third-party DOCX generators can place ordinary lxml elements such as content
    controls or AlternateContent containers between python-docx elements.  Their
    ``xpath`` method does not install python-docx's namespace map, so prefixed
    expressions can fail even when the document itself declares the prefix.
    Iterating Clark-notation tags works for both element types.
    """
    tags = {qn(name) for name in names}
    return sum(node.tag in tags for node in element.iter())


def has_descendant(element, *names) -> bool:
    return descendant_count(element, *names) > 0


def has_page_boundary(element) -> bool:
    """Detect a page/section break on python-docx or raw lxml elements."""
    if element is None:
        return False
    for node in element.iter():
        if node.tag == qn("w:sectPr"):
            return True
        if node.tag == qn("w:br") and node.get(qn("w:type")) == "page":
            return True
    return False


def descendant_text(element, name="w:t") -> str:
    tag = qn(name)
    return "".join(node.text or "" for node in element.iter(tag))


def section_text_width(section):
    """Return usable width without forcing an implicit Word paper size.

    Some generators omit ``w:pgSz`` and rely on the word processor's default.
    Keep that omission in the document, but use Letter width as a conservative
    calculation fallback for header tabs and image fitting. The ambiguity is
    recorded for visual review by ``Formatter.sections``.
    """
    page_width = section.page_width if section.page_width is not None else Inches(8.5)
    left = section.left_margin if section.left_margin is not None else Inches(1)
    right = section.right_margin if section.right_margin is not None else Inches(1)
    return page_width - left - right


def child(parent, name, attrs=None):
    node = parent.find(qn(name))
    if node is None:
        node = OxmlElement(name)
        parent.append(node)
    for k, v in (attrs or {}).items():
        node.set(qn(k), str(v))
    return node


def remove(parent, *names):
    for name in names:
        for node in list(parent.findall(qn(name))):
            parent.remove(node)


def all_runs(p):
    for r in p._p.xpath(".//w:r"):
        # Textboxes and drawings are not paragraph text; leave them intact.
        if r.xpath("ancestor::w:txbxContent") or r.xpath(".//w:drawing | .//w:object | .//w:pict"):
            continue
        yield Run(r, p)


def split_plain_run_at(p, boundary):
    """Split only plain text runs, preserving all run properties and characters."""
    if p._p.xpath("./*[not(self::w:pPr or self::w:r)]") or p._p.xpath("./w:r/*[not(self::w:rPr or self::w:t)]"):
        return False
    offset = 0
    for run in list(p.runs):
        text = run.text
        cut = boundary - offset
        if 0 < cut < len(text):
            new = copy.deepcopy(run._r)
            for elem in list(new):
                if elem.tag != qn("w:rPr"):
                    new.remove(elem)
            child(new, "w:t", {"xml:space": "preserve"}).text = text[cut:]
            for elem in list(run._r):
                if elem.tag != qn("w:rPr"):
                    run._r.remove(elem)
            child(run._r, "w:t", {"xml:space": "preserve"}).text = text[:cut]
            run._r.addnext(new)
            break
        offset += len(text)
    return True


def typography(p, font, *, bold=None, italic=None):
    for run in all_runs(p):
        # Preserve explicitly selected symbol/code fonts.
        if run.font.name in {"Symbol", "Wingdings", "Cambria Math", "Courier New", "Consolas"}:
            continue
        run.font.name = font
        run.font.size = Pt(FONTS[font])
        fonts = run._r.get_or_add_rPr().rFonts
        for a in ("asciiTheme", "hAnsiTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + a), None)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        if bold is not None or italic is not None:
            run.underline = False
            run.font.all_caps = False
            run.font.small_caps = False


def paragraph_format(p, *, indent=0.5, left=0, align=WD_ALIGN_PARAGRAPH.LEFT,
                     spacing=2.0, keep=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.left_indent, pf.right_indent = Inches(left), Inches(0)
    pf.first_line_indent = Inches(indent)
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = spacing
    pf.keep_with_next = keep
    pf.keep_together = False
    pf.widow_control = True
    pr = p._p.get_or_add_pPr()
    # Word's grid and character-based indents can override point values.
    ind = pr.find(qn("w:ind"))
    for attr in ("firstLineChars", "hangingChars", "leftChars", "rightChars", "start", "end"):
        ind.attrib.pop(qn("w:" + attr), None)
    sp = pr.find(qn("w:spacing"))
    for attr in ("beforeAutospacing", "afterAutospacing", "beforeLines", "afterLines"):
        sp.attrib.pop(qn("w:" + attr), None)
    child(pr, "w:snapToGrid", {"w:val": "0"})
    child(pr, "w:contextualSpacing", {"w:val": "0"})


def _configure_style_font(style, font, *, bold, italic):
    style.font.name = font
    style.font.size = Pt(FONTS[font])
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    style.font.all_caps = False
    style.font.small_caps = False
    fonts = style.element.get_or_add_rPr().rFonts
    for attr in ("asciiTheme", "hAnsiTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)


def _configure_paragraph_style(style, font, spec):
    _configure_style_font(style, font, bold=spec["bold"], italic=spec["italic"])
    pf = style.paragraph_format
    pf.alignment = spec["align"]
    pf.left_indent, pf.right_indent = Inches(spec["left"]), Inches(0)
    pf.first_line_indent = Inches(spec["indent"])
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 2.0
    pf.keep_with_next = spec["keep"]
    pf.keep_together = False
    pf.widow_control = True
    pr = style.element.get_or_add_pPr()
    child(pr, "w:snapToGrid", {"w:val": "0"})
    child(pr, "w:contextualSpacing", {"w:val": "0"})


def _author_key(author):
    """Return a conservative first-author/group-author key for a review hint."""
    if re.search(r"[=<>]", author):
        return ""
    author = re.sub(r"^\s*(?:see|e\.g\.,?|cf\.)\s+", "", author, flags=re.I)
    author = re.sub(r"\s+et\s+al\.?\s*$", "", author, flags=re.I)
    author = re.split(r"\s+(?:&|and)\s+|,", author, maxsplit=1, flags=re.I)[0]
    author = re.sub(r"[^\w'’.-]+", " ", author, flags=re.UNICODE).strip(" .-'’_")
    author = re.sub(r"\s+", " ", author)
    return author.casefold() if any(character.isalpha() for character in author) else ""


def citation_reference_check(paragraphs, roles):
    """Lightweight author-year correspondence check; never edits bibliography text."""
    citations, references, unparsed = {}, {}, []
    for index, paragraph in enumerate(paragraphs):
        text = visible_text(paragraph).strip()
        if not text:
            continue
        role = roles.get(index, "body")
        if role == "reference":
            year_match = _REFERENCE_YEAR.search(text)
            if not year_match:
                unparsed.append(index + 1)
                continue
            author_display = text[:year_match.start()].strip()
            key = (_author_key(author_display), year_match.group("year").casefold())
            if not key[0]:
                unparsed.append(index + 1)
                continue
            record = references.setdefault(key, {"author": author_display, "year": year_match.group("year"), "paragraphs": []})
            record["paragraphs"].append(index + 1)
            continue
        if role in {"title", "title_meta", "caption_number", "caption_title", "appendix", "appendix_title"}:
            continue
        for group in re.findall(r"\(([^()]*)\)", text):
            for match in _PARENTHETICAL_CITATION.finditer(group):
                author_display, year = match.group("author").strip(), match.group("year")
                key = (_author_key(author_display), year.casefold())
                if key[0]:
                    record = citations.setdefault(key, {"author": author_display, "year": year, "paragraphs": []})
                    if index + 1 not in record["paragraphs"]:
                        record["paragraphs"].append(index + 1)
        without_parentheses = re.sub(r"\([^()]*\)", "", text)
        for match in _NARRATIVE_CITATION.finditer(text):
            # The match itself contains the year parentheses, so reject only if
            # the author text disappeared while stripping other parentheticals.
            if match.group("author") not in without_parentheses:
                continue
            author_display, year = match.group("author").strip(), match.group("year")
            key = (_author_key(author_display), year.casefold())
            if key[0]:
                record = citations.setdefault(key, {"author": author_display, "year": year, "paragraphs": []})
                if index + 1 not in record["paragraphs"]:
                    record["paragraphs"].append(index + 1)
    unmatched = [value for key, value in citations.items() if key not in references]
    uncited = [value for key, value in references.items() if key not in citations]
    if not citations and not references and not unparsed:
        status = "not_applicable"
    elif unmatched or uncited or unparsed:
        status = "needs_review"
    else:
        status = "passed"
    return {"status": status, "citations_found": len(citations), "reference_entries": len(references) + len(unparsed),
            "unmatched_citations": unmatched, "uncited_references": uncited,
            "unparsed_reference_paragraphs": unparsed,
            "limitation": "Author-year pattern check only; group authors, translated works, secondary citations, personal communications and bibliographic facts still require review.",
            "source_url": SOURCES["citation_match"]["url"]}


def _reference_sort_key(author):
    """Return a conservative Latin-script key, or None when collation is uncertain."""
    letters = [character for character in author if character.isalpha()]
    if not letters or any("LATIN" not in unicodedata.name(character, "") for character in letters):
        return None
    normalized = unicodedata.normalize("NFKD", author).casefold()
    normalized = "".join(character for character in normalized if not unicodedata.combining(character))
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized).strip()
    return normalized or None


def _canonical_doi(text):
    match = _ANY_DOI.search(text)
    if not match:
        return None
    doi = match.group("doi").rstrip(".,;:!?")
    while doi.endswith(")") and doi.count("(") < doi.count(")"):
        doi = doi[:-1]
    return doi.casefold()


def reference_quality_check(paragraphs, roles):
    """Audit reference-list order and duplicates without moving or rewriting entries."""
    entries, issues = [], []

    def issue(code, message, records):
        issues.append({
            "code": code,
            "message": message,
            "paragraphs": [record["paragraph"] for record in records],
            "locations": [record["location"] for record in records],
        })

    for index, paragraph in enumerate(paragraphs):
        if roles.get(index) != "reference":
            continue
        text = visible_text(paragraph).strip()
        record = {
            "paragraph": index + 1,
            "location": f"p{index + 1}",
            "text": text,
            "normalized_text": re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip().casefold(),
            "doi": _canonical_doi(text),
            "author": None,
            "author_key": None,
            "sort_key": None,
            "base_year": None,
            "suffix": "",
        }
        date = _REFERENCE_DATE.search(text)
        if date:
            author = text[:date.start()].strip()
            record["author"] = author
            record["author_key"] = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", author)).strip().casefold()
            record["sort_key"] = _reference_sort_key(author)
            record["base_year"] = date.group("base").casefold()
            record["suffix"] = (date.group("suffix") or "").lstrip("-").casefold()
        else:
            issue(
                "unparsed_reference",
                f"第 {index + 1} 段未识别出清晰的作者和年份，无法检查排序或同作者同年后缀。",
                [record],
            )
        entries.append(record)

    by_text, by_doi = {}, {}
    for record in entries:
        if record["normalized_text"]:
            by_text.setdefault(record["normalized_text"], []).append(record)
        if record["doi"]:
            by_doi.setdefault(record["doi"], []).append(record)
    for records in by_text.values():
        if len(records) > 1:
            issue(
                "duplicate_entry",
                f"第 {_ranges(record['paragraph'] for record in records)} 段内容相同，可能是重复参考文献。",
                records,
            )
    for doi, records in by_doi.items():
        if len(records) > 1:
            issue(
                "duplicate_doi",
                f"第 {_ranges(record['paragraph'] for record in records)} 段使用了相同 DOI（{doi}）。",
                records,
            )

    sortable = [record for record in entries if record["sort_key"]]
    for previous, current in zip(sortable, sortable[1:]):
        if previous["sort_key"] > current["sort_key"]:
            issue(
                "alphabetical_order",
                f"第 {previous['paragraph']} 段排在第 {current['paragraph']} 段之前，但可识别的作者字母顺序相反。",
                [previous, current],
            )

    author_year_groups = {}
    for record in entries:
        if record["author_key"] and record["base_year"]:
            author_year_groups.setdefault((record["author_key"], record["base_year"]), []).append(record)
    repeated_groups = 0
    for records in author_year_groups.values():
        if len(records) < 2:
            continue
        repeated_groups += 1
        expected = [chr(ord("a") + number) for number in range(len(records))]
        actual = [record["suffix"] for record in records]
        if actual != expected:
            issue(
                "same_author_year_suffix",
                f"第 {_ranges(record['paragraph'] for record in records)} 段为同作者同年份条目，年份后缀应按参考文献表顺序连续使用 a、b 等字母。",
                records,
            )

    status = "not_applicable" if not entries else "needs_review" if issues else "passed"
    return {
        "status": status,
        "reference_entries_checked": len(entries),
        "same_author_year_groups": repeated_groups,
        "issues": issues,
        "source_urls": {
            "alphabetical_order": SOURCES["reference_order"]["url"],
            "same_author_year": SOURCES["same_author_date"]["url"],
        },
        "limitation": (
            "Conservative pattern check only. It does not move or rewrite entries, validate bibliographic facts, "
            "classify source types, verify title capitalization or italics, collate non-Latin scripts, or prove "
            "that same-year letters follow title order."
        ),
    }


def front_matter_check(paragraphs, roles, profile, cover=None):
    """Check whether profile-specific front-matter candidates are present; never invent fields."""
    issues = []

    def issue(code, message, paragraph_numbers=()):
        issues.append({
            "code": code,
            "message": message,
            "paragraphs": list(paragraph_numbers),
            "locations": [f"p{number}" for number in paragraph_numbers],
        })

    if cover:
        start, end = cover
        cover_numbers = list(range(start + 1, end + 2))
        nonblank_cover = [index for index in range(start, end + 1)
                          if visible_text(paragraphs[index]).strip()]
    else:
        cover_numbers, nonblank_cover = [], []
        issue("unconfirmed_title_page", "没有确认标题页范围；不会猜测或补写作者与课程／投稿信息。")

    title_numbers = [index + 1 for index in nonblank_cover if roles.get(index) == "title"]
    metadata_numbers = [index + 1 for index in nonblank_cover if roles.get(index) == "title_meta"]
    if cover and len(title_numbers) != 1:
        issue("title_candidate_count", "标题页需要一个明确的论文标题候选。", cover_numbers)

    if cover and profile == "student" and len(metadata_numbers) < 5:
        issue(
            "student_title_fields",
            "学生标题页没有识别到足够的作者、单位、课程、教师和截止日期候选行；请逐项核对。",
            cover_numbers,
        )
    if cover and profile == "professional" and len(metadata_numbers) < 2:
        issue(
            "professional_title_fields",
            "专业论文标题页没有识别到足够的作者和单位候选行；作者注及期刊要求仍需核对。",
            cover_numbers,
        )

    abstract_labels = [index + 1 for index, paragraph in enumerate(paragraphs)
                       if visible_text(paragraph).strip().casefold() == "abstract"]
    keyword_numbers = [index + 1 for index, role in roles.items() if role == "keywords"]
    if len(abstract_labels) > 1:
        issue("multiple_abstract_labels", "识别到多个 Abstract 标签，需要确认哪一个属于正式摘要。", abstract_labels)
    if keyword_numbers and not abstract_labels:
        issue("keywords_without_abstract", "识别到关键词，但没有找到明确的 Abstract 标签。", keyword_numbers)

    outside_cover_titles = []
    for index, role in roles.items():
        if role != "title" or (cover and cover[0] <= index <= cover[1]):
            continue
        outside_cover_titles.append(index + 1)
    body_present = any(role in {"body", "heading1", "heading2", "heading3", "heading4", "heading5"}
                       and visible_text(paragraphs[index]).strip() for index, role in roles.items())
    if cover and body_present and not outside_cover_titles:
        issue(
            "repeated_title_not_identified",
            "已确认标题页，但没有在正文首页识别到重复的论文标题；需结合分页确认。",
        )

    return {
        "status": "needs_review" if issues else "passed",
        "profile": profile,
        "title_page_confirmed": bool(cover),
        "title_page_paragraphs": cover_numbers,
        "title_candidates": title_numbers,
        "metadata_candidates": metadata_numbers,
        "abstract_labels": abstract_labels,
        "keyword_paragraphs": keyword_numbers,
        "repeated_title_candidates": outside_cover_titles,
        "required_title_elements": (
            ["paper title", "author", "affiliation", "course", "instructor", "due date", "page number"]
            if profile == "student" else
            ["paper title", "author byline", "author-affiliation correspondence", "page number", "running head"]
        ),
        "issues": issues,
        "source_urls": {
            "title_page": SOURCES["title"]["url"],
            "student_elements": SOURCES["student_title_elements"]["url"],
        },
        "limitation": (
            "Candidate-count and label check only. It cannot prove that a line contains the correct person, "
            "institution, course, date, author note, or journal-required wording."
        ),
    }


def _identifier_parts(identifier):
    match = re.fullmatch(r"([A-Z]?)(\d+)", identifier.upper())
    return (match.group(1), int(match.group(2))) if match else ("", -1)


def _expanded_identifiers(text):
    """Read explicit lists/ranges such as 1, 2, and 4 or A1–A3."""
    matches = list(NUMBERED_IDENTIFIER.finditer(text))
    if not matches:
        return []
    identifiers = [matches[0].group(0).upper()]
    for previous, current in zip(matches, matches[1:]):
        current_id = current.group(0).upper()
        connector = text[previous.end():current.start()].casefold()
        start_prefix, start_number = _identifier_parts(previous.group(0))
        end_prefix, end_number = _identifier_parts(current_id)
        is_range = bool(re.search(r"(?:\bto\b|\bthrough\b|[-–—])", connector))
        if is_range and start_prefix == end_prefix and 0 < end_number - start_number <= 100:
            identifiers.extend(f"{start_prefix}{number}" for number in range(start_number + 1, end_number + 1))
        else:
            identifiers.append(current_id)
    return identifiers


def _paragraph_math_text(paragraph):
    """Read normal and native-math text without rewriting the paragraph."""
    return "".join(paragraph._p.xpath(".//w:t/text() | .//m:t/text()"))


def numbered_object_check(doc, roles, table_roles=None):
    """Check table, figure and equation labels/callouts without renumbering them."""
    labels = {"table": [], "figure": [], "equation": []}
    callouts = {"table": [], "figure": [], "equation": []}
    issues = []

    def issue(code, kind, message, *, identifiers=(), locations=()):
        issues.append({"code": code, "kind": kind, "identifiers": list(identifiers),
                       "locations": list(locations), "message": message})

    for index, paragraph in enumerate(doc.paragraphs):
        role = roles.get(index, "body")
        text = visible_text(paragraph).strip()
        location = f"p{index + 1}"
        if role == "caption_number":
            match = CAPTION.fullmatch(text)
            if match:
                kind = match.group(1).casefold()
                labels[kind].append({"identifier": match.group(2).upper(), "paragraph": index + 1,
                                     "location": location, "text": text})
            else:
                issue("invalid_caption_label", "table_or_figure",
                      f"第 {index + 1} 段被标记为图表编号，但没有识别出独立的 Table/Figure 编号。",
                      locations=[location])

        full_text = _paragraph_math_text(paragraph).strip()
        if role == "equation" or paragraph._p.xpath(".//m:oMathPara"):
            number = re.search(r"\(([A-Z]?\d+)\)\s*[.,;:]?\s*$", full_text, re.I)
            if number:
                labels["equation"].append({"identifier": number.group(1).upper(), "paragraph": index + 1,
                                           "location": location, "text": full_text[:180]})

        if role in {"caption_number", "caption_title", "note", "reference", "title", "title_meta",
                    "preserve", "equation"}:
            continue
        seen_here = set()
        for match in NUMBERED_CALLOUT.finditer(text):
            kind = match.group("kind").casefold().rstrip("s")
            for identifier in _expanded_identifiers(match.group("identifiers")):
                key = (kind, identifier)
                if key in seen_here:
                    continue
                seen_here.add(key)
                callouts[kind].append({"identifier": identifier, "paragraph": index + 1,
                                       "location": location, "text": match.group(0)})

    for kind in ("table", "figure", "equation"):
        label_ids = [record["identifier"] for record in labels[kind]]
        callout_ids = [record["identifier"] for record in callouts[kind]]
        counts = Counter(label_ids)
        for identifier, count in counts.items():
            if count > 1:
                locations = [record["location"] for record in labels[kind]
                             if record["identifier"] == identifier]
                issue("duplicate_label", kind,
                      f"{kind.title()} {identifier} 出现了 {count} 次编号标签（{'、'.join(locations)}）。",
                      identifiers=[identifier], locations=locations)

        groups = {}
        for identifier in label_ids:
            prefix, number = _identifier_parts(identifier)
            if number >= 0 and identifier not in groups.setdefault(prefix, []):
                groups[prefix].append(identifier)
        for prefix, identifiers in groups.items():
            numbers = [_identifier_parts(identifier)[1] for identifier in identifiers]
            missing = [number for number in range(1, max(numbers) + 1) if number not in numbers]
            if missing:
                missing_ids = [f"{prefix}{number}" for number in missing]
                issue("sequence_gap", kind,
                      f"{kind.title()} 编号跳过了 {', '.join(missing_ids)}；不会自动重新编号。",
                      identifiers=missing_ids)
            if numbers != sorted(numbers):
                issue("label_order", kind,
                      f"{kind.title()} 编号在文档中的顺序为 {', '.join(identifiers)}，不是递增顺序。",
                      identifiers=identifiers,
                      locations=[record["location"] for record in labels[kind]
                                 if record["identifier"] in identifiers])

        label_set, callout_set = set(label_ids), set(callout_ids)
        for identifier in sorted(callout_set - label_set, key=_identifier_parts):
            locations = sorted({record["location"] for record in callouts[kind]
                                if record["identifier"] == identifier})
            issue("callout_without_label", kind,
                  f"正文提到 {kind.title()} {identifier}（{'、'.join(locations)}），但没有找到对应编号标签。",
                  identifiers=[identifier], locations=locations)
        for identifier in sorted(label_set - callout_set, key=_identifier_parts):
            locations = [record["location"] for record in labels[kind]
                         if record["identifier"] == identifier]
            issue("label_without_callout", kind,
                  f"{kind.title()} {identifier}（{'、'.join(locations)}）未在正文中找到明确提及。",
                  identifiers=[identifier], locations=locations)

        first_mentions = []
        for record in callouts[kind]:
            if record["identifier"] in label_set and record["identifier"] not in first_mentions:
                first_mentions.append(record["identifier"])
        mention_groups = {}
        for identifier in first_mentions:
            prefix, number = _identifier_parts(identifier)
            if number >= 0:
                mention_groups.setdefault(prefix, []).append(identifier)
        for identifiers in mention_groups.values():
            numbers = [_identifier_parts(identifier)[1] for identifier in identifiers]
            if len(numbers) > 1 and numbers != sorted(numbers):
                locations = []
                for identifier in identifiers:
                    locations.append(next(record["location"] for record in callouts[kind]
                                          if record["identifier"] == identifier))
                issue("first_mention_order", kind,
                      f"正文首次提及 {kind.title()} 的顺序为 {', '.join(identifiers)}，需要核对编号。",
                      identifiers=identifiers, locations=locations)

    data_table_count = None if table_roles is None else sum(role == "data" for role in table_roles.values())
    drawing_count = descendant_count(doc._element, "wp:inline", "wp:anchor")
    displayed_equation_count = len(doc._element.xpath(".//m:oMathPara"))
    if data_table_count is not None and data_table_count > len(labels["table"]):
        issue("data_table_without_label", "table",
              f"已确认 {data_table_count} 个数据表，但只识别到 {len(labels['table'])} 个 Table 编号；需要核对缺失的编号和标题。")
    elif table_roles is None and doc.tables and not labels["table"]:
        issue("possible_table_without_label", "table",
              f"检测到 {len(doc.tables)} 个 Word 表格，但没有识别到 Table 编号；请先判断它们是数据表还是布局表。")
    if drawing_count and not labels["figure"]:
        issue("possible_figure_without_label", "figure",
              f"检测到 {drawing_count} 个内嵌或浮动图形对象，但没有识别到 Figure 编号；请判断是否属于需要图题的论文图。")

    detected = sum(len(records) for records in labels.values()) + sum(len(records) for records in callouts.values())
    detected += (data_table_count or 0) + drawing_count + displayed_equation_count
    status = "not_applicable" if not detected and not doc.tables else "needs_review" if issues else "passed"
    return {
        "status": status,
        "labels": labels,
        "callouts": callouts,
        "counts": {
            kind: {"labels": len(labels[kind]), "callouts": len(callouts[kind])}
            for kind in ("table", "figure", "equation")
        },
        "objects": {"top_level_tables": len(doc.tables), "confirmed_data_tables": data_table_count,
                    "drawings": drawing_count, "displayed_equations": displayed_equation_count},
        "issues": issues,
        "source_urls": {kind: SOURCES[kind]["url"] for kind in ("tables", "figures", "equations")},
        "limitation": (
            "Pattern-based label and callout check only. It does not infer whether a drawing is a research figure, "
            "pair every caption to an object, follow field updates, or renumber any object."
        ),
    }


def caption_object_check(doc, roles, table_roles=None, cover=None):
    """Pair top-level tables/drawings with nearby APA number, title, and optional note paragraphs."""
    paragraph_ids = {paragraph._p: index for index, paragraph in enumerate(doc.paragraphs)}
    table_ids = {table._tbl: index for index, table in enumerate(doc.tables, 1)}
    items = []
    for element in doc._element.body:
        if element in paragraph_ids:
            index = paragraph_ids[element]
            drawing_count = descendant_count(element, "wp:inline", "wp:anchor")
            items.append({
                "kind": "paragraph",
                "paragraph": index + 1,
                "role": roles.get(index, "body"),
                "text": visible_text(doc.paragraphs[index]).strip(),
                "drawing_count": drawing_count,
            })
        elif element in table_ids:
            table_number = table_ids[element]
            items.append({
                "kind": "table",
                "table": table_number,
                "role": None if table_roles is None else table_roles.get(str(table_number), "data"),
            })

    issues, objects, used_caption_paragraphs = [], [], set()

    def issue(code, kind, message, locations=()):
        issues.append({"code": code, "kind": kind, "message": message, "locations": list(locations)})

    def nearby_paragraphs(position, direction):
        records = []
        cursor = position + direction
        while 0 <= cursor < len(items) and len(records) < 3:
            item = items[cursor]
            if item["kind"] == "table" or item.get("drawing_count"):
                break
            if item.get("text"):
                records.append(item)
            cursor += direction
        return records if direction > 0 else list(reversed(records))

    for position, item in enumerate(items):
        kind = None
        object_id = None
        confirmed = False
        drawing_count = 0
        if item["kind"] == "table":
            if item.get("role") == "preserve":
                continue
            kind, object_id = "table", f"table{item['table']}"
            confirmed = item.get("role") == "data"
        elif item.get("drawing_count"):
            paragraph_index = item["paragraph"] - 1
            if cover and cover[0] <= paragraph_index <= cover[1]:
                continue
            kind, object_id, drawing_count = "figure", f"p{item['paragraph']}", item["drawing_count"]
        else:
            continue

        before = nearby_paragraphs(position, -1)
        after = nearby_paragraphs(position, 1)
        number = before[-2] if len(before) >= 2 and before[-2]["role"] == "caption_number" else None
        title = before[-1] if before and before[-1]["role"] == "caption_title" else None
        if not number and before and before[-1]["role"] == "caption_number":
            number = before[-1]
        expected_label = "Table" if kind == "table" else "Figure"
        label_matches = bool(number and re.match(rf"^{expected_label}\s+[A-Z]?\d+\s*$", number["text"], re.I))
        if label_matches:
            confirmed = True
        note = after[0] if after and after[0]["role"] == "note" else None

        if number:
            used_caption_paragraphs.add(number["paragraph"])
        if title:
            used_caption_paragraphs.add(title["paragraph"])
        if note:
            used_caption_paragraphs.add(note["paragraph"])

        record = {
            "kind": kind,
            "object": object_id,
            "drawing_count": drawing_count,
            "confirmed_research_object": confirmed,
            "number_paragraph": number["paragraph"] if number else None,
            "title_paragraph": title["paragraph"] if title else None,
            "note_paragraph": note["paragraph"] if note else None,
        }
        objects.append(record)

        if not confirmed:
            if kind == "table":
                issue(
                    "table_classification_required",
                    kind,
                    f"{object_id} 尚未确认是数据表还是布局／问卷表格；确认前不会把缺少图题当作 APA 错误。",
                    [object_id],
                )
                continue
            issue(
                "drawing_classification_required",
                kind,
                f"{object_id} 含 {drawing_count} 个图形对象，但附近没有明确的 Figure 编号；请判断它是论文图还是校徽／装饰对象。",
                [object_id],
            )
            continue
        if not number or not label_matches:
            issue(
                "missing_or_mismatched_number",
                kind,
                f"{object_id} 附近没有匹配的 {expected_label} 编号。",
                [object_id],
            )
        if not title:
            issue(
                "missing_title",
                kind,
                f"{object_id} 附近没有识别到独立的图表标题段落。",
                [object_id],
            )
        if drawing_count > 1 and label_matches:
            issue(
                "multiple_drawings_one_caption",
                kind,
                f"{object_id} 同一段含 {drawing_count} 个图形对象并共用一个说明；请确认是否为一个组合图。",
                [object_id],
            )

    for item in items:
        if item.get("kind") != "paragraph" or item.get("role") not in {"caption_number", "caption_title", "note"}:
            continue
        if item["paragraph"] in used_caption_paragraphs:
            continue
        issue(
            "orphan_caption_part",
            "caption",
            f"第 {item['paragraph']} 段被识别为图表说明，但没有与相邻的表格或图形对象配对。",
            [f"p{item['paragraph']}"],
        )

    detected = len(objects) + sum(item.get("role") in {"caption_number", "caption_title", "note"} for item in items)
    return {
        "status": "not_applicable" if not detected else "needs_review" if issues else "passed",
        "objects_checked": len(objects),
        "paired_objects": sum(bool(item["number_paragraph"] and item["title_paragraph"]) for item in objects),
        "objects": objects,
        "issues": issues,
        "source_urls": {"tables": SOURCES["tables"]["url"], "figures": SOURCES["figures"]["url"]},
        "limitation": (
            "Adjacency and role check only. It does not prove that a drawing is a research figure, that a caption "
            "describes the correct object, or that optional notes are scientifically complete."
        ),
    }


def combined_preflight_summary(**reports):
    """Create one compact overview for AI review and user-facing workflow decisions."""
    sections = []
    for key, report in reports.items():
        if not isinstance(report, dict):
            continue
        issues = report.get("issues", [])
        sections.append({
            "id": key,
            "status": report.get("status", "needs_review"),
            "issue_count": len(issues),
        })
    relevant = [section for section in sections if section["status"] != "not_applicable"]
    needs_review = [section for section in relevant if section["status"] == "needs_review"]
    return {
        "status": "needs_review" if needs_review else "passed",
        "checks_run": len(sections),
        "relevant_checks": len(relevant),
        "issue_count": sum(section["issue_count"] for section in sections),
        "sections": sections,
    }


def _clean_url_candidate(raw):
    """Remove surrounding sentence punctuation without changing displayed text."""
    clean = raw
    while clean.endswith((".", ",", ";", ":", "!", "?")):
        clean = clean[:-1]
    for opening, closing in (("(", ")"), ("[", "]")):
        while clean.endswith(closing) and clean.count(opening) < clean.count(closing):
            clean = clean[:-1]
    return clean


def _url_candidates(text):
    for match in _WEB_URL.finditer(text):
        target = _clean_url_candidate(match.group(0))
        if not target:
            continue
        parsed = urlsplit(target)
        if parsed.scheme.lower() in {"http", "https"} and parsed.netloc:
            yield match.start(), match.start() + len(target), target


def _run_fragment(source, text):
    """Copy one plain run's properties and replace only its text content."""
    node = copy.deepcopy(source)
    for element in list(node):
        if element.tag != qn("w:rPr"):
            node.remove(element)
    content = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        content.set(qn("xml:space"), "preserve")
    content.text = text
    node.append(content)
    return node


def _external_hyperlinks(paragraph, location):
    records = []
    for node in paragraph._p.xpath(".//w:hyperlink"):
        relationship_id = node.get(qn("r:id"))
        if not relationship_id or relationship_id not in paragraph.part.rels:
            continue
        relationship = paragraph.part.rels[relationship_id]
        if relationship.reltype != RT.HYPERLINK or not relationship.is_external:
            continue
        records.append({
            "location": location,
            "text": "".join(node.xpath(".//w:t/text()")),
            "target": relationship.target_ref,
        })
    return records


def document_hyperlinks(doc):
    """Return external body hyperlinks using targets rather than unstable rIds."""
    records = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        records.extend(_external_hyperlinks(paragraph, f"p{index}"))
    return records


def _link_urls_in_plain_run(paragraph, run_node, location):
    children = list(run_node)
    if not children or any(child.tag not in {qn("w:rPr"), qn("w:t")} for child in children):
        return []
    text = "".join(child.text or "" for child in children if child.tag == qn("w:t"))
    candidates = list(_url_candidates(text))
    if not candidates:
        return []
    parent = run_node.getparent()
    if parent is not paragraph._p:
        return []
    position = parent.index(run_node)
    cursor, added = 0, []
    for start, end, target in candidates:
        if cursor < start:
            parent.insert(position, _run_fragment(run_node, text[cursor:start]))
            position += 1
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
        hyperlink.set(qn("w:history"), "1")
        hyperlink.append(_run_fragment(run_node, text[start:end]))
        parent.insert(position, hyperlink)
        position += 1
        added.append({"location": location, "text": text[start:end], "target": target})
        cursor = end
    if cursor < len(text):
        parent.insert(position, _run_fragment(run_node, text[cursor:]))
    parent.remove(run_node)
    return added


def reference_link_check(doc, roles, apply_safe_changes=False):
    """Create live links only from explicit URLs in confirmed reference paragraphs."""
    references = [(index, paragraph) for index, paragraph in enumerate(doc.paragraphs)
                  if roles.get(index) == "reference"]
    before = [record for index, paragraph in references
              for record in _external_hyperlinks(paragraph, f"p{index + 1}")]
    added, protected_locations = [], []
    for index, paragraph in references:
        location = f"p{index + 1}"
        if paragraph._p.xpath(".//w:instrText | .//w:fldChar | .//w:fldSimple | .//w:sdt"):
            if list(_url_candidates(visible_text(paragraph))):
                protected_locations.append(location)
            continue
        if apply_safe_changes:
            for run_node in list(paragraph._p.xpath("./w:r")):
                added.extend(_link_urls_in_plain_run(paragraph, run_node, location))

    after = [record for index, paragraph in references
             for record in _external_hyperlinks(paragraph, f"p{index + 1}")]
    unlinked, bare_dois = [], []
    for index, paragraph in references:
        location, text = f"p{index + 1}", visible_text(paragraph)
        linked = Counter()
        for record in _external_hyperlinks(paragraph, location):
            linked[record["target"]] += 1
            for _, _, displayed_target in _url_candidates(record["text"]):
                if displayed_target != record["target"]:
                    linked[displayed_target] += 1
        candidates = Counter(target for _, _, target in _url_candidates(text))
        for target, count in (candidates - linked).items():
            unlinked.extend({"location": location, "target": target} for _ in range(count))
        url_ranges = [(start, end) for start, end, _ in _url_candidates(text)]
        for match in _BARE_DOI.finditer(text):
            if any(match.start() < end and match.end() > start for start, end in url_ranges):
                continue
            bare_dois.append({"location": location, "text": _clean_url_candidate(match.group(0))})

    detected = len(before) + len(added) + len(unlinked) + len(bare_dois)
    issues = []
    if unlinked:
        issues.append(f"{len(unlinked)} 个完整 DOI／URL 仍不是可点击链接。")
    if bare_dois:
        issues.append(f"{len(bare_dois)} 个 DOI 使用了非 https://doi.org/ 格式，需要核对后转换。")
    if protected_locations:
        issues.append("文献管理器／Word 域中的链接已保留，需在原管理器中更新。")
    status = "not_applicable" if not detected else "needs_review" if issues else "passed"
    return {
        "status": status,
        "reference_paragraphs_checked": len(references),
        "live_links_before": len(before),
        "links_added": added,
        "live_links_after": len(after),
        "linked_locations": sorted({item["location"] for item in added}),
        "unlinked_urls": unlinked,
        "bare_dois": bare_dois,
        "protected_locations": sorted(set(protected_locations)),
        "issues": issues,
        "source_url": SOURCES["reference_links"]["url"],
        "limitation": "Only explicit URLs in confirmed reference paragraphs are linked; missing or incorrect DOI/URL data are never invented.",
    }


def content_signature(doc):
    """Verify content, fields, drawings, math and anchors independently of styles."""
    root = doc._element.body
    tags = ["w:t", "w:instrText", "w:fldChar", "w:fldSimple", "w:tab", "w:br",
            "w:bookmarkStart", "w:bookmarkEnd", "w:hyperlink", "w:commentRangeStart",
            "w:commentRangeEnd", "w:commentReference", "w:footnoteReference", "w:endnoteReference"]
    result = []
    for tag in tags:
        result.append((tag, [(e.text, tuple(sorted(e.attrib.items()))) for e in root.iter(qn(tag))]))
    # Drawing extents may change, but images/relations and native math must persist.
    for uri, local in [("http://schemas.openxmlformats.org/drawingml/2006/main", "blip"),
                       ("http://schemas.openxmlformats.org/officeDocument/2006/math", "oMath")]:
        result.append((local, [etree.tostring(e, method="c14n") for e in root.iter("{" + uri + "}" + local)]))
    return result


def _ordered_subsequence(original, updated):
    """True when every original record survives in order among allowed additions."""
    position = 0
    for record in original:
        while position < len(updated) and updated[position] != record:
            position += 1
        if position == len(updated):
            return False
        position += 1
    return True


def package_payloads(path):
    xml_parts = {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}
    parser = etree.XMLParser(resolve_entities=False, no_network=True, remove_blank_text=True)
    with ZipFile(path) as z:
        result = {}
        for name in z.namelist():
            if not (name.startswith(("word/media/", "word/embeddings/", "word/charts/", "customXml/"))
                    or name in xml_parts):
                continue
            payload = z.read(name)
            if name in xml_parts:
                payload = etree.tostring(etree.fromstring(payload, parser), method="c14n")
            result[name] = hashlib.sha256(payload).hexdigest()
        return result


class Formatter:
    def __init__(self, doc, profile="student", font="Times New Roman", running_head="", config=None,
                 add_styles=False):
        self.doc, self.profile, self.font = doc, profile, font
        self.config = config or {}
        self.running_head = running_head.upper().strip()
        self.events = []
        self.paragraphs = doc.paragraphs
        self.roles = {}
        self.counts = Counter()
        self.cover = None
        self.applied_styles = {}
        self.add_styles = add_styles
        self.statistics_report = None
        self.reference_link_report = None
        self.reference_quality_report = None
        self.numbered_object_report = None
        self.front_matter_report = None
        self.caption_object_report = None
        self.expected_body_text = None

    def event(self, status, message, rule=None, location="document"):
        self.events.append({"status": status, "location": location, "message": message,
                            "rule": rule, "source_url": SOURCES[rule]["url"] if rule else None})

    def reusable_styles(self):
        """Create formatter-owned styles without changing user styles used by preserved content."""
        styles = self.doc.styles
        created = {}
        for name, spec in APA7_PARAGRAPH_STYLES.items():
            try:
                style = styles[name]
            except KeyError:
                style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if style.type != WD_STYLE_TYPE.PARAGRAPH:
                raise ValueError(f"现有样式“{name}”类型冲突，无法建立可继续编辑的 APA7 样式。")
            _configure_paragraph_style(style, self.font, spec)
            style.quick_style = True
            style.hidden = False
            style.unhide_when_used = True
            created[name] = style
        for name, spec in APA7_CHARACTER_STYLES.items():
            try:
                style = styles[name]
            except KeyError:
                style = styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
            if style.type != WD_STYLE_TYPE.CHARACTER:
                raise ValueError(f"现有样式“{name}”类型冲突，无法建立可继续编辑的 APA7 样式。")
            _configure_style_font(style, self.font, bold=spec["bold"], italic=spec["italic"])
            style.quick_style = True
            style.hidden = False
            style.unhide_when_used = True
            created[name] = style
        body = created["APA7 Body"]
        for name in ("APA7 Heading 1", "APA7 Heading 2", "APA7 Heading 3", "APA7 Abstract", "APA7 Keywords",
                     "APA7 Caption Title", "APA7 Note", "APA7 Block Quote", "APA7 Run-in Heading"):
            created[name].next_paragraph_style = body
        created["APA7 Title"].next_paragraph_style = created["APA7 Title Metadata"]
        created["APA7 Title Metadata"].next_paragraph_style = created["APA7 Title Metadata"]
        created["APA7 Body"].next_paragraph_style = body
        created["APA7 Reference"].next_paragraph_style = created["APA7 Reference"]
        created["APA7 Caption Number"].next_paragraph_style = created["APA7 Caption Title"]
        self.event("applied", "已加入可继续编辑的 APA7 正文、标题、参考文献、图表说明和块引用样式。", "paragraph")

    def apply_role_style(self, index, paragraph, role):
        if not self.add_styles:
            return
        name = ROLE_STYLE_NAMES.get(role)
        if name:
            paragraph.style = self.doc.styles[name]
            self.applied_styles[index] = name

    def detect_roles(self):
        state = "body"
        texts = [visible_text(p).strip() for p in self.paragraphs]
        title_indices = []
        for i, p in enumerate(self.paragraphs):
            text = texts[i]
            style = p.style.name if p.style is not None else ""
            next_element = p._p.getnext()
            followed_by_visual = next_element is not None and (
                next_element.tag == qn("w:tbl") or has_descendant(next_element, "w:drawing")
            )
            previous = self.paragraphs[i - 1] if i else None
            if state == "abstract" and (p.paragraph_format.page_break_before or
                    (previous is not None and previous._p.xpath(".//w:br[@w:type='page'] | ./w:pPr/w:sectPr"))):
                state = "body"
                self.event("review", "摘要在显式分页处结束；后续普通样式正文的标题需确认。", "paragraph", f"p{i+1}")
            role = state
            if not text:
                role = "preserve"
            elif text.casefold() in SECTION_LABELS:
                role = "section"
                state = "abstract" if text.casefold() == "abstract" else "reference" if text.casefold() in {"references", "reference"} else "body"
            elif re.fullmatch(r"Appendix(?: [A-Z])?", text):
                role, state = "appendix", "body"
            elif style in STYLE_ROLE_HINTS:
                role, state = STYLE_ROLE_HINTS[style], "body"
                if role == "title":
                    title_indices.append(i)
            elif style == "APA7 Run-in Heading":
                role, state = ("heading5" if any(run.style and run.style.name == "APA7 Level 5 Prefix" for run in p.runs)
                               else "heading4"), "body"
            elif re.fullmatch(r"Heading [1-5]", style):
                role, state = "heading" + style[-1], "body"
            elif style == "Title":
                role, state = "title", "body"
                title_indices.append(i)
            elif CAPTION.fullmatch(text):
                role = "caption_number"
                state = "body"
            elif COMBINED_CAPTION.match(text) and state != "reference" and (style == "Caption" or followed_by_visual):
                role = "preserve"
                self.event("review", "疑似编号和标题合在一段的图表说明，原格式保留；需拆分编号与下一行标题。", "figures" if text.lower().startswith("figure") else "tables", f"p{i+1}")
            elif style in {"Quote", "Intense Quote", "APA Block Quote"}:
                role = "quote"
            elif style.casefold().startswith(("toc ", "index ", "list ", "code", "source code")) or p._p.xpath("./w:pPr/w:numPr"):
                role = "preserve"
                self.event("review", "列表／目录结构保留；请确认其格式及是否需要保留。", "headings", f"p{i+1}")
            elif p._p.xpath(".//m:oMathPara"):
                role = "equation"
            elif p._p.xpath(".//m:oMath"):
                role = "body" if p._p.xpath(".//w:t[normalize-space()]") else "equation"
            elif p._p.xpath(".//w:object | .//w:pict | .//wp:anchor"):
                role = "preserve"
            elif text.startswith("Keywords:"):
                role = "keywords"
                state = "body"
            self.roles[i] = role

        # Recognize a cover only if Title is followed by an explicit page boundary.
        # Pagination cannot be derived from character counts.
        if title_indices and title_indices[0] <= 4:
            start = title_indices[0]
            boundary = None
            for j in range(start + 1, min(len(self.paragraphs), start + 20)):
                p = self.paragraphs[j]
                if p.paragraph_format.page_break_before or (self.roles[j] == "title" and texts[j] == texts[start]) or texts[j].casefold() == "abstract":
                    boundary = j
                    break
                if p._p.xpath(".//w:br[@w:type='page'] | ./w:pPr/w:sectPr"):
                    boundary = j + 1
                    break
            if boundary is not None:
                self.cover = (start, boundary - 1)
                for j in range(start + 1, boundary):
                    if texts[j] and self.roles[j] == "body":
                        self.roles[j] = "title_meta"
                self.event("assumption", "由 Title 样式和显式分页推定标题页范围；作者、单位及课程字段仍需核对。", "title")
            else:
                self.event("review", "找到 Title 样式，但无法确认标题页边界；可在配置中指定 title_page 段落范围。", "title")
        else:
            self.event("review", "未确认标题页。请标记 Word 的 Title 样式或配置 roles/title_page；程序不会编造作者信息。", "title")

        for i in range(len(self.paragraphs)):
            role = self.roles[i]
            j = i + 1
            if j < len(self.paragraphs) and role in {"caption_number", "appendix"} and texts[j] and self.roles[j] in {"body", "reference"}:
                next_el = self.paragraphs[j]._p.getnext()
                has_visual = next_el is not None and (
                    next_el.tag == qn("w:tbl") or has_descendant(next_el, "w:drawing")
                )
                if role == "appendix" or has_visual:
                    self.roles[j] = "appendix_title" if role == "appendix" else "caption_title"
            if texts[i].startswith("Note.") and i > 0:
                prev = self.paragraphs[i]._p.getprevious()
                if prev is not None and (prev.tag == qn("w:tbl") or has_descendant(prev, "w:drawing")):
                    self.roles[i] = "note"

        for key, value in self.config.get("roles", {}).items():
            idx = int(key) - 1
            if idx < 0 or idx >= len(self.paragraphs) or value not in ROLES:
                raise ValueError(f"无效 roles 配置：{key}={value}")
            self.roles[idx] = value
        cover = self.config.get("title_page")
        if cover:
            start, end = int(cover[0]) - 1, int(cover[1]) - 1
            if not 0 <= start <= end < len(self.paragraphs):
                raise ValueError("title_page 必须为有效的首尾段落编号 [start, end]。")
            self.cover = (start, end)
            for i in range(start, end + 1):
                if texts[i]:
                    self.roles[i] = "title" if i == start else "title_meta"

    def preflight(self):
        xml = self.doc._element
        revision_names = {"ins", "del", "moveFrom", "moveTo", "moveFromRangeStart", "moveToRangeStart",
                          "cellIns", "cellDel", "cellMerge", "numberingChange", "customXmlInsRangeStart", "customXmlDelRangeStart"}
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        for part in self.doc.part.package.parts:
            if not str(part.partname).startswith("/word/") or not str(part.partname).endswith(".xml"):
                continue
            tree = etree.fromstring(part.blob, parser)
            for node in tree.iter():
                if not isinstance(node.tag, str) or not node.tag.startswith("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"):
                    continue
                local = etree.QName(node).localname
                if local in revision_names or local.endswith("PrChange"):
                    raise ValueError("检测到尚未处理的修订（" + str(part.partname) + "）。请先在 Word 的副本中审阅／接受或拒绝修订，再运行格式化；不会自动替你接受修订。")
        for expr, message in [
            (".//w:sdt", "存在内容控件，其中的段落可能未被自动格式化。"),
            (".//wp:anchor", "存在浮动图片；保留锚点，需检查环绕、位置与跨页。"),
            (".//w:txbxContent", "存在文本框；其中的标题或图注需单独检查。"),
            (".//m:oMath | .//m:oMathPara", "原生公式保留；公式编号、统计符号及上下间距需检查。"),
            (".//w:instrText | .//w:fldSimple", "保留 Word／文献管理器域；更新域可能重新应用文献管理器样式，需在管理器中也选择 APA 7。"),
        ]:
            if xml.xpath(expr):
                self.event("review", message)
        self.event("review", "学校或期刊可能有额外规定。请核对目标要求。", "scope")
        self.event("review", "最终仍需在 Word 中逐页检查分页、图表位置和字体替代；本次程序运行未做视觉验收。")

    def sections(self):
        if not self.doc.sections:
            # Some third-party generators omit the final body-level ``sectPr``
            # entirely. Word tolerates that package, but python-docx then
            # exposes zero sections and later page-dependent work cannot
            # resolve margins, headers, or image width. Add the missing
            # formatting container without changing any manuscript content.
            self.doc._body._body.get_or_add_sectPr()
            self.event(
                "changed",
                "原稿没有 Word 节属性；已加入默认节以设置页边距、页眉和图片可用宽度。",
                "margins",
                "section1",
            )
        for n, section in enumerate(self.doc.sections, 1):
            section.top_margin = section.bottom_margin = Inches(1)
            section.left_margin = section.right_margin = Inches(1)
            section.gutter = 0
            # Keep portrait/landscape and paper size; institutional preferences vary.
            remove(section._sectPr, "w:docGrid", "w:pgBorders")
            cols = section._sectPr.find(qn("w:cols"))
            if cols is not None and int(cols.get(qn("w:num"), "1")) > 1:
                self.event("review", "多栏版式保留；APA 投稿通常需要单栏，需按目标期刊确认。", "scope", f"section{n}")
            if section.page_width is None:
                self.event(
                    "review",
                    "原稿未显式写入纸张尺寸；未强制改变纸张，页眉与图片宽度暂按 Letter 计算，需逐页核对。",
                    "scope",
                    f"section{n}",
                )
            elif round(section.page_width.inches, 2) not in {8.5, 11.0}:
                self.event("review", "保留现有纸张尺寸；默认新建测试文档使用 Letter，学校指定 A4 时应遵循学校要求。", "scope", f"section{n}")
        self.event("applied", "各节四边页边距设为 1 英寸，保留纸张方向与尺寸。", "margins")

    def header_footer(self):
        if self.profile == "professional" and not self.running_head:
            self.event("review", "投稿模式尚未提供 running_head：先生成页码，短标题待补。", "header")
        if len(self.running_head) > 50:
            raise ValueError("Running head 含空格及标点不得超过 50 个字符，请提供较短标题。")
        # Formatting copy: explicit replacement is optional; unknown existing content
        # remains untouched by default. Never silently erase a logo or running text.
        replace = bool(self.config.get("replace_headers", False))
        # python-docx creates explicit default/first/even header parts when the
        # three variants below are accessed.  LibreOffice can otherwise render
        # a stale even-page layout even though all three OOXML parts contain the
        # same APA header.  Make the intended odd/even variants active and keep
        # their content identical so page numbers stay at the upper right in
        # both Word and LibreOffice renders.
        self.doc.settings.odd_and_even_pages_header_footer = True
        seen = set()
        previous_width = None
        for n, section in enumerate(self.doc.sections, 1):
            current_width = section_text_width(section)
            for container in (section.header, section.first_page_header, section.even_page_header):
                root = container._element
                text = "".join(root.xpath(".//w:t/text()"))
                instructions = " ".join(root.xpath(".//w:instrText/text()"))
                managed = bool(root.xpath(".//w:bookmarkStart[@w:name='APA7Header']"))
                complex_content = bool(root.xpath(".//w:drawing | .//w:tbl | .//w:pict"))
                only_page = not complex_content and re.fullmatch(r"\s*\d*\s*", text) and (not instructions.strip() or instructions.strip() == "PAGE")
                if previous_width is not None and current_width != previous_width and container.is_linked_to_previous and (managed or only_page or replace):
                    container.is_linked_to_previous = False
                    root = container._element
                if root in seen:
                    continue
                seen.add(root)
                if not (managed or only_page or replace):
                    self.event("review", "现有页眉有内容，已保留；要重建 APA 页眉，可设置 replace_headers=true。", "header", f"section{n}")
                    continue
                if text.strip() and replace and not managed and not only_page:
                    self.event("changed", "按配置重建副本页眉；原页眉文字：" + text, "header", f"section{n}")
                for elem in list(root):
                    root.remove(elem)
                p = container.add_paragraph()
                paragraph_format(p, indent=0)
                bookmark = OxmlElement("w:bookmarkStart")
                bookmark.set(qn("w:id"), str(800000 + n * 10 + len(seen)))
                bookmark.set(qn("w:name"), "APA7Header")
                p._p.append(bookmark)
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), bookmark.get(qn("w:id")))
                p._p.append(end)
                if self.profile == "professional" and self.running_head:
                    p.add_run(self.running_head)
                    p.paragraph_format.tab_stops.add_tab_stop(section_text_width(section), WD_TAB_ALIGNMENT.RIGHT)
                    p.add_run("\t")
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = p.add_run()
                for tag, val in [("w:fldChar", "begin"), ("w:instrText", " PAGE "), ("w:fldChar", "end")]:
                    node = OxmlElement(tag)
                    if tag == "w:instrText":
                        node.set(qn("xml:space"), "preserve")
                        node.text = val
                    else:
                        node.set(qn("w:fldCharType"), val)
                    run._r.append(node)
                typography(p, self.font, bold=False, italic=False)
            section.header_distance = Inches(0.5)
            pg = child(section._sectPr, "w:pgNumType", {"w:fmt": "decimal"})
            if n == 1:
                pg.set(qn("w:start"), "1")
            else:
                pg.attrib.pop(qn("w:start"), None)
            for footer in (section.footer, section.first_page_footer, section.even_page_footer):
                if footer._element.xpath(".//w:t | .//w:drawing | .//w:fldChar"):
                    self.event("review", "现有页脚内容保留，请检查是否有多余页码。", "header", f"section{n}")
            previous_width = current_width
        self.event("applied", "空白或本工具生成的页眉设为自动 PAGE 域；现有复杂页眉另列检查项。", "header")

    def run_in(self, i, p, role):
        prefix = self.config.get("run_in_headings", {}).get(str(i + 1))
        if not prefix or not visible_text(p).startswith(prefix) or not prefix.endswith("."):
            self.event("review", "四／五级标题需明确标题文本（含结尾句号）并与正文同段。可用 run_in_headings 指定；此段保持原样。", "headings", f"p{i+1}")
            return False
        if not split_plain_run_at(p, len(prefix)):
            self.event("review", "行内标题含复杂域或嵌入对象，已保留。", "headings", f"p{i+1}")
            return False
        if self.add_styles:
            p.style = self.doc.styles["APA7 Run-in Heading"]
            self.applied_styles[i] = "APA7 Run-in Heading"
            prefix_style = self.doc.styles["APA7 Level 5 Prefix" if role == "heading5" else "APA7 Level 4 Prefix"]
        else:
            p.style = self.doc.styles["Normal"]
            prefix_style = None
        offset = 0
        for run in p.runs:
            if offset < len(prefix):
                if prefix_style is not None:
                    run.style = prefix_style
                run.bold = True
                run.italic = role == "heading5"
            else:
                # Normal paragraph style removes heading inheritance; preserve
                # explicit semantic emphasis in the continuation text.
                if run.style and re.fullmatch(r"Heading [45] Char", run.style.name):
                    run.style = None
            offset += len(run.text)
        paragraph_format(p)
        typography(p, self.font)
        return True

    def paragraphs_format(self):
        for i, p in enumerate(self.paragraphs):
            role = self.roles[i]
            text = visible_text(p).strip()
            original_alignment = p.paragraph_format.alignment
            if role == "preserve":
                if text:
                    self.event("review", "此段保留原格式（特殊对象、列表或明确 preserve）。", location=f"p{i+1}")
                continue
            if role in {"heading4", "heading5"}:
                if self.run_in(i, p, role):
                    self.counts[role] += 1
                continue
            if p._p.xpath(".//wp:inline") and not text:
                continue
            self.apply_role_style(i, p, role)
            # Do not discard scientific italic/superscript/bold in ordinary prose.
            typography(p, self.font)
            paragraph_format(p)
            if role in {"title", "section", "appendix", "appendix_title", "heading1"}:
                paragraph_format(p, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
                typography(p, self.font, bold=True, italic=False)
                if role in {"section", "appendix"} and text.casefold() != "author note":
                    # Avoid an extra blank page after an existing explicit page break.
                    previous = p._p.getprevious()
                    already_breaks = has_page_boundary(previous)
                    if not already_breaks:
                        p.paragraph_format.page_break_before = True
            elif role in {"heading2", "heading3"}:
                paragraph_format(p, indent=0, keep=True)
                typography(p, self.font, bold=True, italic=role == "heading3")
            elif role == "title_meta":
                paragraph_format(p, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
                typography(p, self.font, bold=False, italic=False)
            elif role == "abstract":
                paragraph_format(p, indent=0)
            elif role == "reference":
                paragraph_format(p, indent=-0.5, left=0.5)
            elif role in {"caption_number", "caption_title", "note"}:
                paragraph_format(p, indent=0, keep=role != "note")
                if role != "note":
                    typography(p, self.font, bold=role == "caption_number", italic=role == "caption_title")
                elif visible_text(p).startswith("Note.") and split_plain_run_at(p, 5):
                    offset = 0
                    for run in p.runs:
                        if offset < 5:
                            run.italic = True
                        offset += len(run.text)
            elif role in {"quote", "quote_continuation"}:
                paragraph_format(p, indent=0.5 if role == "quote_continuation" else 0, left=0.5)
                self.event("review", "按已标记块引用排版；请核对引用词数、引号、出处定位及是否为同一引文的后续段。", "quotations", f"p{i+1}")
            elif role == "equation":
                paragraph_format(p, indent=0)
                p.paragraph_format.alignment = original_alignment
                self.event("review", "已识别独立公式并规范段落缩进、双倍行距和段落间距；公式编号、右侧位置、变量及标点仍需核对。", "equations", f"p{i+1}")
            if role in {"title", "section", "appendix", "appendix_title", "caption_number", "caption_title"} or role.startswith("heading"):
                # A nil direct border overrides decorative borders inherited from
                # built-in/custom Title styles without rewriting document styles.
                border = child(p._p.get_or_add_pPr(), "w:pBdr")
                for side in ("top", "left", "bottom", "right", "between", "bar"):
                    child(border, "w:" + side, {"w:val": "nil"})
            if role == "caption_number" and i > 0 and self.roles[i - 1] == "body" and not p.paragraph_format.page_break_before:
                p.paragraph_format.space_before = Pt(FONTS[self.font] * 2)
            if role == "note" and i + 1 < len(self.paragraphs) and self.roles[i + 1] == "body" and not self.paragraphs[i + 1].paragraph_format.page_break_before:
                p.paragraph_format.space_after = Pt(FONTS[self.font] * 2)
            if role not in {"title", "title_meta", "abstract", "keywords"}:
                previous_index = i - 1
                while previous_index >= 0 and not visible_text(self.paragraphs[previous_index]).strip():
                    previous_index -= 1
                if previous_index >= 0 and self.roles[previous_index] in {"abstract", "keywords"}:
                    boundary_exists = bool(p.paragraph_format.page_break_before)
                    boundary_exists = boundary_exists or any(
                        self.paragraphs[index]._p.xpath(".//w:br[@w:type='page'] | ./w:pPr/w:sectPr")
                        for index in range(previous_index, i)
                    )
                    if not boundary_exists:
                        p.paragraph_format.page_break_before = True
                        self.event(
                            "changed",
                            "摘要／关键词后的论文正文从新页开始。",
                            "paragraph",
                            f"p{i+1}",
                        )
            if role.startswith("heading"):
                if p._p.xpath("./w:pPr/w:numPr") or re.match(r"^\s*(?:\d+(?:\.\d+)*|[A-Z])[.)]?\s+", text):
                    self.event("review", "标题包含自动或文字编号；需确认学校要求后取消编号。", "headings", f"p{i+1}")
                if text.casefold() == "introduction":
                    self.event("review", "正文开头通常由论文标题起标题作用，无须 Introduction 标题。", "headings", f"p{i+1}")
            if role in {"title", "caption_title", "appendix_title"} or role.startswith("heading"):
                self.event("review", "大小写未重写；请核对 title case、缩写、专名及科学符号。", "case", f"p{i+1}")
            if COMBINED_CAPTION.match(text) and role != "reference":
                self.event("review", "疑似编号和标题在同段的图表说明；请拆成编号、下一行标题，或明确 roles 配置。", "figures" if text.lower().startswith("figure") else "tables", f"p{i+1}")
            if role == "body" and len(text.split()) < 14 and text and all(r.bold for r in p.runs if r.text.strip()):
                self.event("review", "疑似用粗体模拟的标题；未猜测其层级，请应用 Heading 样式。", "headings", f"p{i+1}")
            if text.startswith(("\t", "  ")) or visible_text(p).startswith(("\t", "  ")):
                self.event("review", "存在手动空格／制表符缩进，可能与段落缩进叠加。", "paragraph", f"p{i+1}")
            self.counts[role] += 1
        self.event("applied", "已识别段落应用正文／摘要／参考文献／标题等对应行距、缩进与字体。保留普通正文中的强调及上下标。", "paragraph")
        if self.cover and self.roles.get(self.cover[0]) == "title":
            start, end = self.cover
            title = self.paragraphs[start]
            if start == 0:
                title.paragraph_format.space_before = Pt(FONTS[self.font] * 6)
            elif all(not visible_text(p).strip() for p in self.paragraphs[:start]):
                for p in self.paragraphs[:start]:
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = Pt(FONTS[self.font] * 2)
            if start + 1 <= end and visible_text(self.paragraphs[start + 1]).strip():
                title.paragraph_format.space_after = Pt(FONTS[self.font] * 2)
            elif start + 1 <= end:
                blank = self.paragraphs[start + 1]
                blank.paragraph_format.space_before = blank.paragraph_format.space_after = Pt(0)
                blank.paragraph_format.line_spacing = Pt(FONTS[self.font] * 2)
            self.event("applied", "已确认标题页的标题起始留白及标题后缺失的作者间隔已设置；作者注位置仍需逐页核对。", "title")

    def tables(self):
        table_roles = self.config.get("table_roles", {})
        if not isinstance(table_roles, dict) or any(not str(k).isdigit() or not 1 <= int(k) <= len(self.doc.tables)
                                                  or v not in {"data", "preserve"} for k, v in table_roles.items()):
            raise ValueError("table_roles 必须把有效表格编号映射为 data 或 preserve。")
        for i, table in enumerate(self.doc.tables, 1):
            location = f"table{i}"
            if table_roles.get(str(i)) == "preserve":
                self.event("review", "按结构配置保留此表格；可能为问卷、布局或尚未确定用途，不套用数据表样式。", "tables", location)
                continue
            if table._tbl.xpath(".//w:gridSpan | .//w:vMerge | .//w:hMerge | ./w:tr/w:trPr/w:gridBefore | ./w:tr/w:trPr/w:gridAfter | .//w:tbl"):
                self.event("review", "合并单元格／嵌套或不规则表格保留原结构与格式；需要确定多级表头及必要横线后处理。", "tables", location)
                continue
            header_rows = int(self.config.get("table_header_rows", {}).get(str(i), 1))
            if not 1 <= header_rows <= len(table.rows):
                raise ValueError(f"table_header_rows[{i}] 超出表格行数。")
            if str(i) not in self.config.get("table_header_rows", {}):
                explicit = [r for r in table.rows if r._tr.xpath("./w:trPr/w:tblHeader")]
                if explicit:
                    header_rows = len(explicit)
                else:
                    self.event("assumption", "本表默认第 1 行为表头；可用 table_header_rows 配置确认。", "tables", location)
            tblpr = table._tbl.tblPr
            remove(tblpr, "w:tblStyle", "w:shd", "w:tblBorders")
            borders = child(tblpr, "w:tblBorders")
            for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
                child(borders, "w:" + edge, {"w:val": "single" if edge in {"top", "bottom"} else "nil", "w:sz": "8", "w:color": "000000"})
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            seen = set()
            for row_num, row in enumerate(table.rows):
                trpr = row._tr.get_or_add_trPr()
                remove(trpr, "w:trHeight")
                if row_num < header_rows:
                    child(trpr, "w:tblHeader", {"w:val": "1"})
                else:
                    remove(trpr, "w:tblHeader")
                for col_num, cell in enumerate(row.cells):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    pr = cell._tc.get_or_add_tcPr()
                    remove(pr, "w:shd", "w:tcBorders")
                    cb = child(pr, "w:tcBorders")
                    for edge in ("top", "bottom", "left", "right", "start", "end", "insideH", "insideV"):
                        val = "single" if edge == "bottom" and row_num == header_rows - 1 else "nil"
                        # Avoid nil cell borders masking outer table rules.
                        if edge == "top" and row_num == 0 or edge == "bottom" and row_num == len(table.rows) - 1:
                            val = "single"
                        child(cb, "w:" + edge, {"w:val": val, "w:sz": "8", "w:color": "000000"})
                    for p in cell.paragraphs:
                        align = WD_ALIGN_PARAGRAPH.CENTER
                        if row_num >= header_rows and (col_num == 0 or len(visible_text(p)) > 45):
                            align = WD_ALIGN_PARAGRAPH.LEFT
                        paragraph_format(p, indent=0, align=align, spacing=1.0)
                        typography(p, self.font)
            prev = table._tbl.getprevious()
            if prev is None or prev.tag != qn("w:p"):
                self.event("review", "表格上方未找到标题段落。", "tables", location)
            self.event("applied", "简单 Word 表应用必要横线、无竖线、表头重复、单倍单元格行距；保留列宽与内容。", "tables", location)
            self.event("review", "请核对表头行数、标题、注释、必要分组横线、列宽和页面位置；表格可能用于问卷或布局而非数据。", "tables", location)

    def images(self):
        # Use the actual section of each inline shape, not the narrowest section.
        width_by_p = {}
        section_idx = 0
        for p in self.paragraphs:
            section = self.doc.sections[section_idx]
            width_by_p[p._p] = section_text_width(section)
            if p._p.xpath("./w:pPr/w:sectPr") and section_idx + 1 < len(self.doc.sections):
                section_idx += 1
        for i, shape in enumerate(self.doc.inline_shapes, 1):
            parents = shape._inline.xpath("ancestor::w:p[1]")
            if not parents or parents[0] not in width_by_p:
                self.event("review", "表内或特殊容器中的图片未自动缩放。", "figures", f"image{i}")
                continue
            width = width_by_p[parents[0]]
            if shape.width > width:
                old = shape.width
                shape.width = width
                shape.height = round(shape.height * width / old)
                self.event("changed", "超出页边距的内嵌图片等比例缩小到正文宽度；请重新核对图中文字的物理字号。", "figures", f"image{i}")
            self.event("review", "图像原始字节保留；请检查图内 8–14 pt 无衬线文字、清晰度、坐标、图例、编号、版权说明及首次提及位置。", "figures", f"image{i}")

    def statistics_and_equations(self):
        """Apply conservative presentation fixes and record all content-sensitive checks."""
        import apa7_statistics

        self.statistics_report = apa7_statistics.format_statistics_and_equations(
            self.doc, self.roles, self.config.get("table_roles", {}), self.profile
        )
        report = self.statistics_report
        equations = report["equations"]
        if report["expressions_found"]:
            self.event(
                "applied",
                f"识别 {report['expressions_found']} 处统计表达；规范了 {len(report['safe_text_edits'])} 处安全文本格式和 "
                f"{report['symbols_formatted']} 个统计符号的斜体／正体。统计数值未改动。",
                "statistics",
            )
        if equations["native_math_paragraphs"] or equations["plain_text_formula_candidates"]:
            self.event(
                "applied",
                f"识别 {equations['native_math_paragraphs']} 个含 Word 原生公式的段落和 "
                f"{equations['plain_text_formula_candidates']} 个纯文本公式候选；原生公式内容未重写。",
                "equations",
            )
        if report["issues"]:
            self.event(
                "review",
                f"统计数据与公式有 {len(report['issues'])} 项需要作者确认；不会自动舍入、重算、补写或改变显著性。",
                "statistics",
            )

    def reference_hyperlinks(self):
        """Make explicit URLs in confirmed references live without changing text."""
        self.reference_link_report = reference_link_check(self.doc, self.roles, apply_safe_changes=True)
        report = self.reference_link_report
        if report["links_added"]:
            self.event(
                "applied",
                f"将参考文献中 {len(report['links_added'])} 个已有 DOI／URL 设为可点击链接；显示文字未改变。",
                "reference_links",
            )
        if report["issues"]:
            self.event(
                "review",
                "参考文献链接仍有需要确认的内容：" + " ".join(report["issues"]),
                "reference_links",
            )

    def reference_quality(self):
        """Audit clear reference-list quality signals without changing entries."""
        self.reference_quality_report = reference_quality_check(self.doc.paragraphs, self.roles)
        report = self.reference_quality_report
        if report["status"] != "not_applicable":
            self.event(
                "applied",
                f"已检查 {report['reference_entries_checked']} 条参考文献的重复、字母顺序和同作者同年后缀；没有移动或改写条目。",
                "reference_order",
            )
        if report["issues"]:
            self.event(
                "review",
                f"参考文献质量检查发现 {len(report['issues'])} 项需要确认。",
                "reference_order",
            )

    def numbered_objects(self):
        """Audit labels and body callouts; never renumber document content."""
        self.numbered_object_report = numbered_object_check(
            self.doc, self.roles, self.config.get("table_roles")
        )
        report = self.numbered_object_report
        if report["status"] != "not_applicable":
            total_labels = sum(item["labels"] for item in report["counts"].values())
            total_callouts = sum(item["callouts"] for item in report["counts"].values())
            self.event(
                "applied",
                f"已检查 {total_labels} 个图表／公式编号和 {total_callouts} 处正文提及；没有自动重新编号。",
                "figures",
            )
        if report["issues"]:
            self.event(
                "review",
                f"图、表、公式的编号与正文提及有 {len(report['issues'])} 项需要核对。",
                "figures",
            )

    def structural_audits(self):
        """Run profile/front-matter and caption/object checks without rewriting content."""
        self.front_matter_report = front_matter_check(
            self.doc.paragraphs, self.roles, self.profile, self.cover
        )
        self.caption_object_report = caption_object_check(
            self.doc, self.roles, self.config.get("table_roles"), self.cover
        )
        if self.front_matter_report["issues"]:
            self.event(
                "review",
                f"标题页、摘要或关键词有 {len(self.front_matter_report['issues'])} 项需要确认。",
                "title",
            )
        if self.caption_object_report["status"] != "not_applicable":
            self.event(
                "applied",
                f"已尝试配对 {self.caption_object_report['objects_checked']} 个表格／图形对象与编号、标题和注释；没有移动对象。",
                "figures",
            )
        if self.caption_object_report["issues"]:
            self.event(
                "review",
                f"图表与说明配对有 {len(self.caption_object_report['issues'])} 项需要确认。",
                "figures",
            )
    def run(self):
        self.preflight()
        self.detect_roles()
        self.sections()
        self.header_footer()
        if self.add_styles:
            self.reusable_styles()
        self.paragraphs_format()
        self.reference_hyperlinks()
        self.reference_quality()
        self.tables()
        self.images()
        self.statistics_and_equations()
        self.numbered_objects()
        self.structural_audits()
        self.expected_body_text = "".join(self.doc._element.body.xpath(".//w:t/text()"))
        self.event("review", "参考文献只处理已识别条目的段落格式、明确 URL 的链接和保守的重复／排序提示。书目事实、文献类型、标题大小写、斜体位置和缺失 DOI 仍需核对。", "references")
        self.event("review", "标题页的必需信息、标题位置、作者间空行及投稿作者注必须检查；不根据缺失信息生成内容。", "title")
        self.event("review", "附录单图表例外及字母编号仍需结合附录结构核对。", "appendices")


def unique_output(source, explicit=None, save_feedback=False):
    if explicit:
        out = Path(explicit).expanduser().resolve()
        if out == source.resolve() or out.exists():
            raise ValueError("输出必须是一个尚不存在的新文件，不能覆盖原稿或已有结果。")
        if out.suffix.lower() != ".docx":
            raise ValueError("输出扩展名必须为 .docx。")
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        out = source.with_name(f"{source.stem}_APA7_{stamp}.docx")
    if not out.parent.is_dir():
        raise ValueError("输出文件夹不存在。")
    if save_feedback and out.with_suffix(".feedback.html").exists():
        raise ValueError("同名简要反馈已存在，请换一个输出文件名。")
    return out


def review_checklist(profile):
    """Human review tasks, deliberately distinct from machine preservation checks."""
    title_fields = (["论文标题", "作者", "单位", "课程编号与名称", "教师", "截止日期"]
                    if profile == "student" else ["论文标题", "作者", "作者与单位的对应关系", "适用的作者注信息", "running head"])
    items = [
        ("target", "目标要求", "scope", ["核对学校、课程或期刊的具体要求与例外"]),
        ("title_metadata", "标题页信息", "title", title_fields),
        ("structure", "结构与标题", "headings", ["确认标题层级和 title case", "四／五级标题在正文同段，且只缩进首行"]),
        ("references", "引用与参考文献", "references", ["核对作者、年份、文献类型、标点和斜体", "核对排序、DOI、URL 可点击状态与正文引用的对应关系"]),
        ("statistics", "统计数据汇报", "statistics", ["核对检验统计量、自由度、精确 p 值、效应量和置信区间", "确认任何舍入、显著性和原始分析输出一致"]),
        ("equations", "公式", "equations", ["核对公式变量、上下标、括号和标点", "核对独立公式编号顺序及编号在右侧的位置"]),
        ("tables", "表格", "tables", ["确认实际表头和必要的横线", "核对跨页、编号、标题、注释与正文首次提及顺序"]),
        ("figures", "图片与图表", "figures", ["核对编号、图题、注释、数据、单位、图例和版权", "重绘 SVG 核对缓存与数据源是否一致", "按最终插入尺寸核对图内字体和清晰度"]),
        ("render", "逐页视觉检查", "scope", ["在 Word 或渲染器中检查所有页面", "确认页码、标题页位置、分页、字体替代和图表溢出"]),
    ]
    return [{"id": key, "label": label, "status": "pending_review", "requirements": requirements,
             "source_url": SOURCES[rule]["url"]} for key, label, rule, requirements in items]


def prepare_config(path, profile="student"):
    """Inspect without formatting; proposed roles are never silently confirmed."""
    source = Path(path).expanduser().resolve()
    if not source.is_file() or source.suffix.lower() != ".docx":
        raise ValueError("结构配置只支持 .docx；旧 .doc 请先转换并核对。")
    if profile not in {"student", "professional"}:
        raise ValueError("无效论文模式。")
    before = digest(source)
    doc = Document(source)
    formatter = Formatter(doc, profile=profile)
    formatter.preflight()
    formatter.detect_roles()
    paragraphs = [{"paragraph": i + 1, "style": p.style.name if p.style else "",
                   "text": visible_text(p), "inferred_role": formatter.roles[i]}
                  for i, p in enumerate(doc.paragraphs)]
    tables = [{"table": i, "rows": len(t.rows), "columns": len(t.columns),
               "complex": bool(t._tbl.xpath(".//w:gridSpan | .//w:vMerge | .//w:hMerge | .//w:tbl | ./w:tr/w:trPr/w:gridBefore | ./w:tr/w:trPr/w:gridAfter")),
               "cells": [[c.text for c in row.cells] for row in t.rows]}
              for i, t in enumerate(doc.tables, 1)]
    object_summary = {
        "top_level_tables": len(tables),
        "inline_drawings": descendant_count(doc._element, "wp:inline"),
        "floating_drawings": descendant_count(doc._element, "wp:anchor"),
        "native_charts": descendant_count(doc._element, "c:chart"),
        "suggested_abstract_items": sum(role in {"abstract", "keywords"} for role in formatter.roles.values()),
        "suggested_heading_items": sum(role in {"heading1", "heading2", "heading3", "heading4", "heading5"}
                                       for role in formatter.roles.values()),
        "suggested_reference_entries": sum(role == "reference" for role in formatter.roles.values()),
        "suggested_appendix_labels": sum(role == "appendix" for role in formatter.roles.values()),
    }
    structural_roles = {"title", "section", "heading1", "heading2", "heading3", "heading4", "heading5",
                        "appendix", "appendix_title", "caption_number", "caption_title"}
    hierarchy_outline = [{"paragraph": i + 1, "role": formatter.roles[i], "text": visible_text(p)}
                         for i, p in enumerate(doc.paragraphs) if formatter.roles[i] in structural_roles]
    citation_review = citation_reference_check(doc.paragraphs, formatter.roles)
    reference_links = reference_link_check(doc, formatter.roles, apply_safe_changes=False)
    reference_quality = reference_quality_check(doc.paragraphs, formatter.roles)
    numbered_objects = numbered_object_check(doc, formatter.roles)
    front_matter = front_matter_check(doc.paragraphs, formatter.roles, profile, formatter.cover)
    caption_objects = caption_object_check(doc, formatter.roles, None, formatter.cover)
    import apa7_statistics
    statistics_review = apa7_statistics.format_statistics_and_equations(
        doc, formatter.roles, {}, profile, apply_safe_changes=False
    )
    object_summary["statistical_expressions"] = statistics_review["expressions_found"]
    object_summary["native_math_paragraphs"] = statistics_review["equations"]["native_math_paragraphs"]
    object_summary["plain_text_formula_candidates"] = statistics_review["equations"]["plain_text_formula_candidates"]
    object_summary["reference_live_links"] = reference_links["live_links_after"]
    object_summary["reference_unlinked_urls"] = len(reference_links["unlinked_urls"])
    object_summary["reference_bare_dois"] = len(reference_links["bare_dois"])
    object_summary["reference_quality_issues"] = len(reference_quality["issues"])
    object_summary["numbered_table_labels"] = numbered_objects["counts"]["table"]["labels"]
    object_summary["numbered_figure_labels"] = numbered_objects["counts"]["figure"]["labels"]
    object_summary["numbered_equation_labels"] = numbered_objects["counts"]["equation"]["labels"]
    object_summary["caption_object_issues"] = len(caption_objects["issues"])
    object_summary["front_matter_issues"] = len(front_matter["issues"])
    preflight_summary = combined_preflight_summary(
        front_matter=front_matter,
        citation_reference=citation_review,
        reference_links=reference_links,
        reference_quality=reference_quality,
        statistics_formula=statistics_review,
        numbered_objects=numbered_objects,
        caption_objects=caption_objects,
    )
    paragraph_ids = {p._p: i for i, p in enumerate(doc.paragraphs, 1)}
    table_ids = {t._tbl: i for i, t in enumerate(doc.tables, 1)}
    body_order = []
    for element in doc._element.body:
        if element in paragraph_ids:
            body_order.append({"kind": "paragraph", "paragraph": paragraph_ids[element]})
        elif element in table_ids:
            body_order.append({"kind": "table", "table": table_ids[element]})
        elif element.tag != qn("w:sectPr"):
            body_order.append({"kind": "unsupported_container", "tag": etree.QName(element).localname,
                               "text": descendant_text(element)})
    if digest(source) != before:
        raise RuntimeError("读取期间原稿发生变化，请重新生成配置。")
    return {"source_sha256": before, "roles": {}, "run_in_headings": {},
            "table_header_rows": {}, "table_roles": {}, "replace_headers": False,
            "_review": {"source": str(source), "profile": profile,
                        "instructions": "下方仅为推定，不会自动写入 roles。确认后把需要覆盖的段落编号与角色填入顶层 roles；原稿变化后必须重新生成配置。",
                        "suggested_title_page": [i + 1 for i in formatter.cover] if formatter.cover else None,
                        "allowed_roles": sorted(ROLES), "paragraphs": paragraphs,
                        "tables": tables, "body_order": body_order, "object_summary": object_summary,
                        "hierarchy_outline": hierarchy_outline, "citation_reference_check": citation_review,
                        "reference_link_check": reference_links,
                        "reference_quality_check": reference_quality,
                        "front_matter_check": front_matter,
                        "numbered_object_check": numbered_objects,
                        "caption_object_check": caption_objects,
                        "statistics_formula_check": statistics_review,
                        "preflight_summary": preflight_summary,
                        "events": formatter.events, "checklist": review_checklist(profile)}}


def save_config_draft(source, destination, profile="student"):
    destination = Path(destination).expanduser().absolute()
    if destination.suffix.lower() != ".json":
        raise ValueError("结构配置输出必须使用 .json 扩展名。")
    if destination.exists() or destination.is_symlink():
        raise ValueError("配置输出已存在，不能覆盖原稿或已有配置。")
    draft = prepare_config(source, profile)
    with destination.open("x", encoding="utf-8") as output:
        json.dump(draft, output, ensure_ascii=False, indent=2)
    return destination


def verify_saved_format(doc, formatter):
    """Verify core formatting after the DOCX has been saved and reopened."""
    checks, failures = [], []

    def result(key, label, status, details, rule):
        checks.append({"id": key, "label": label, "status": status, "details": details,
                       "source_url": SOURCES[rule]["url"]})

    margin_issues = []
    for number, section in enumerate(doc.sections, 1):
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            value = getattr(section, side)
            if value is None or abs(value.inches - 1.0) > 0.002:
                margin_issues.append(f"section{number}.{side}")
    if margin_issues:
        failures.append("页边距未保存到 1 英寸：" + "、".join(margin_issues))
    else:
        result("margins", "页边距", "passed", f"{len(doc.sections)} 个节的四边页边距均为 1 英寸。", "margins")

    centered = {"title", "title_meta", "section", "appendix", "appendix_title", "heading1"}
    no_indent = centered | {"heading2", "heading3", "title_meta", "abstract",
                            "caption_number", "caption_title", "note"}
    paragraph_issues, font_issues, verified_paragraphs = [], [], 0
    for index, paragraph in enumerate(doc.paragraphs):
        role = formatter.roles.get(index, "preserve")
        text = visible_text(paragraph).strip()
        if role == "preserve" or role in {"heading4", "heading5", "equation"} or (paragraph._p.xpath(".//wp:inline") and not text):
            continue
        verified_paragraphs += 1
        expected_align = WD_ALIGN_PARAGRAPH.CENTER if role in centered else WD_ALIGN_PARAGRAPH.LEFT
        expected_left = 0.5 if role in {"reference", "quote", "quote_continuation"} else 0.0
        if role == "reference":
            expected_first = -0.5
        elif role == "quote":
            expected_first = 0.0
        elif role == "quote_continuation":
            expected_first = 0.5
        elif role in no_indent:
            expected_first = 0.0
        else:
            expected_first = 0.5
        pf = paragraph.paragraph_format
        actual_left = 0.0 if pf.left_indent is None else pf.left_indent.inches
        actual_first = 0.0 if pf.first_line_indent is None else pf.first_line_indent.inches
        if pf.alignment != expected_align or abs(actual_left - expected_left) > 0.002 or abs(actual_first - expected_first) > 0.002:
            paragraph_issues.append(str(index + 1))
        if not isinstance(pf.line_spacing, float) or abs(pf.line_spacing - 2.0) > 0.001:
            paragraph_issues.append(str(index + 1))
        for run in all_runs(paragraph):
            if not run.text or run.font.name in {"Symbol", "Wingdings", "Cambria Math", "Courier New", "Consolas"}:
                continue
            size = run.font.size.pt if run.font.size is not None else None
            if run.font.name != formatter.font or size is None or abs(size - FONTS[formatter.font]) > 0.05:
                font_issues.append(str(index + 1))
                break
    paragraph_issues = sorted(set(paragraph_issues), key=int)
    font_issues = sorted(set(font_issues), key=int)
    if paragraph_issues:
        failures.append("段落对齐、缩进或双倍行距保存异常：第 " + _ranges(map(int, paragraph_issues)) + " 段")
    else:
        result("paragraph_layout", "段落格式", "passed", f"已重新打开并核验 {verified_paragraphs} 个已处理段落。", "paragraph")
    if font_issues:
        failures.append("字体或字号保存异常：第 " + _ranges(map(int, font_issues)) + " 段")
    else:
        result("typography", "字体与字号", "passed", f"已处理段落使用 {formatter.font} {FONTS[formatter.font]} pt；特殊符号和代码字体保留。", "font")

    table_issues, verified_tables = [], 0
    table_roles = formatter.config.get("table_roles", {})
    for number, table in enumerate(doc.tables, 1):
        if table_roles.get(str(number)) == "preserve" or table._tbl.xpath(
                ".//w:gridSpan | .//w:vMerge | .//w:hMerge | ./w:tr/w:trPr/w:gridBefore | ./w:tr/w:trPr/w:gridAfter | .//w:tbl"):
            continue
        verified_tables += 1
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        values = {} if borders is None else {etree.QName(edge).localname: edge.get(qn("w:val")) for edge in borders}
        if values.get("top") != "single" or values.get("bottom") != "single" or any(
                values.get(edge) != "nil" for edge in ("left", "right", "insideV")):
            table_issues.append(str(number))
    if table_issues:
        failures.append("表格边框保存异常：表格 " + "、".join(table_issues))
    else:
        result("table_rules", "表格横线", "passed",
               f"已重新打开并核验 {verified_tables} 个简单数据表；保留型或复杂表不计入自动通过。", "tables")

    if formatter.add_styles:
        style_issues = []
        for name, spec in APA7_PARAGRAPH_STYLES.items():
            try:
                style = doc.styles[name]
            except KeyError:
                style_issues.append(name + " 缺失")
                continue
            pf, sf = style.paragraph_format, style.font
            left = 0.0 if pf.left_indent is None else pf.left_indent.inches
            first = 0.0 if pf.first_line_indent is None else pf.first_line_indent.inches
            size = sf.size.pt if sf.size is not None else None
            if (style.type != WD_STYLE_TYPE.PARAGRAPH or sf.name != formatter.font or size is None
                    or abs(size - FONTS[formatter.font]) > 0.05 or sf.bold is not spec["bold"]
                    or sf.italic is not spec["italic"] or pf.alignment != spec["align"]
                    or abs(left - spec["left"]) > 0.002 or abs(first - spec["indent"]) > 0.002
                    or not isinstance(pf.line_spacing, float) or abs(pf.line_spacing - 2.0) > 0.001):
                style_issues.append(name)
        for name, spec in APA7_CHARACTER_STYLES.items():
            try:
                style = doc.styles[name]
            except KeyError:
                style_issues.append(name + " 缺失")
                continue
            size = style.font.size.pt if style.font.size is not None else None
            if (style.type != WD_STYLE_TYPE.CHARACTER or style.font.name != formatter.font or size is None
                    or abs(size - FONTS[formatter.font]) > 0.05 or style.font.bold is not spec["bold"]
                    or style.font.italic is not spec["italic"]):
                style_issues.append(name)
        for index, name in formatter.applied_styles.items():
            if index >= len(doc.paragraphs) or doc.paragraphs[index].style.name != name:
                style_issues.append(f"第 {index + 1} 段未绑定 {name}")
        if style_issues:
            failures.append("可继续编辑的 APA7 样式保存异常：" + "、".join(style_issues))
        else:
            result("reusable_styles", "可继续编辑样式", "passed",
                   f"已重新打开并确认 {len(APA7_PARAGRAPH_STYLES)} 个段落样式和 {len(APA7_CHARACTER_STYLES)} 个同行标题样式。", "paragraph")

    missing_page, missing_head = [], []
    if not doc.settings.odd_and_even_pages_header_footer:
        failures.append("奇偶页页眉变体未激活；不同渲染器可能错误使用旧的偶数页页眉。")
    seen = set()
    for section_number, section in enumerate(doc.sections, 1):
        for label, container in (("default", section.header), ("first", section.first_page_header),
                                 ("even", section.even_page_header)):
            root = container._element
            if root in seen:
                continue
            seen.add(root)
            instructions = " ".join(root.xpath(".//w:instrText/text()") + root.xpath(".//w:fldSimple/@w:instr"))
            if not re.search(r"\bPAGE\b", instructions, re.I):
                missing_page.append(f"section{section_number}.{label}")
            if formatter.profile == "professional" and formatter.running_head:
                header_text = "".join(root.xpath(".//w:t/text()"))
                if formatter.running_head not in header_text:
                    missing_head.append(f"section{section_number}.{label}")
    header_notes = []
    if missing_page:
        header_notes.append("以下页眉变体没有确认到自动页码：" + "、".join(missing_page))
    if formatter.profile == "professional" and not formatter.running_head:
        header_notes.append("专业论文尚未提供 running head。")
    elif missing_head:
        header_notes.append("以下页眉变体没有确认到 running head：" + "、".join(missing_head))
    result("page_header", "页码与页眉", "needs_review" if header_notes else "passed",
           "；".join(header_notes) if header_notes else "已重新打开并确认自动 PAGE 域及所选模式的页眉。", "header")

    if failures:
        raise RuntimeError("保存后格式核验失败，未交付文件；" + "；".join(failures))
    return checks


def _ranges(numbers):
    """Turn paragraph numbers into short labels such as 2–4、7."""
    numbers = sorted(set(numbers))
    groups = []
    for number in numbers:
        if not groups or number != groups[-1][-1] + 1:
            groups.append([number])
        else:
            groups[-1].append(number)
    return "、".join(str(group[0]) if len(group) == 1 else f"{group[0]}–{group[-1]}" for group in groups)


def simple_feedback(report):
    """Build the four short, student-facing sections used by the CLI and Skill."""
    roles = report.get("paragraph_roles", {})
    by_role = {}
    for number, role in roles.items():
        by_role.setdefault(role, []).append(int(number))
    inventory = report.get("visuals", {}).get("inventory", {}).get("counts", {})
    table_count = inventory.get("native_tables", 0)
    figure_count = inventory.get("native_charts", 0) + inventory.get("raster_media", 0) + inventory.get("vector_media", 0)
    citation_check = report.get("citation_reference_check") or {}
    reference_links = report.get("reference_link_check") or {}
    reference_quality = report.get("reference_quality_check") or {}
    front_matter = report.get("front_matter_check") or {}
    numbered = report.get("numbered_object_check") or {}
    caption_objects = report.get("caption_object_check") or {}
    preflight = report.get("preflight_summary") or {}
    statistics = report.get("statistical_reporting") or {}
    equations = statistics.get("equations") or {}

    changed = ["页面和正文：统一页边距、字体、双倍行距、段落间距和页码。"]
    if report.get("reusable_styles_added"):
        changed.append("继续写作：已按你的要求加入 APA7 正文、标题、参考文献和图表说明样式。")
    if report["profile"] == "student":
        changed.append("页眉：按学生论文模式处理页码。")
    elif report.get("running_head"):
        changed.append("页眉：按专业论文模式处理 running head 和页码。")
    else:
        changed.append("页眉：已处理页码；尚未添加缺失的 running head。")
    if by_role.get("title") or by_role.get("title_meta"):
        changed.append("标题页：统一标题和作者信息的对齐、粗体及间距。")
    if any(by_role.get(f"heading{i}") for i in range(1, 6)):
        changed.append("标题：按识别出的层级设置 APA 标题样式。")
    if by_role.get("reference"):
        changed.append("参考文献：设置双倍行距和 0.5 英寸悬挂缩进。")
    if reference_links.get("links_added"):
        changed.append(f"参考文献链接：将 {len(reference_links['links_added'])} 个已有 DOI／URL 设为可点击链接，显示文字未改变。")
    if reference_quality.get("status") != "not_applicable":
        changed.append(
            f"参考文献质量：检查了 {reference_quality.get('reference_entries_checked', 0)} 条文献的重复、字母顺序和同作者同年后缀；没有移动或改写条目。"
        )
    if table_count:
        changed.append(f"表格：处理了 {table_count} 个可编辑表格的对齐和边框。")
    if figure_count:
        changed.append(f"图和图表：识别了 {figure_count} 个对象；过宽图片才会等比例缩小。")
    if statistics.get("expressions_found"):
        changed.append(
            f"统计数据：识别 {statistics['expressions_found']} 处统计表达，规范了 "
            f"{len(statistics.get('safe_text_edits', []))} 处写法和 {statistics.get('symbols_formatted', 0)} 个统计符号；数值未改变。"
        )
    if equations.get("native_math_paragraphs") or equations.get("plain_text_formula_candidates"):
        changed.append(
            f"公式：识别 {equations.get('native_math_paragraphs', 0)} 个含 Word 原生公式的段落和 "
            f"{equations.get('plain_text_formula_candidates', 0)} 个纯文本公式候选；没有重写公式内容。"
        )
    if numbered.get("status") != "not_applicable":
        total_labels = sum(item.get("labels", 0) for item in numbered.get("counts", {}).values())
        total_callouts = sum(item.get("callouts", 0) for item in numbered.get("counts", {}).values())
        changed.append(
            f"编号检查：核对了 {total_labels} 个图表／公式编号和 {total_callouts} 处正文提及；没有自动重新编号。"
        )
    if preflight:
        changed.append(
            f"完整预检：一次检查标题页、引用、参考文献、统计、公式和图表，共有 {preflight.get('issue_count', 0)} 项需要人工判断。"
        )

    where = ["全文：页面、正文和页眉格式。"]
    role_labels = {
        "title": "论文标题", "title_meta": "标题页信息", "section": "章节标题",
        "heading1": "一级标题", "heading2": "二级标题", "heading3": "三级标题",
        "heading4": "四级标题", "heading5": "五级标题", "reference": "参考文献条目",
        "caption_number": "图表编号", "caption_title": "图表标题", "note": "图表注释",
        "abstract": "摘要", "keywords": "关键词", "quote": "块引用",
        "quote_continuation": "连续块引用", "appendix": "附录标签", "appendix_title": "附录标题",
        "equation": "独立公式",
    }
    for role, label in role_labels.items():
        if by_role.get(role):
            where.append(f"原稿第 {_ranges(by_role[role])} 段：{label}格式。")
    if table_count:
        where.append(f"原稿中的 {table_count} 个表格：保留数据内容；统计表达可能只调整符号、空格和前导零。")
    if figure_count:
        where.append(f"原稿中的 {figure_count} 个图或图表：保留原始内容，仅检查尺寸和位置。")
    if statistics.get("changed_locations"):
        locations = "、".join(statistics["changed_locations"][:12])
        suffix = "等位置" if len(statistics["changed_locations"]) > 12 else ""
        where.append(f"原稿 {locations}{suffix}：统计符号、空格、前导零或公式段落格式。")
    if reference_links.get("linked_locations"):
        locations = "、".join(reference_links["linked_locations"])
        where.append(f"原稿 {locations}：已有 DOI／URL 增加可点击链接，文字未改变。")

    source_keys = ["font", "margins", "spacing", "paragraph", "header"]
    if by_role.get("title") or by_role.get("title_meta"):
        source_keys.append("title")
    if any(by_role.get(f"heading{i}") for i in range(1, 6)):
        source_keys.append("headings")
    if by_role.get("reference"):
        source_keys.append("references")
    if citation_check.get("citations_found") or citation_check.get("reference_entries"):
        source_keys.append("citation_match")
    if reference_links.get("status") != "not_applicable":
        source_keys.append("reference_links")
    if reference_quality.get("status") != "not_applicable":
        source_keys.append("reference_order")
        if reference_quality.get("same_author_year_groups"):
            source_keys.append("same_author_date")
    if front_matter:
        source_keys.append("title")
        if report.get("profile") == "student":
            source_keys.append("student_title_elements")
    if table_count:
        source_keys.append("tables")
    if figure_count:
        source_keys.append("figures")
    if statistics.get("expressions_found"):
        source_keys.append("statistics")
        if report.get("profile") == "professional":
            source_keys.append("jars")
    if equations.get("native_math_paragraphs") or equations.get("plain_text_formula_candidates"):
        source_keys.append("equations")
    numbered_counts = numbered.get("counts", {})
    if any(numbered_counts.get("table", {}).values()) and "tables" not in source_keys:
        source_keys.append("tables")
    if any(numbered_counts.get("figure", {}).values()) and "figures" not in source_keys:
        source_keys.append("figures")
    if any(numbered_counts.get("equation", {}).values()) and "equations" not in source_keys:
        source_keys.append("equations")
    if caption_objects.get("status") != "not_applicable":
        if any(item.get("kind") == "table" for item in caption_objects.get("objects", [])) and "tables" not in source_keys:
            source_keys.append("tables")
        if any(item.get("kind") == "figure" for item in caption_objects.get("objects", [])) and "figures" not in source_keys:
            source_keys.append("figures")
    source_keys = list(dict.fromkeys(source_keys))
    sources = [{"title": SOURCES[key]["title"], "url": SOURCES[key]["url"],
                "quote": SOURCES[key]["quote"]} for key in source_keys]

    ai_review = report.get("ai_review") or {}
    compliance = ai_review.get("compliance_review") or {}
    labels = {"target_requirements": "目标要求", "title_page": "标题页",
              "abstract_keywords": "摘要与关键词", "headings": "标题层级",
              "citations_references": "引文与参考文献", "statistics": "统计数据",
              "equations": "公式", "tables": "表格", "figures": "图片与图表", "appendices": "附录"}
    if compliance:
        review = [f"{labels.get(key, key)}：{item['reason']}" for key, item in compliance.items()
                  if item.get("status") == "needs_review"]
        review.extend(ai_review.get("unresolved", []))
    else:
        review = ["确认标题页姓名、学校、课程、教师和日期等信息是否完整。"
                  if report["profile"] == "student" else
                  "确认作者单位、作者注和 running head 是否符合投稿要求。",
                  "确认标题层级是否判断正确。",
                  "核对正文引用与参考文献内容、顺序、DOI 和斜体。"]
        if table_count:
            review.append("核对表格编号、标题、表头、注释及正文提及顺序。")
        if figure_count:
            review.append("核对图的编号、标题、清晰度、图例、单位和版权说明。")
    review.extend(check["details"] for check in report.get("machine_checks", [])
                  if check.get("status") == "needs_review"
                  and check.get("id") not in {"reference_links", "statistical_reporting",
                                               "reference_quality", "numbered_object_references",
                                               "front_matter", "caption_object_pairing"})
    unmatched = citation_check.get("unmatched_citations", [])
    uncited = citation_check.get("uncited_references", [])
    unparsed = citation_check.get("unparsed_reference_paragraphs", [])
    if unmatched:
        paragraphs = [p for item in unmatched for p in item.get("paragraphs", [])]
        review.append(f"引文与参考文献：{len(unmatched)} 组正文作者—年份未找到明显对应条目（第 {_ranges(paragraphs)} 段）。")
    if uncited:
        paragraphs = [p for item in uncited for p in item.get("paragraphs", [])]
        review.append(f"引文与参考文献：{len(uncited)} 条参考文献未找到明显正文引文（第 {_ranges(paragraphs)} 段）。")
    if unparsed and reference_quality.get("status") == "not_applicable":
        review.append(f"引文与参考文献：第 {_ranges(unparsed)} 段未识别出清晰的作者—年份，需要人工核对。")
    if reference_links.get("issues"):
        review.append("参考文献链接：" + " ".join(reference_links["issues"]))
    reference_issues = reference_quality.get("issues", [])
    if reference_issues:
        examples = "；".join(item["message"] for item in reference_issues[:2])
        suffix = "；其余问题请在 Word 中核对。" if len(reference_issues) > 2 else ""
        review.append(f"参考文献质量：发现 {len(reference_issues)} 项需要确认。{examples}{suffix}")
    front_issues = front_matter.get("issues", [])
    if front_issues:
        examples = "；".join(item["message"] for item in front_issues[:2])
        suffix = "；其余项目请结合标题页和分页核对。" if len(front_issues) > 2 else ""
        review.append(f"标题页与前置页：发现 {len(front_issues)} 项需要确认。{examples}{suffix}")
    statistic_issues = statistics.get("issues", [])
    if statistic_issues:
        examples = "；".join(item["message"] for item in statistic_issues[:2])
        review.append(f"统计数据与公式：发现 {len(statistic_issues)} 项需要确认。{examples}")
    numbered_issues = numbered.get("issues", [])
    if numbered_issues:
        examples = "；".join(item["message"] for item in numbered_issues[:2])
        suffix = "；其余问题请结合 Word 页面核对。" if len(numbered_issues) > 2 else ""
        review.append(f"图表与公式编号：发现 {len(numbered_issues)} 项需要确认。{examples}{suffix}")
    caption_issues = caption_objects.get("issues", [])
    if caption_issues:
        examples = "；".join(item["message"] for item in caption_issues[:2])
        suffix = "；其余对象请结合 Word 页面核对。" if len(caption_issues) > 2 else ""
        review.append(f"图表与说明：发现 {len(caption_issues)} 项需要确认。{examples}{suffix}")
    review.append("最后在 Word 中逐页看一遍分页和学校或期刊的特殊要求。")
    if report.get("visuals", {}).get("export_error"):
        review.append("矢量导出没有完成，需要重新处理。")

    review = list(dict.fromkeys(item for item in review if item))
    return {"summary": "已生成 1 个新的 Word 格式副本；原稿、研究内容和统计数值没有改变。",
            "changed": changed, "apa_sources": sources, "locations": where,
            "needs_review": review}


def feedback_text(feedback):
    lines = [feedback["summary"], "", "改了什么："]
    lines.extend("- " + item for item in feedback["changed"])
    lines.extend(["", "APA 来源："])
    lines.extend(f"- {item['title']}：{item['url']}\n  原文：{item['quote']}" for item in feedback["apa_sources"])
    lines.extend(["", "改了原稿哪里：", "- 没有改变研究内容或统计数值，只修改下面位置的安全格式："])
    lines.extend("- " + item for item in feedback["locations"])
    lines.extend(["", "还要审核："])
    lines.extend("- " + item for item in feedback["needs_review"])
    return "\n".join(lines)


def report_html(report):
    """Optional concise HTML; normal runs create only the DOCX."""
    esc = html.escape
    feedback = report["feedback"]
    items = lambda values: "".join(f"<li>{esc(value)}</li>" for value in values)
    sources = "".join(f"<li><a href='{esc(source['url'], quote=True)}'>{esc(source['title'])}</a>"
                      f"<blockquote>{esc(source['quote'])}</blockquote></li>" for source in feedback["apa_sources"])
    return f"""<!doctype html><html lang='zh-CN'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>
<title>APA 7 简要反馈</title><style>body{{max-width:760px;margin:36px auto;padding:0 20px;font:16px/1.65 system-ui;color:#17212b}}blockquote{{border-left:3px solid #7288a2;padding-left:12px;margin-left:0}}a{{color:#155a91}}li{{margin:8px 0}}</style>
<h1>APA 7 简要反馈</h1><p>{esc(feedback['summary'])}</p>
<h2>改了什么</h2><ul>{items(feedback['changed'])}</ul>
<h2>APA 来源</h2><ul>{sources}</ul>
<h2>改了原稿哪里</h2><p>没有改变研究内容或统计数值，只修改下面位置的安全格式：</p><ul>{items(feedback['locations'])}</ul>
<h2>还要审核</h2><ul>{items(feedback['needs_review'])}</ul></html>"""


def convert_legacy(source, temp, soffice):
    executable = soffice or shutil.which("soffice")
    if not executable:
        raise ValueError("旧版 .doc 需先用 Word 另存为 .docx，或通过 --soffice 指定 LibreOffice 命令。")
    profile = Path(temp) / "lo_profile"
    proc = subprocess.run([str(executable), "-env:UserInstallation=" + profile.as_uri(), "--headless", "--convert-to", "docx", "--outdir", str(temp), str(source)], capture_output=True, text=True, timeout=120)
    out = Path(temp) / (source.stem + ".docx")
    if proc.returncode or not out.is_file():
        raise ValueError("旧版 Word 转换失败：" + (proc.stderr or proc.stdout)[-1000:])
    return out


def format_file(source, *, output=None, profile=None, font="Times New Roman", running_head="", config=None,
                soffice=None, export_visuals=False, add_styles=False, save_feedback=False):
    source = Path(source).expanduser().resolve()
    if not source.is_file() or source.suffix.lower() not in {".docx", ".doc"}:
        raise ValueError("请选择现有 .docx 或 .doc 文件。加密文档、.docm 宏文件及 RTF 暂不支持。")
    if profile not in {"student", "professional"}:
        raise ValueError("请先选择论文类型：student（学生论文）或 professional（专业／投稿论文）。")
    if font not in FONTS:
        raise ValueError("无效字体。")
    out = unique_output(source, output, save_feedback)
    before = digest(source)
    config = {} if config is None else config
    if not isinstance(config, dict):
        raise ValueError("配置文件顶层必须是 JSON 对象。")
    if "source_sha256" in config:
        bound_hash = config["source_sha256"]
        if not isinstance(bound_hash, str) or not re.fullmatch(r"[0-9a-fA-F]{64}", bound_hash):
            raise ValueError("配置中的 source_sha256 必须为 64 位 SHA256。")
        if bound_hash.lower() != before:
            raise ValueError("配置不属于当前版本原稿，段落编号可能已变化。请重新生成并确认配置；未输出文档。")
    visuals_api = None
    try:
        import apa7_visuals as visuals_api
    except ImportError:
        if export_visuals:
            raise ValueError("矢量功能需要把 apa7_visuals.py 放在主程序旁边，并安装 reportlab。")
    visual_dir = out.with_suffix(".visuals")
    if export_visuals and visual_dir.exists():
        raise ValueError("同名图形导出文件夹已存在，请更换输出名。")
    with tempfile.TemporaryDirectory(prefix="apa7_") as tmp:
        converted = source.suffix.lower() == ".doc"
        input_docx = convert_legacy(source, tmp, soffice) if converted else source
        doc = Document(input_docx)
        signature = content_signature(doc)
        source_links = document_hyperlinks(doc)
        resources = package_payloads(input_docx)
        formatter = Formatter(doc, profile, font, running_head, config, add_styles=add_styles)
        formatter.run()
        expected_signature = content_signature(doc)
        expected_links = document_hyperlinks(doc)
        static_source = [item for item in signature if item[0] not in {"w:t", "w:hyperlink"}]
        static_expected = [item for item in expected_signature if item[0] not in {"w:t", "w:hyperlink"}]
        source_link_counts = Counter((item["location"], item["text"], item["target"]) for item in source_links)
        expected_link_counts = Counter((item["location"], item["text"], item["target"]) for item in expected_links)
        recorded_link_counts = Counter(
            (item["location"], item["text"], item["target"])
            for item in (formatter.reference_link_report or {}).get("links_added", [])
        )
        source_hyperlink_nodes = dict(signature)["w:hyperlink"]
        expected_hyperlink_nodes = dict(expected_signature)["w:hyperlink"]
        if (static_source != static_expected
                or not _ordered_subsequence(source_hyperlink_nodes, expected_hyperlink_nodes)
                or len(expected_hyperlink_nodes) - len(source_hyperlink_nodes) != sum(recorded_link_counts.values())
                or expected_link_counts - source_link_counts != recorded_link_counts
                or source_link_counts - expected_link_counts):
            raise RuntimeError("格式处理改变了未获允许的内容结构，未交付文件；原稿未改动。")
        if converted:
            formatter.event("review", "输入经 LibreOffice 从 .doc 转为 .docx；内容保留检查以转换后版本为基准，旧格式转换保真需人工检查。")
        saved = Path(tmp) / "formatted.docx"
        doc.save(saved)
        check = Document(saved)
        # Run splitting may alter w:t boundaries. Statistical presentation and
        # newly wrapped reference hyperlinks are verified against the exact
        # in-memory document that passed the allowlist checks above.
        after_signature = content_signature(check)
        after_text = "".join(x[0] or "" for x in after_signature[0][1])
        if (after_text != formatter.expected_body_text or expected_signature != after_signature
                or document_hyperlinks(check) != expected_links):
            raise RuntimeError("内容保留检查失败，未交付格式化文件；原稿未改动。")
        if resources != package_payloads(saved) or digest(source) != before:
            raise RuntimeError("原稿或资源保留检查失败，未交付格式化文件。")
        machine_checks = [{"id": "content_preservation", "label": "原文与资源",
                           "status": "passed", "details": "研究内容、统计数值、域、公式、图片及嵌入资源通过保存后保留检查；安全的统计呈现与参考文献链接修改已单独记录。",
                           "source_url": None}]
        machine_checks.extend(verify_saved_format(check, formatter))
        statistics_report = formatter.statistics_report or {}
        if statistics_report.get("status") != "not_applicable":
            machine_checks.append({
                "id": "statistical_reporting",
                "label": "统计数据与公式",
                "status": statistics_report.get("status", "needs_review"),
                "details": (
                    f"识别 {statistics_report.get('expressions_found', 0)} 处统计表达和 "
                    f"{statistics_report.get('equations', {}).get('native_math_paragraphs', 0)} 个含原生公式的段落；"
                    f"{len(statistics_report.get('issues', []))} 项需确认。"
                ),
                "source_url": SOURCES["statistics"]["url"],
            })
        reference_link_report = formatter.reference_link_report or {}
        if reference_link_report.get("status") != "not_applicable":
            machine_checks.append({
                "id": "reference_links",
                "label": "参考文献 DOI／URL 链接",
                "status": reference_link_report.get("status", "needs_review"),
                "details": (
                    f"新增 {len(reference_link_report.get('links_added', []))} 个可点击链接；"
                    f"仍有 {len(reference_link_report.get('unlinked_urls', []))} 个完整 URL 和 "
                    f"{len(reference_link_report.get('bare_dois', []))} 个裸 DOI 需要确认。"
                ),
                "source_url": SOURCES["reference_links"]["url"],
            })
        reference_quality_report = formatter.reference_quality_report or reference_quality_check(
            check.paragraphs, formatter.roles
        )
        if reference_quality_report.get("status") != "not_applicable":
            machine_checks.append({
                "id": "reference_quality",
                "label": "参考文献质量",
                "status": reference_quality_report["status"],
                "details": (
                    f"检查 {reference_quality_report.get('reference_entries_checked', 0)} 条参考文献；"
                    f"{len(reference_quality_report.get('issues', []))} 项需确认，不会自动移动或改写条目。"
                ),
                "source_url": SOURCES["reference_order"]["url"],
            })
        numbered_report = numbered_object_check(
            check, formatter.roles, formatter.config.get("table_roles")
        )
        if numbered_report.get("status") != "not_applicable":
            total_labels = sum(item["labels"] for item in numbered_report["counts"].values())
            total_callouts = sum(item["callouts"] for item in numbered_report["counts"].values())
            machine_checks.append({
                "id": "numbered_object_references",
                "label": "图表与公式编号",
                "status": numbered_report["status"],
                "details": (
                    f"检查 {total_labels} 个编号标签和 {total_callouts} 处正文提及；"
                    f"{len(numbered_report['issues'])} 项需确认，不会自动重新编号。"
                ),
                "source_url": SOURCES["figures"]["url"],
            })
        front_matter_report = formatter.front_matter_report or front_matter_check(
            check.paragraphs, formatter.roles, profile, formatter.cover
        )
        machine_checks.append({
            "id": "front_matter",
            "label": "标题页、摘要与关键词",
            "status": front_matter_report["status"],
            "details": (
                f"标题页{'已确认' if front_matter_report['title_page_confirmed'] else '未确认'}；"
                f"{len(front_matter_report['issues'])} 项需确认，不会补写缺失信息。"
            ),
            "source_url": SOURCES["title"]["url"],
        })
        caption_object_report = formatter.caption_object_report or caption_object_check(
            check, formatter.roles, formatter.config.get("table_roles"), formatter.cover
        )
        if caption_object_report.get("status") != "not_applicable":
            machine_checks.append({
                "id": "caption_object_pairing",
                "label": "图表与说明配对",
                "status": caption_object_report["status"],
                "details": (
                    f"检查 {caption_object_report['objects_checked']} 个表格／图形对象；"
                    f"配对 {caption_object_report['paired_objects']} 个，"
                    f"{len(caption_object_report['issues'])} 项需确认。"
                ),
                "source_url": SOURCES["figures"]["url"],
            })
        citation_check = citation_reference_check(check.paragraphs, formatter.roles)
        preflight_summary = combined_preflight_summary(
            front_matter=front_matter_report,
            citation_reference=citation_check,
            reference_links=reference_link_report,
            reference_quality=reference_quality_report,
            statistics_formula=statistics_report,
            numbered_objects=numbered_report,
            caption_objects=caption_object_report,
        )
        visual_result = {}
        if visuals_api:
            try:
                visual_result["inventory"] = visuals_api.inspect_visuals(input_docx)
            except Exception as exc:
                formatter.event("review", "图形清单读取未完成：" + str(exc), "figures")
        if export_visuals:
            try:
                visual_result["exports"] = visuals_api.export_visuals(input_docx, visual_dir)
                formatter.event("applied", "图形导出清单已生成：" + str(visual_dir) + "。请查看各对象的成功、保留位图或不支持状态。", "figures")
            except Exception as exc:
                formatter.event("review", "Word 格式处理成功，但图形导出未完成：" + str(exc), "figures")
                visual_result["export_error"] = str(exc)
        report = {"version": VERSION, "created_at": datetime.now(timezone.utc).isoformat(),
                  "status": "FORMAT_APPLIED_REVIEW_REQUIRED", "input": str(source), "output": str(out),
                  "profile": profile, "font": font, "running_head": formatter.running_head or None,
                  "reusable_styles_added": add_styles,
                  "source_sha256": before,
                  "preservation": "passed (legacy conversion excluded)" if converted else "passed",
                  "paragraph_roles": {str(i + 1): r for i, r in formatter.roles.items()},
                  "counts": dict(formatter.counts), "events": formatter.events, "sources": SOURCES,
                  "configuration_source_bound": "source_sha256" in config,
                  "ai_review": config.get("ai_review"),
                  "machine_checks": machine_checks,
                  "citation_reference_check": citation_check,
                  "reference_link_check": reference_link_report,
                  "reference_quality_check": reference_quality_report,
                  "front_matter_check": front_matter_report,
                  "numbered_object_check": numbered_report,
                  "caption_object_check": caption_object_report,
                  "preflight_summary": preflight_summary,
                  "statistical_reporting": statistics_report,
                  "review_checklist": review_checklist(profile),
                  "visuals": visual_result}
        report["feedback"] = simple_feedback(report)
        # Exclusive creation keeps an existing output safe even during concurrent runs.
        with out.open("xb") as dst, saved.open("rb") as src:
            shutil.copyfileobj(src, dst)
        if save_feedback:
            with out.with_suffix(".feedback.html").open("x", encoding="utf-8") as f:
                f.write(report_html(report))
    return out, report


def inspect_file(path):
    doc = Document(path)
    return [{"paragraph": i, "style": p.style.name if p.style else "", "text": visible_text(p)} for i, p in enumerate(doc.paragraphs, 1)]


def gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
        root = tk.Tk()
    except Exception as e:
        raise SystemExit(f"当前 Python 没有可用的 Tk 界面（{e}）。请用命令行：python3 apa7_format.py 论文.docx --profile student")
    root.title("APA 7 Word 格式助手")
    root.geometry("700x520")
    frame = ttk.Frame(root, padding=24)
    frame.pack(fill="both", expand=True)
    ttk.Label(frame, text="APA 7 Word 格式助手", font=("Helvetica", 20)).pack(anchor="w")
    ttk.Label(frame, text="选择文档和论文类型后，只生成 1 个新的 Word 格式副本。\n保留原稿和论文文字；图片内部文字、引用内容及最终分页仍需检查。", wraplength=620).pack(anchor="w", pady=12)
    path = tk.StringVar()
    row = ttk.Frame(frame)
    row.pack(fill="x", pady=8)
    ttk.Entry(row, textvariable=path).pack(side="left", fill="x", expand=True)
    ttk.Button(row, text="选择 Word", command=lambda: path.set(filedialog.askopenfilename(filetypes=[("Word", "*.docx *.doc")]))).pack(side="right", padx=(10, 0))
    mode, font, head = tk.StringVar(value=""), tk.StringVar(value="Times New Roman"), tk.StringVar()
    for label, variable, options in [("论文模式", mode, ("student", "professional")), ("字体", font, tuple(FONTS))]:
        row = ttk.Frame(frame)
        row.pack(fill="x", pady=5)
        ttk.Label(row, text=label, width=16).pack(side="left")
        ttk.Combobox(row, textvariable=variable, values=options, state="readonly").pack(side="left")
    ttk.Label(frame, text="student = 学生论文；professional = 投稿论文").pack(anchor="w")
    row = ttk.Frame(frame)
    row.pack(fill="x", pady=10)
    ttk.Label(row, text="投稿短标题（≤50 字符）").pack(side="left")
    ttk.Entry(row, textvariable=head).pack(side="left", fill="x", expand=True, padx=10)
    vector_choice, style_choice = tk.BooleanVar(value=False), tk.BooleanVar(value=False)
    ttk.Checkbutton(frame, text="同时提取图片／原始矢量，并将支持的原生图表导出 SVG", variable=vector_choice).pack(anchor="w", pady=8)
    ttk.Checkbutton(frame, text="可选：加入 APA7 样式，方便在成品中继续写作", variable=style_choice).pack(anchor="w", pady=4)
    status = tk.StringVar(value="无需联网，无需 API Key。")
    ttk.Label(frame, textvariable=status, wraplength=620).pack(anchor="w", pady=8)
    def start():
        values = (path.get(), mode.get(), font.get(), head.get(), vector_choice.get(), style_choice.get())
        if not values[0]:
            messagebox.showerror("请选择文件", "请先选择一个 Word 文档。")
            return
        if values[1] not in {"student", "professional"}:
            messagebox.showerror("请选择论文类型", "请先选择 student（学生论文）或 professional（专业／投稿论文）。")
            return
        button.config(state="disabled")
        status.set("正在处理…")
        def done(out, report):
            button.config(state="normal")
            status.set("已生成：" + out.name)
            feedback = report["feedback"]
            short_review = "\n".join("• " + item for item in feedback["needs_review"][:3])
            messagebox.showinfo("格式副本已生成", f"只生成了这 1 个 Word 副本：\n{out}\n\n还要审核：\n{short_review}")
        def failed(error):
            button.config(state="normal")
            status.set("处理未完成；原稿保留。")
            messagebox.showerror("未完成", error)
        def work():
            try:
                result = format_file(values[0], profile=values[1], font=values[2], running_head=values[3],
                                     export_visuals=values[4], add_styles=values[5])
                root.after(0, lambda: done(*result))
            except Exception as exc:
                error = str(exc)
                root.after(0, lambda: failed(error))
        threading.Thread(target=work, daemon=True).start()
    button = ttk.Button(frame, text="生成 APA 7 格式副本", command=start)
    button.pack(anchor="w", pady=10)
    root.mainloop()


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        gui()
        return 0
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", nargs="?", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--profile", choices=("student", "professional"), help="必选：student 学生论文；professional 专业／投稿论文")
    parser.add_argument("--font", choices=tuple(FONTS), default="Times New Roman")
    parser.add_argument("--running-head", default="")
    parser.add_argument("--config", type=Path, help="JSON: roles, title_page, run_in_headings, table_header_rows, replace_headers")
    parser.add_argument("--soffice", help="旧 .doc 转换器的可执行文件路径")
    parser.add_argument("--inspect", action="store_true", help="只打印带编号的正文段落，供配置 roles 使用")
    parser.add_argument("--prepare-config", type=Path, metavar="NEW_JSON", help="只读分析 .docx，生成绑定原稿的新结构配置草稿；不修改 Word")
    parser.add_argument("--inspect-visuals", action="store_true", help="只读取原生表格、图表和图片类型清单（.docx）")
    parser.add_argument("--export-visuals", action="store_true", help="同时提取原始图片／矢量，并把支持的原生图表重绘为 SVG")
    parser.add_argument("--add-styles", action="store_true", help="可选：加入 APA7 Word 样式，方便在成品中继续写作")
    parser.add_argument("--save-feedback", action="store_true", help="另存 1 个简短 HTML 反馈；默认不生成额外报告文件")
    parser.add_argument("--sources", action="store_true", help="打印官网原文短引文与链接")
    args = parser.parse_args(argv)
    try:
        if args.sources:
            print(json.dumps(SOURCES, ensure_ascii=False, indent=2))
            return 0
        if not args.input:
            parser.error("需要输入 Word 文件路径，或使用 --sources。")
        if not args.profile and not (args.inspect or args.inspect_visuals):
            raise ValueError("请先选择论文类型：--profile student（学生论文）或 --profile professional（专业／投稿论文）。")
        if args.prepare_config:
            destination = save_config_draft(args.input, args.prepare_config, args.profile)
            print(f"已生成结构配置草稿：{destination}\n推定结果在 _review 中；请确认 roles 后再用于格式处理。Word 原稿未改动。")
            return 0
        if args.inspect:
            print(json.dumps(inspect_file(args.input), ensure_ascii=False, indent=2))
            return 0
        if args.inspect_visuals:
            from apa7_visuals import inspect_visuals
            print(json.dumps(inspect_visuals(args.input), ensure_ascii=False, indent=2))
            return 0
        config = json.loads(args.config.read_text(encoding="utf-8")) if args.config else {}
        if not isinstance(config, dict):
            raise ValueError("配置文件顶层必须是 JSON 对象。")
        out, report = format_file(args.input, output=args.output, profile=args.profile, font=args.font,
                                  running_head=args.running_head, config=config, soffice=args.soffice,
                                  export_visuals=args.export_visuals, add_styles=args.add_styles,
                                  save_feedback=args.save_feedback)
        print(feedback_text(report["feedback"]))
        print(f"\nWord 副本：{out}")
        if args.save_feedback:
            print(f"简要反馈：{out.with_suffix('.feedback.html')}")
        if args.export_visuals:
            print("图形导出：" + ("未完成，需要重新处理" if report["visuals"].get("export_error") else str(out.with_suffix(".visuals"))))
        return 0
    except Exception as exc:
        print(f"未完成：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

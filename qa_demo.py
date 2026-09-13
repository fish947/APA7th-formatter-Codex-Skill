"""Generate an entirely synthetic fixture for render QA. Not a user paper."""
from pathlib import Path
import argparse
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from apa7_format import format_file


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    args.directory.mkdir(parents=True, exist_ok=True)
    image = args.directory / "fixture_plot.png"
    # A small deterministic plot fixture; production figures are left intact.
    canvas = Image.new("RGB", (900, 396), "white")
    drawing = ImageDraw.Draw(canvas)
    font = ImageFont.load_default(size=28)
    drawing.line([(90, 45), (90, 320), (850, 320)], fill="black", width=3)
    points = [(150, 270), (350, 150), (550, 210), (750, 90)]
    drawing.line(points, fill="black", width=4)
    for j, (x, y) in enumerate(points):
        drawing.ellipse((x - 6, y - 6, x + 6, y + 6), fill="black")
        drawing.text((x - 8, 325), str(j), fill="black", font=font)
    drawing.text((345, 360), "Trial Block", fill="black", font=font)
    drawing.text((100, 5), "Response (a.u.)", fill="black", font=font)
    canvas.save(image, dpi=(180, 180))
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 2
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph("Response Patterns in a Synthetic Task", "Title")
    doc.add_paragraph()
    for line in ["Alex Example", "Department of Psychology, Example University", "PSY 101: Research Methods", "Dr. Taylor Example", "September 12, 2026"]:
        doc.add_paragraph(line)
    p = doc.add_paragraph("Response Patterns in a Synthetic Task", "Title")
    p.paragraph_format.page_break_before = True
    doc.add_paragraph("This document contains synthetic test material for checking the formatter. It includes ordinary prose, a table, a figure, and a reference list. None of the results represent a real experiment. The formatting engine should retain every word, number, and image while changing the requested paragraph properties.")
    doc.add_paragraph("Method", "Heading 1")
    doc.add_paragraph("Participants", "Heading 2")
    p = doc.add_paragraph("All values are illustrative. The sample size was ")
    p.add_run("N").italic = True
    p.add_run(" = 40. The figure and table are provided to test layout and should remain readable after the conversion.")
    doc.add_paragraph("Results", "Heading 1")
    doc.add_paragraph("Table 1 summarizes two artificial conditions. Figure 1 displays an illustrative sequence.")
    doc.add_paragraph("Table 1")
    doc.add_paragraph("Synthetic Response Summary")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for row, values in zip(table.rows, [("Condition", "Mean", "SD"), ("A", "3.2", "0.8"), ("B", "2.4", "0.6")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    p = doc.add_paragraph()
    p.add_run("Note.").italic = True
    p.add_run(" Values are synthetic and used only for software testing.")
    p = doc.add_paragraph("Figure 1")
    p.paragraph_format.page_break_before = True
    doc.add_paragraph("Illustrative Response Across Trial Blocks")
    p = doc.add_paragraph()
    p.add_run().add_picture(str(image), width=Inches(5))
    p = doc.add_paragraph()
    p.add_run("Note.").italic = True
    p.add_run(" The figure contains synthetic values.")
    doc.add_paragraph("References")
    p = doc.add_paragraph("Example, A. (2026). ")
    p.add_run("Synthetic document for testing a formatter").italic = True
    p.add_run(". Example Publisher.")
    source = args.directory / "fixture.docx"
    doc.save(source)
    for profile, head in [("student", ""), ("professional", "SYNTHETIC RESPONSE PATTERNS")]:
        out, _ = format_file(source, output=args.directory / (profile + ".docx"), profile=profile, running_head=head)
        print(out)


if __name__ == "__main__":
    main()

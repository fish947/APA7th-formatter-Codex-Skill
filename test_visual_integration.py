"""Formatter/API/CLI integration tests with a real DOCX OPC package.

All generated manuscripts, chart parts, reports, and exports live in temp dirs.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches
from lxml import etree

import apa7_format as formatter
import apa7_visuals as visuals


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
)
SVG = b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><path d="M1 1L9 9" stroke="black"/></svg>'
CHART = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">
 <c:chart><c:plotArea><c:layout/><c:barChart><c:barDir val="col"/><c:grouping val="clustered"/>
  <c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:v>Response</c:v></c:tx>
   <c:cat><c:strLit><c:ptCount val="2"/><c:pt idx="0"><c:v>A</c:v></c:pt><c:pt idx="1"><c:v>B</c:v></c:pt></c:strLit></c:cat>
   <c:val><c:numLit><c:formatCode>General</c:formatCode><c:ptCount val="2"/><c:pt idx="0"><c:v>3</c:v></c:pt><c:pt idx="1"><c:v>7</c:v></c:pt></c:numLit></c:val>
  </c:ser><c:axId val="1"/><c:axId val="2"/></c:barChart>
  <c:catAx><c:axId val="1"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:axPos val="b"/><c:crossAx val="2"/></c:catAx>
  <c:valAx><c:axId val="2"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:axPos val="l"/><c:crossAx val="1"/></c:valAx>
 </c:plotArea><c:plotVisOnly val="1"/></c:chart>
</c:chartSpace>'''


def build_document(path):
    doc = Document()
    doc.add_paragraph("This document contains editable native objects and original image assets.")
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip((table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)),
                           ("Condition", "Response", "A", "3")):
        cell.text = value
    doc.add_paragraph().add_run().add_picture(io.BytesIO(PNG), width=Inches(1))
    chart_part = Part(PackURI("/word/charts/chart1.xml"),
                      "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", CHART, doc.part.package)
    relation_id = doc.part.relate_to(chart_part, RT.CHART)
    drawing = OxmlElement("w:drawing")
    inline = OxmlElement("wp:inline")
    extent = OxmlElement("wp:extent")
    extent.set("cx", str(Inches(4)))
    extent.set("cy", str(Inches(3)))
    inline.append(extent)
    properties = OxmlElement("wp:docPr")
    properties.set("id", "2")
    properties.set("name", "Native Chart")
    inline.append(properties)
    graphic, data = OxmlElement("a:graphic"), OxmlElement("a:graphicData")
    data.set("uri", visuals.NS["c"])
    chart = OxmlElement("c:chart")
    chart.set(qn("r:id"), relation_id)
    data.append(chart)
    graphic.append(data)
    inline.append(graphic)
    drawing.append(inline)
    doc.add_paragraph().add_run()._r.append(drawing)
    vector_part = Part(PackURI("/word/media/diagram.svg"), "image/svg+xml", SVG, doc.part.package)
    doc.part.relate_to(vector_part, RT.IMAGE)
    doc.save(path)
    return path


class VisualIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="apa7_visual_integration_")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = build_document(self.root / "source.docx")
        self.source_bytes = self.source.read_bytes()

    def test_api_exports_and_feedback_preserve_source_native_chart_and_image_bytes(self):
        out, report = formatter.format_file(self.source, output=self.root / "formatted.docx",
                                            profile="student", export_visuals=True)
        self.assertEqual(self.source.read_bytes(), self.source_bytes)
        with ZipFile(out) as package:
            self.assertEqual(package.read("word/charts/chart1.xml"), CHART)
            self.assertEqual(package.read("word/media/diagram.svg"), SVG)
            self.assertEqual(package.read("word/media/image1.png"), PNG)
        output_doc = Document(out)
        self.assertEqual(len(output_doc.tables), 1, "Native table must remain editable, not become an image")
        self.assertEqual(output_doc.tables[0].cell(1, 1).text, "3")
        chart_relations = [relation.target_part.partname for relation in output_doc.part.rels.values() if relation.reltype == RT.CHART]
        self.assertEqual(chart_relations, ["/word/charts/chart1.xml"])
        self.assertEqual(len(output_doc._element.xpath(".//c:chart")), 1)
        self.assertEqual(report["visuals"]["inventory"]["counts"]["native_charts"], 1)
        exports = report["visuals"]["exports"]
        self.assertEqual(exports["source_sha256"], hashlib.sha256(self.source_bytes).hexdigest())
        folder = out.with_suffix(".visuals")
        manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
        by_part = {item["source_part"]: item for item in manifest["items"]}
        for part, content in (("word/media/image1.png", PNG), ("word/media/diagram.svg", SVG)):
            self.assertEqual((folder / by_part[part]["output"]).read_bytes(), content)
        chart_item = by_part["word/charts/chart1.xml"]
        self.assertEqual(chart_item["method"], "reconstructed_from_chart_cache", chart_item)
        self.assertEqual(chart_item["data"]["series"][0]["points"], [["A", 3.0], ["B", 7.0]])
        svg = etree.fromstring((folder / chart_item["output"]).read_bytes())
        self.assertFalse(svg.findall(".//svg:image", visuals.NS))
        self.assertTrue(svg.findall(".//svg:path", visuals.NS) or svg.findall(".//svg:rect", visuals.NS))
        self.assertEqual(json.loads(json.dumps(report["visuals"]["exports"])), manifest)
        self.assertTrue(any("图和图表" in item for item in report["feedback"]["changed"]))
        self.assertFalse(out.with_suffix(".feedback.html").exists())

    def test_cli_inspection_is_read_only_and_cli_export_produces_expected_artifacts(self):
        script = Path(formatter.__file__).resolve()
        before_names = {p.name for p in self.root.iterdir()}
        inspected = subprocess.run([sys.executable, str(script), str(self.source), "--inspect-visuals"],
                                   capture_output=True, text=True, check=False, timeout=30)
        self.assertEqual(inspected.returncode, 0, inspected.stderr)
        inventory = json.loads(inspected.stdout)
        self.assertEqual(inventory["counts"]["native_charts"], 1)
        self.assertEqual(inventory["counts"]["native_tables"], 1)
        self.assertEqual(inventory["counts"]["raster_media"], 1)
        self.assertEqual(inventory["counts"]["vector_media"], 1)
        self.assertEqual({p.name for p in self.root.iterdir()}, before_names)
        self.assertEqual(self.source.read_bytes(), self.source_bytes)
        out = self.root / "cli.docx"
        exported = subprocess.run([sys.executable, str(script), str(self.source), "--profile", "student",
                                   "--output", str(out), "--export-visuals"],
                                  capture_output=True, text=True, check=False, timeout=30)
        self.assertEqual(exported.returncode, 0, exported.stderr)
        self.assertIn(str(out.with_suffix(".visuals")), exported.stdout)
        for path in (out, out.with_suffix(".visuals") / "manifest.json"):
            self.assertTrue(path.is_file(), str(path))
        self.assertFalse(out.with_suffix(".feedback.html").exists())
        self.assertEqual(self.source.read_bytes(), self.source_bytes)

    def test_existing_visual_export_directory_is_rejected_before_any_deliverable_is_written(self):
        out = self.root / "collision.docx"
        destination = out.with_suffix(".visuals")
        destination.mkdir()
        sentinel = destination / "manifest.json"
        sentinel.write_bytes(b"existing export must be preserved")
        with self.assertRaisesRegex(ValueError, "导出文件夹已存在"):
            formatter.format_file(self.source, output=out, profile="student", export_visuals=True)
        self.assertEqual(sentinel.read_bytes(), b"existing export must be preserved")
        self.assertEqual(self.source.read_bytes(), self.source_bytes)
        self.assertEqual(list(destination.iterdir()), [sentinel])
        for path in (out, out.with_suffix(".feedback.html")):
            self.assertFalse(path.exists(), str(path))

    def test_optional_export_failure_is_reported_while_successful_word_output_is_retained(self):
        with patch.object(visuals, "export_visuals", side_effect=OSError("simulated export storage failure")):
            out, report = formatter.format_file(self.source, output=self.root / "partial.docx",
                                                profile="student", export_visuals=True)
        self.assertTrue(out.is_file())
        self.assertEqual(self.source.read_bytes(), self.source_bytes)
        self.assertEqual(report["preservation"], "passed")
        self.assertEqual(report["visuals"]["export_error"], "simulated export storage failure")
        self.assertNotIn("exports", report["visuals"])
        self.assertTrue(any(event["status"] == "review" and "simulated export storage failure" in event["message"] for event in report["events"]))
        self.assertIn("矢量导出没有完成，需要重新处理。", report["feedback"]["needs_review"])
        self.assertFalse(out.with_suffix(".feedback.html").exists())
        self.assertFalse(out.with_suffix(".visuals").exists())


if __name__ == "__main__":
    unittest.main()

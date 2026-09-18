"""Tests for review gating and provenance, not a claim of automated AI judgment."""
import copy
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from lxml import etree

import apa7_format as engine
import apa7_workflow as workflow


class AIWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="apa7_ai_gate_")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "paper.docx"
        doc = Document()
        doc.add_paragraph("Method", "Heading 1")
        doc.add_paragraph("Participants", "Heading 2")
        doc.add_paragraph("The synthetic sample contains forty participants.")
        data = doc.add_table(rows=2, cols=2)
        data.style = "Table Grid"
        for row, values in zip(data.rows, [("Condition", "Mean"), ("A", "3")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        layout = doc.add_table(rows=1, cols=2)
        layout.cell(0, 0).text = "Name"
        layout.cell(0, 1).text = "Alex Example"
        doc.save(self.source)
        self.before = self.source.read_bytes()

    def reviewed(self):
        config = workflow.prepare(self.source, "student")
        config["review_status"] = "reviewed"
        config["roles"] = {"1": "heading1", "2": "heading2", "3": "body"}
        config["table_roles"] = {"1": "data", "2": "preserve"}
        config["table_header_rows"] = {"1": 1}
        config["decision_notes"] = {
            "paragraphs": {
                "1": {"confidence": "high", "reason": "Method is the top-level methods section."},
                "2": {"confidence": "high", "reason": "Participants is nested within Method."},
                "3": {"confidence": "high", "reason": "A complete narrative sentence about the sample."}},
            "tables": {
                "1": {"confidence": "high", "reason": "One header row labels condition and mean values."},
                "2": {"confidence": "uncertain", "reason": "A name field, not clearly a research data table."}}}
        config["compliance_review"] = {
            "target_requirements": {"status": "pass", "reason": "No course or journal deviations were supplied for this synthetic fixture."},
            "title_page": {"status": "needs_review", "reason": "The synthetic fixture has no title page."},
            "abstract_keywords": {"status": "not_applicable", "reason": "This short fixture contains no abstract."},
            "headings": {"status": "pass", "reason": "Method and Participants form a valid two-level hierarchy."},
            "citations_references": {"status": "not_applicable", "reason": "The fixture contains no citations or references."},
            "statistics": {"status": "not_applicable", "reason": "The fixture contains no statistical test expression."},
            "equations": {"status": "not_applicable", "reason": "The fixture contains no native or plain-text equation."},
            "tables": {"status": "needs_review", "reason": "The second table may be a layout object rather than a data table."},
            "figures": {"status": "not_applicable", "reason": "The fixture contains no figures."},
            "appendices": {"status": "not_applicable", "reason": "The fixture contains no appendices."}}
        config["unresolved"] = ["Confirm the intended layout of the name table."]
        return config

    def test_prepare_exposes_cells_order_and_no_review_claim(self):
        config = workflow.prepare(self.source, "student")
        self.assertEqual(config["review_status"], "pending")
        self.assertEqual(config["roles"], {})
        self.assertEqual({item["status"] for item in config["compliance_review"].values()}, {"pending"})
        self.assertEqual(config["_review"]["tables"][0]["cells"], [["Condition", "Mean"], ["A", "3"]])
        self.assertEqual([x["kind"] for x in config["_review"]["body_order"]], ["paragraph"] * 3 + ["table"] * 2)
        self.assertEqual(config["_review"]["object_summary"]["top_level_tables"], 2)
        self.assertEqual(config["_review"]["citation_reference_check"]["status"], "not_applicable")
        self.assertIn("front_matter_check", config["_review"])
        self.assertIn("caption_object_check", config["_review"])
        self.assertEqual(config["_review"]["preflight_summary"]["checks_run"], 7)
        self.assertEqual(self.source.read_bytes(), self.before)

    def test_pending_and_missing_object_classifications_rejected(self):
        with self.assertRaisesRegex(ValueError, "尚未完成"):
            workflow.validate_review(self.source, workflow.prepare(self.source, "student"), "student")
        config = self.reviewed()
        del config["roles"]["3"]
        with self.assertRaisesRegex(ValueError, "覆盖全部"):
            workflow.validate_review(self.source, config, "student")

    def test_uncertain_role_must_preserve_and_reasons_required(self):
        config = self.reviewed()
        config["table_roles"]["2"] = "data"
        with self.assertRaisesRegex(ValueError, "不确定对象"):
            workflow.validate_review(self.source, config, "student")
        config = self.reviewed()
        config["decision_notes"]["paragraphs"]["1"]["reason"] = ""
        with self.assertRaisesRegex(ValueError, "分类依据"):
            workflow.validate_review(self.source, config, "student")

    def test_reviewed_apply_preserves_layout_table_and_source(self):
        before_table = etree.tostring(Document(self.source).tables[1]._tbl, method="c14n")
        out, report = workflow.apply_reviewed(self.source, self.reviewed(), profile="student", output=self.root / "output.docx")
        self.assertEqual(self.source.read_bytes(), self.before)
        self.assertEqual(etree.tostring(Document(out).tables[1]._tbl, method="c14n"), before_table)
        self.assertEqual(report["ai_review"]["status"], "classification_and_compliance_record_validated")
        self.assertEqual(len(report["ai_review"]["unresolved"]), 1)
        self.assertTrue(all(item["status"] == "passed" for item in report["machine_checks"]
                            if item["id"] not in {"page_header", "numbered_object_references",
                                                  "front_matter", "caption_object_pairing"}))
        numbered = next(item for item in report["machine_checks"]
                        if item["id"] == "numbered_object_references")
        self.assertEqual(numbered["status"], "needs_review")
        self.assertTrue(any(issue["code"] == "data_table_without_label"
                            for issue in report["numbered_object_check"]["issues"]))
        self.assertTrue(any(e["location"] == "table2" and "保留此表格" in e["message"] for e in report["events"]))

    def test_compliance_review_must_be_complete_and_match_present_objects(self):
        config = self.reviewed()
        del config["compliance_review"]["headings"]
        with self.assertRaisesRegex(ValueError, "覆盖全部"):
            workflow.validate_review(self.source, config, "student")
        config = self.reviewed()
        config["compliance_review"]["tables"] = {"status": "not_applicable", "reason": "Incorrect test claim."}
        with self.assertRaisesRegex(ValueError, "存在表格"):
            workflow.validate_review(self.source, config, "student")
        config = self.reviewed()
        config["compliance_review"]["headings"] = {"status": "not_applicable", "reason": "Incorrect test claim."}
        with self.assertRaisesRegex(ValueError, "检测到标题层级"):
            workflow.validate_review(self.source, config, "student")
        config = self.reviewed()
        config["compliance_review"]["title_page"] = {"status": "not_applicable", "reason": "Incorrect test claim."}
        with self.assertRaisesRegex(ValueError, "不能标为不适用"):
            workflow.validate_review(self.source, config, "student")

    def test_detected_statistics_and_equations_require_compliance_review(self):
        doc = Document(self.source)
        doc.paragraphs[2].text = "The difference was t(28) = 2.40, p = .023."
        doc.save(self.source)
        config = self.reviewed()
        config["compliance_review"]["statistics"] = {"status": "not_applicable", "reason": "Incorrect test claim."}
        with self.assertRaisesRegex(ValueError, "统计数据审核不能标为不适用"):
            workflow.validate_review(self.source, config, "student")

        doc = Document(self.source)
        doc.add_paragraph("The model was y = mx + b.")
        doc.save(self.source)
        fresh = engine.prepare_config(self.source, "student")
        review = {key: {"status": "pass", "reason": "Synthetic evidence."} for key in workflow.COMPLIANCE_AREAS}
        review["equations"] = {"status": "not_applicable", "reason": "Incorrect test claim."}
        with self.assertRaisesRegex(ValueError, "公式审核不能标为不适用"):
            workflow.validate_compliance_review({"compliance_review": review}, fresh)

    def test_explicit_header_count_and_matching_profile_required(self):
        config = self.reviewed()
        config["table_header_rows"] = {}
        with self.assertRaisesRegex(ValueError, "表头行数"):
            workflow.validate_review(self.source, config, "student")
        with self.assertRaisesRegex(ValueError, "模式"):
            workflow.validate_review(self.source, self.reviewed(), "professional")

    def test_complex_table_cannot_be_forced_through_data_path(self):
        doc = Document(self.source)
        doc.tables[0].cell(0, 0).merge(doc.tables[0].cell(0, 1))
        doc.save(self.source)
        config = self.reviewed()
        with self.assertRaisesRegex(ValueError, "复杂表格"):
            workflow.validate_review(self.source, config, "student")

    def test_cover_range_cannot_override_uncertain_or_conflicting_roles(self):
        config = self.reviewed()
        config["title_page"] = [1, 3]
        with self.assertRaisesRegex(ValueError, "title_page 与已判断角色冲突"):
            workflow.validate_review(self.source, config, "student")

    def qa_fixture(self):
        folder = self.root / "render"
        folder.mkdir()
        pages = []
        for number in (1, 2):
            # Arbitrary deterministic bytes test integrity binding, not image rendering.
            path = folder / f"page-{number}.png"
            path.write_bytes(f"synthetic page bytes {number}".encode())
            pages.append({"page": number, "file": path.name, "sha256": engine.digest(path)})
        current = engine.digest(self.source)
        workflow.write_new_json(folder / "render_manifest.json", {"document_sha256": current, "pages": pages})
        review = {"document_sha256": current, "pages": [
            {"page": i, "status": "pass", "notes": "Synthetic test of a complete review record."} for i in (1, 2)]}
        return folder, review

    def test_qa_requires_all_pages_and_preserves_issue_status(self):
        folder, review = self.qa_fixture()
        self.assertEqual(workflow.record_qa(self.source, folder, review)["status"], "VISUAL_REVIEW_RECORDED")
        review["pages"][1]["status"] = "issue"
        self.assertEqual(workflow.record_qa(self.source, folder, review)["status"], "REVISION_REQUIRED")
        incomplete = copy.deepcopy(review)
        incomplete["pages"].pop()
        with self.assertRaisesRegex(ValueError, "全部页面"):
            workflow.record_qa(self.source, folder, incomplete)

    def test_qa_rejects_changed_images_or_document(self):
        folder, review = self.qa_fixture()
        (folder / "page-1.png").write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "渲染图片"):
            workflow.record_qa(self.source, folder, review)
        review["document_sha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "Word 版本不一致"):
            workflow.record_qa(self.source, folder, review)


if __name__ == "__main__":
    unittest.main()

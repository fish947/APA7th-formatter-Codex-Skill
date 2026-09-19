"""Regression tests for structured results to APA tables and figures."""
from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document

import apa7_results as results


class ResultsGenerationTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="apa7_results_test_")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def spec(self):
        return {
            "schema_version": 1,
            "document_title": "Results Tables and Figures",
            "tables": [{
                "number": 1,
                "title": "Descriptive Statistics by Condition",
                "columns": ["Condition", "n", "M", "SD", "p"],
                "rows": [["Control", "30", "4.10", "0.82", ".032"], ["Treatment", "30", "4.68", "0.77", ".000"]],
                "numeric_columns": [1, 2, 3, 4],
                "note": "M = mean; SD = standard deviation; p = .000 requires author review.",
            }],
            "figures": [{
                "number": 1,
                "title": "Mean Score by Session",
                "type": "line",
                "x_label": "Session",
                "y_label": "Mean score",
                "series": [
                    {"name": "Control", "x": ["1", "2", "3"], "y": ["4.10", "4.20", "4.32"]},
                    {"name": "Treatment", "x": ["1", "2", "3"], "y": ["4.18", "4.45", "4.68"]},
                ],
                "alt_text": "Two lines show mean scores increasing across three sessions.",
                "note": "Error bars were not supplied and were not invented.",
            }],
        }

    def test_builds_editable_table_and_data_derived_figure_without_recomputing_values(self):
        output = self.root / "results.docx"
        svg_dir = self.root / "svg"
        path, report = results.build_results_document(self.spec(), output, export_svg=svg_dir)
        self.assertEqual(path, output)
        self.assertFalse(report["values_recomputed"])
        self.assertEqual(report["tables_created"], 1)
        self.assertEqual(report["figures_created"], 1)
        self.assertTrue(any(item["code"] == "p_zero" for item in report["statistical_warnings"]))
        doc = Document(output)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual([cell.text for cell in table.rows[0].cells], ["Condition", "n", "M", "SD", "p"])
        self.assertEqual([cell.text.strip() for cell in table.rows[1].cells], ["Control", "30", "4.10", "0.82", ".032"])
        self.assertEqual([cell.text.strip() for cell in table.rows[2].cells], ["Treatment", "30", "4.68", "0.77", ".000"])
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        self.assertIn("Table 1", text)
        self.assertIn("Figure 1", text)
        self.assertIn("Error bars were not supplied and were not invented.", text)
        with ZipFile(output) as package:
            self.assertTrue(any(name.endswith(".png") for name in package.namelist()))
            document_xml = package.read("word/document.xml")
            self.assertIn(b"Two lines show mean scores increasing", document_xml)
            self.assertNotIn(b"w:insideV w:val=\"single\"", document_xml)
        svg = (svg_dir / "figure-1.svg").read_text(encoding="utf-8")
        self.assertIn("<path", svg)
        self.assertIn("<circle", svg)
        self.assertNotIn("<image", svg)
        manifest = json.loads((svg_dir / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["items"][0]["method"], "reconstructed_from_supplied_data")

    def test_json_loader_preserves_reported_decimal_precision(self):
        source = self.root / "results.json"
        source.write_text('{"schema_version":1,"tables":[{"title":"T","columns":["M"],"rows":[[0.10]]}],"figures":[]}', encoding="utf-8")
        loaded = results.load_spec(source)
        self.assertEqual(loaded["tables"][0]["rows"][0][0], "0.10")
        output = self.root / "precision.docx"
        results.build_results_document(loaded, output)
        self.assertEqual(Document(output).tables[0].cell(1, 0).text.strip(), "0.10")

    def test_csv_cli_creates_one_table_and_rejects_overwrite(self):
        source = self.root / "values.csv"
        source.write_text("Condition,n,M,SD\nControl,30,4.10,0.82\nTreatment,30,4.68,0.77\n", encoding="utf-8")
        output = self.root / "table.docx"
        command = [sys.executable, str(Path(results.__file__).resolve()), str(source), "--output", str(output), "--table-title", "Descriptive Statistics"]
        first = subprocess.run(command, capture_output=True, text=True, check=False)
        self.assertEqual(first.returncode, 0, first.stderr)
        self.assertEqual(json.loads(first.stdout)["tables_created"], 1)
        second = subprocess.run(command, capture_output=True, text=True, check=False)
        self.assertNotEqual(second.returncode, 0)
        self.assertIn("输出文件已存在", second.stderr)

    def test_invalid_chart_data_is_rejected_before_output(self):
        spec = self.spec()
        spec["figures"][0]["series"][0]["y"] = ["1", "2"]
        output = self.root / "bad.docx"
        with self.assertRaisesRegex(results.ResultsSpecError, "等长"):
            results.build_results_document(spec, output)
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()

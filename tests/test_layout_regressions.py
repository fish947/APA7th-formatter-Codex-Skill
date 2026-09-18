"""Targeted regressions found during the formatter's layout review.

Run with: python -m unittest discover -s tests -p 'test_*.py' -v
All generated documents remain inside temporary directories.
"""
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import apa7_format as apa


class LayoutRegressionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="apa7_layout_tests_")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def format_document(self, document):
        source = self.root / "source.docx"
        document.save(source)
        original = source.read_bytes()
        output, report = apa.format_file(source, output=self.root / "formatted.docx", profile="student")
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(report["preservation"], "passed")
        return Document(output), report

    def test_title_overrides_inherited_blue_border_with_direct_nil(self):
        document = Document()
        style_properties = document.styles["Title"].element.get_or_add_pPr()
        for existing in list(style_properties.findall(qn("w:pBdr"))):
            style_properties.remove(existing)
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for attribute, value in {"val": "single", "sz": "24", "color": "0070C0"}.items():
            bottom.set(qn("w:" + attribute), value)
        borders.append(bottom)
        style_properties.append(borders)
        document.add_paragraph("An Example Research Title", style="Title")

        result, report = self.format_document(document)
        title = result.paragraphs[0]
        self.assertEqual(report["paragraph_roles"]["1"], "title")
        self.assertEqual(title.text, "An Example Research Title")
        self.assertEqual(title.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(title.runs[0].bold)
        direct_borders = title._p.xpath("./w:pPr/w:pBdr")
        self.assertEqual(len(direct_borders), 1)
        for edge in ("top", "left", "bottom", "right", "between", "bar"):
            with self.subTest(edge=edge):
                node = direct_borders[0].find(qn("w:" + edge))
                self.assertIsNotNone(node)
                self.assertEqual(node.get(qn("w:val")), "nil")
        # The paragraph override must work without rewriting the shared Title style.
        inherited_bottom = result.styles["Title"].element.xpath("./w:pPr/w:pBdr/w:bottom")[0]
        self.assertEqual(inherited_bottom.get(qn("w:val")), "single")
        self.assertEqual(inherited_bottom.get(qn("w:color")), "0070C0")

    def test_table_callout_sentence_is_formatted_as_body_before_separate_caption(self):
        document = Document()
        sentence = document.add_paragraph("Table 1 summarizes the results.")
        sentence.paragraph_format.first_line_indent = Inches(0)
        sentence.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sentence.runs[0].font.name = "Arial"
        sentence.runs[0].font.size = Pt(20)
        document.add_paragraph("Table 1")
        document.add_paragraph("Summary of Results")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Measure"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Accuracy"
        table.cell(1, 1).text = ".90"

        result, report = self.format_document(document)
        body = result.paragraphs[0]
        self.assertEqual(report["paragraph_roles"]["1"], "body")
        self.assertEqual(report["paragraph_roles"]["2"], "caption_number")
        self.assertEqual(report["paragraph_roles"]["3"], "caption_title")
        self.assertEqual(body.text, "Table 1 summarizes the results.")
        self.assertEqual(body.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(body.paragraph_format.first_line_indent, Inches(0.5))
        self.assertEqual(body.paragraph_format.line_spacing, 2.0)
        self.assertEqual(body.runs[0].font.name, "Times New Roman")
        self.assertEqual(body.runs[0].font.size, Pt(12))

    def test_note_prefix_split_italicizes_only_prefix_and_preserves_statistics(self):
        document = Document()
        document.add_paragraph("Table 1")
        document.add_paragraph("Statistical Results")
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Statistic"
        table.cell(1, 0).text = ".03"
        note = document.add_paragraph()
        plain = note.add_run("Note. The test reports ")
        plain.italic = False
        plain.bold = False
        note.add_run("p").italic = True
        note.add_run(" = .03 with ")
        note.add_run("strong evidence").bold = True
        note.add_run(".")
        original_text = note.text

        result, report = self.format_document(document)
        formatted = result.paragraphs[-1]
        self.assertEqual(report["paragraph_roles"]["3"], "note")
        self.assertEqual(formatted.text, original_text)
        self.assertEqual(formatted.paragraph_format.first_line_indent, Inches(0))
        self.assertEqual(formatted.runs[0].text, "Note.")
        self.assertTrue(formatted.runs[0].italic)
        self.assertFalse(formatted.runs[0].bold)
        self.assertEqual(formatted.runs[1].text, " The test reports ")
        self.assertFalse(formatted.runs[1].italic)
        self.assertFalse(formatted.runs[1].bold)
        by_text = {run.text: run for run in formatted.runs}
        self.assertTrue(by_text["p"].italic)
        self.assertTrue(by_text["strong evidence"].bold)
        self.assertIsNone(by_text[" = .03 with "].italic)
        self.assertIsNone(by_text["strong evidence"].italic)


if __name__ == "__main__":
    unittest.main()

"""Regression tests for safe APA statistical and equation presentation."""
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

import apa7_format as engine
import apa7_statistics as statistics


def add_native_display_equation(paragraph, text="x=1", number=None):
    math_paragraph = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    content = OxmlElement("m:t")
    content.text = text
    run.append(content)
    math.append(run)
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)
    if number is not None:
        paragraph.add_run(f" ({number})")


class StatisticalReportingTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="apa7_statistics_")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def test_normalization_changes_presentation_but_not_numeric_values(self):
        before = "The result was t(28)=.75, P=0.032, M=.55, and SD=.20."
        after = statistics.normalize_statistical_text(before)
        self.assertEqual(after, "The result was t(28) = 0.75, p = .032, M = 0.55, and SD = 0.20.")
        self.assertEqual(statistics._decimal_values(before), statistics._decimal_values(after))

    def test_statistical_symbols_are_styled_and_review_items_are_reported(self):
        doc = Document()
        paragraph = doc.add_paragraph("The result was t(28)=.75, P=0.032, M=.55, and SD=.20.")
        report = statistics.format_statistics_and_equations(doc, {0: "body"})
        self.assertEqual(paragraph.text, "The result was t(28) = 0.75, p = .032, M = 0.55, and SD = 0.20.")
        by_text = {run.text: run for run in paragraph.runs if run.text in {"t", "p", "M", "SD"}}
        self.assertEqual(set(by_text), {"t", "p", "M", "SD"})
        self.assertTrue(all(run.italic for run in by_text.values()))
        self.assertFalse(report["data_values_changed"])
        self.assertEqual(report["tests_found"], 1)
        self.assertTrue(any(issue["code"] == "test_missing_effect" for issue in report["issues"]))
        self.assertTrue(any(issue["code"] == "test_missing_ci" for issue in report["issues"]))

    def test_p_zero_and_invalid_precision_are_never_rewritten_as_new_results(self):
        doc = Document()
        paragraph = doc.add_paragraph("The results were p=.000 and p=.03245.")
        report = statistics.format_statistics_and_equations(doc, {0: "body"})
        self.assertEqual(paragraph.text, "The results were p = .000 and p = .03245.")
        codes = {issue["code"] for issue in report["issues"]}
        self.assertIn("p_zero", codes)
        self.assertIn("p_precision", codes)
        self.assertNotIn("p < .001", paragraph.text)

    def test_greek_statistical_symbol_is_upright(self):
        doc = Document()
        run = doc.add_paragraph().add_run("β=0.25 and ηp2=0.08")
        run.italic = True
        report = statistics.format_statistics_and_equations(doc, {0: "body"})
        paragraph = doc.paragraphs[0]
        self.assertEqual(paragraph.text, "β = .25 and ηp2 = .08")
        beta = next(run for run in paragraph.runs if run.text == "β")
        self.assertFalse(beta.italic)
        self.assertGreaterEqual(report["symbols_formatted"], 1)
        self.assertTrue(any(issue["code"] == "stat_superscript" for issue in report["issues"]))

    def test_plain_formula_detection_avoids_ordinary_assignment_prose(self):
        doc = Document()
        doc.add_paragraph("The model was y = mx + b.")
        doc.add_paragraph("Condition A = 5 participants in the control group.")
        report = statistics.format_statistics_and_equations(doc, {0: "body", 1: "body"})
        self.assertEqual(report["equations"]["plain_text_formula_candidates"], 1)
        self.assertEqual(report["equations"]["locations"], [])

    def test_statistical_formatting_is_idempotent(self):
        doc = Document()
        paragraph = doc.add_paragraph("The result was F(1, 28)=.75, P=0.032, R2=0.12.")
        first = statistics.format_statistics_and_equations(doc, {0: "body"})
        first_text = paragraph.text
        first_values = statistics._decimal_values(first_text)
        second = statistics.format_statistics_and_equations(doc, {0: "body"})
        self.assertEqual(paragraph.text, first_text)
        self.assertEqual(statistics._decimal_values(paragraph.text), first_values)
        self.assertEqual(second["safe_text_edits"], [])
        self.assertFalse(first["data_values_changed"] or second["data_values_changed"])

    def test_preserved_table_is_inspected_without_being_changed(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "p=0.010"
        report = statistics.format_statistics_and_equations(doc, {}, {"1": "preserve"})
        self.assertEqual(table.cell(0, 0).text, "p=0.010")
        self.assertEqual(report["expressions_found"], 1)
        self.assertEqual(report["safe_text_edits"], [])

    def test_native_equation_is_preserved_and_numbering_is_checked(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        add_native_display_equation(paragraph, "x=1")
        before_math = etree.tostring(paragraph._p.xpath(".//m:oMath")[0], method="c14n")
        report = statistics.format_statistics_and_equations(doc, {0: "equation"})
        self.assertEqual(etree.tostring(paragraph._p.xpath(".//m:oMath")[0], method="c14n"), before_math)
        self.assertEqual(report["equations"]["displayed_equations"], 1)
        self.assertEqual(report["equations"]["paragraphs_formatted"], 1)
        self.assertTrue(any(issue["code"] == "equation_number" for issue in report["issues"]))
        self.assertEqual(paragraph.paragraph_format.line_spacing, 2.0)

    def test_end_to_end_allows_only_recorded_statistical_text_edits(self):
        doc = Document()
        doc.add_paragraph("Results", "Heading 1")
        doc.add_paragraph("The difference was t(28)=2.41, P=0.023, d=.45, 95% CI [0.10, 0.80].")
        source = self.root / "source.docx"
        output = self.root / "output.docx"
        doc.save(source)
        original = source.read_bytes()
        formatted, report = engine.format_file(source, output=output, profile="student")
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(
            Document(formatted).paragraphs[1].text,
            "The difference was t(28) = 2.41, p = .023, d = 0.45, 95% CI [0.10, 0.80].",
        )
        self.assertEqual(report["version"], "0.14.0")
        self.assertFalse(report["statistical_reporting"]["data_values_changed"])
        self.assertTrue(any(source["title"] == "Numbers and Statistics Guide" for source in report["feedback"]["apa_sources"]))


if __name__ == "__main__":
    unittest.main()

"""Structure-review safeguards for document-bound formatting decisions."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

import apa7_format as apa


class StructureReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="apa7_structure_")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        self.doc = Document()
        self.doc.add_paragraph("A Synthetic Paper", "Title")
        self.doc.add_paragraph("Alex Example")
        p = self.doc.add_paragraph("A Synthetic Paper", "Title")
        p.paragraph_format.page_break_before = True
        self.doc.add_paragraph("Results", "Heading 1")
        self.doc.add_paragraph("The expression `y = mx + b` keeps its introduction and explanation.")
        self.doc.save(self.source)
        self.original = self.source.read_bytes()

    def test_draft_is_read_only_and_inferred_roles_are_not_confirmed(self):
        draft = apa.prepare_config(self.source)
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertEqual(draft["source_sha256"], apa.digest(self.source))
        self.assertEqual(draft["roles"], {})
        self.assertEqual(draft["_review"]["suggested_title_page"], [1, 2])
        self.assertEqual(draft["_review"]["paragraphs"][3]["inferred_role"], "heading1")
        self.assertEqual(draft["_review"]["paragraphs"][4]["text"], self.doc.paragraphs[4].text)
        self.assertEqual(list(self.root.iterdir()), [self.source])

    def test_prepare_handles_raw_ooxml_containers_from_third_party_generators(self):
        doc = Document()
        paragraph = doc.add_paragraph("A caption-like paragraph")
        container = etree.Element(qn("w:sdt"))
        content = etree.SubElement(container, qn("w:sdtContent"))
        nested_paragraph = etree.SubElement(content, qn("w:p"))
        run = etree.SubElement(nested_paragraph, qn("w:r"))
        drawing = etree.SubElement(run, qn("w:drawing"))
        etree.SubElement(drawing, qn("wp:inline"))
        paragraph._p.addnext(container)
        source = self.root / "third-party.docx"
        doc.save(source)

        draft = apa.prepare_config(source)

        self.assertEqual(draft["_review"]["object_summary"]["inline_drawings"], 1)
        unsupported = [item for item in draft["_review"]["body_order"]
                       if item["kind"] == "unsupported_container"]
        self.assertEqual(unsupported, [{"kind": "unsupported_container", "tag": "sdt", "text": ""}])

    def test_config_hash_rejects_changed_original_before_creating_outputs(self):
        config = apa.prepare_config(self.source)
        config["roles"] = {"4": "heading2"}
        self.doc.add_paragraph("New content changes the binding.")
        self.doc.save(self.source)
        output = self.root / "blocked.docx"
        with self.assertRaisesRegex(ValueError, "配置不属于当前版本"):
            apa.format_file(self.source, output=output, profile="student", config=config, export_visuals=True)
        self.assertEqual(list(self.root.iterdir()), [self.source])

    def test_matching_config_applies_confirmed_role_and_preserves_formula_sentence(self):
        config = apa.prepare_config(self.source)
        config["roles"] = {"4": "heading2"}
        out, report = apa.format_file(self.source, output=self.root / "formatted.docx", profile="student", config=config)
        self.assertTrue(report["configuration_source_bound"])
        self.assertEqual(report["paragraph_roles"]["4"], "heading2")
        self.assertEqual(Document(out).paragraphs[4].text, self.doc.paragraphs[4].text)
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertEqual({x["status"] for x in report["review_checklist"]}, {"pending_review"})
        self.assertEqual({path.name for path in self.root.iterdir()}, {"source.docx", "formatted.docx"})
        self.assertIn("APA 来源", apa.feedback_text(report["feedback"]))

    def test_cli_creates_new_config_only_and_refuses_existing_destination(self):
        destination = self.root / "structure.json"
        command = [sys.executable, str(Path(apa.__file__).resolve()), str(self.source),
                   "--profile", "professional", "--prepare-config", str(destination)]
        first = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertEqual(first.returncode, 0, first.stderr)
        config_bytes = destination.read_bytes()
        draft = json.loads(config_bytes)
        self.assertEqual(draft["_review"]["profile"], "professional")
        fields = next(x["requirements"] for x in draft["_review"]["checklist"] if x["id"] == "title_metadata")
        self.assertIn("running head", fields)
        self.assertNotIn("教师", fields)
        second = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertEqual(second.returncode, 2)
        self.assertEqual(destination.read_bytes(), config_bytes)
        self.assertEqual({x.name for x in self.root.iterdir()}, {"source.docx", "structure.json"})
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_malformed_hash_rejected_and_student_checklist_has_student_fields(self):
        with self.assertRaisesRegex(ValueError, "64 位 SHA256"):
            apa.format_file(self.source, profile="student", config={"source_sha256": "not-a-hash"})
        fields = next(x["requirements"] for x in apa.review_checklist("student") if x["id"] == "title_metadata")
        self.assertIn("教师", fields)
        self.assertNotIn("running head", fields)
        self.assertEqual(list(self.root.iterdir()), [self.source])


if __name__ == "__main__":
    unittest.main()

#!/usr/bin/env python3
"""Conservative APA statistical-reporting and equation formatter.

The module changes presentation only. It never rounds, recalculates, invents,
or removes a reported numeric value, and it never rewrites native Word math.
"""
from __future__ import annotations

import copy
from decimal import Decimal, InvalidOperation
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.run import Run


_NUMBER = r"[−-]?(?:\d+(?:\.\d+)?|\.\d+)"
_OPERATOR = r"(?:<=|>=|=|<|>|≤|≥)"
P_VALUE = re.compile(
    rf"(?<![\w])(?P<symbol>[pP])\s*(?P<operator>{_OPERATOR})\s*(?P<value>{_NUMBER})(?!\d)"
)
TEST_STATISTIC = re.compile(
    rf"(?<![\w])(?P<symbol>χ²|χ2|t|F|z|Z|r|U|W|H)\s*(?P<degrees>\([^()\n]{{1,50}}\))?\s*"
    rf"(?P<operator>{_OPERATOR})\s*(?P<value>{_NUMBER})(?!\d)"
)
SUMMARY_STATISTIC = re.compile(
    rf"(?<![\w])(?P<symbol>Mdn|IQR|SD|SE|df|OR|RR|R²|R2|M|N|n|d|b|B)\s*"
    rf"(?P<operator>{_OPERATOR})\s*(?P<value>{_NUMBER})(?!\d)"
)
GREEK_STATISTIC = re.compile(
    rf"(?<![\w])(?P<symbol>α|β|χ²|χ2|ηp²|ηp2|η²|η2|ω²|ω2)\s*(?P<operator>{_OPERATOR})\s*"
    rf"(?P<value>{_NUMBER})(?!\d)"
)
STAT_SYMBOL_CONTEXT = re.compile(
    r"(?<![\w])(?P<symbol>Mdn|IQR|SD|SE|df|OR|RR|R²|R2|M|N|n|p|t|F|z|Z|r|U|W|H|d|b|B)"
    r"(?=\s*(?:\([^()\n]{1,50}\))?\s*(?:<=|>=|=|<|>|≤|≥))"
)
GREEK_SYMBOL_CONTEXT = re.compile(
    r"(?<![\w])(?P<symbol>α|β|χ²|χ2|ηp²|ηp2|η²|η2|ω²|ω2)"
    r"(?=\s*(?:<=|>=|=|<|>|≤|≥))"
)
PLAIN_FORMULA = re.compile(
    r"(?<![\w])(?P<lhs>[A-Za-z](?:_[A-Za-z0-9]+)?)\s*=\s*"
    r"(?P<rhs>(?=[^\n]{1,100}(?:[+−\-*/^√]|\d))[^.;\n]{1,100})"
)
EFFECT_SIZE = re.compile(r"(?<![\w])(?:d|r|R²|R2|η²|η2|ηp²|ηp2|ω²|ω2|OR|RR)\s*=", re.I)
CONFIDENCE_INTERVAL = re.compile(r"\b(?:9[05]%\s*)?CI\b", re.I)
NUMERIC_TOKEN = re.compile(_NUMBER)

_SKIP_EDIT_ROLES = {
    "title", "title_meta", "reference", "caption_number", "appendix", "appendix_title", "preserve"
}
_LEADING_ZERO_SYMBOLS = {
    "M", "Mdn", "IQR", "SD", "SE", "t", "F", "z", "Z", "U", "W", "H", "d", "b", "B", "OR", "RR", "χ²", "χ2"
}
_NO_LEADING_ZERO_SYMBOLS = {"p", "r", "R²", "R2", "α", "β", "η²", "η2", "ηp²", "ηp2", "ω²", "ω2"}
_MATH_FUNCTIONS = {"abs", "cos", "exp", "ln", "log", "max", "min", "sin", "sqrt", "tan"}


def _decimal_values(text):
    values = []
    for token in NUMERIC_TOKEN.findall(text):
        try:
            values.append(Decimal(token.replace("−", "-").replace("-.", "-0.").replace(".", "0.", 1)
                                  if token.startswith(".") else token.replace("−", "-")))
        except InvalidOperation:
            values.append(token)
    return values


def _normal_value(symbol, value):
    if symbol in _NO_LEADING_ZERO_SYMBOLS and value.startswith("0."):
        return value[1:]
    if symbol in _NO_LEADING_ZERO_SYMBOLS and value.startswith("-0."):
        return "-." + value[3:]
    if symbol in _NO_LEADING_ZERO_SYMBOLS and value.startswith("−0."):
        return "−." + value[3:]
    if symbol in _LEADING_ZERO_SYMBOLS and value.startswith("."):
        return "0" + value
    if symbol in _LEADING_ZERO_SYMBOLS and value.startswith("-."):
        return "-0." + value[2:]
    if symbol in _LEADING_ZERO_SYMBOLS and value.startswith("−."):
        return "−0." + value[2:]
    return value


def _normalize_p(match):
    return f"p {match.group('operator')} {_normal_value('p', match.group('value'))}"


def _normalize_test(match):
    symbol = match.group("symbol")
    degrees = (match.group("degrees") or "").strip()
    value = _normal_value(symbol, match.group("value"))
    return f"{symbol}{degrees} {match.group('operator')} {value}"


def _normalize_summary(match):
    symbol = match.group("symbol")
    value = _normal_value(symbol, match.group("value"))
    return f"{symbol} {match.group('operator')} {value}"


def _normalize_greek(match):
    symbol = match.group("symbol")
    value = _normal_value(symbol, match.group("value"))
    return f"{symbol} {match.group('operator')} {value}"


def normalize_statistical_text(text):
    """Return an APA-presentational normalization without changing numeric values."""
    changed = P_VALUE.sub(_normalize_p, text)
    changed = TEST_STATISTIC.sub(_normalize_test, changed)
    changed = SUMMARY_STATISTIC.sub(_normalize_summary, changed)
    changed = GREEK_STATISTIC.sub(_normalize_greek, changed)
    if _decimal_values(text) != _decimal_values(changed):
        raise RuntimeError("统计格式处理试图改变数值，已停止。")
    return changed


def _iter_tables(table, prefix):
    seen_cells = set()
    for row_number, row in enumerate(table.rows, 1):
        for column_number, cell in enumerate(row.cells, 1):
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            cell_prefix = f"{prefix}.r{row_number}.c{column_number}"
            for paragraph_number, paragraph in enumerate(cell.paragraphs, 1):
                yield paragraph, f"{cell_prefix}.p{paragraph_number}"
            for nested_number, nested in enumerate(cell.tables, 1):
                yield from _iter_tables(nested, f"{cell_prefix}.table{nested_number}")


def iter_document_paragraphs(doc, roles=None, table_roles=None):
    roles, table_roles = roles or {}, table_roles or {}
    for index, paragraph in enumerate(doc.paragraphs):
        yield paragraph, f"p{index + 1}", roles.get(index), roles.get(index) not in _SKIP_EDIT_ROLES
    for table_number, table in enumerate(doc.tables, 1):
        editable = table_roles.get(str(table_number)) != "preserve"
        for paragraph, location in _iter_tables(table, f"table{table_number}"):
            yield paragraph, location, None, editable


def _simple_run(run):
    children = list(run._r)
    return bool(run.text) and all(child.tag in {qn("w:rPr"), qn("w:t")} for child in children)


def _new_text_run(source, text, italic):
    node = OxmlElement("w:r")
    properties = source._r.find(qn("w:rPr"))
    if properties is not None:
        node.append(copy.deepcopy(properties))
    content = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        content.set(qn("xml:space"), "preserve")
    content.text = text
    node.append(content)
    result = Run(node, source._parent)
    result.italic = italic
    return node


def _format_symbols_in_run(run):
    if not _simple_run(run):
        return 0
    targets = [(match.start("symbol"), match.end("symbol"), True) for match in STAT_SYMBOL_CONTEXT.finditer(run.text)]
    targets.extend((match.start("symbol"), match.end("symbol"), False) for match in GREEK_SYMBOL_CONTEXT.finditer(run.text))
    targets.sort()
    if not targets:
        return 0
    filtered, last = [], -1
    for start, end, italic in targets:
        if start >= last:
            filtered.append((start, end, italic))
            last = end
    parent, position = run._r.getparent(), run._r.getparent().index(run._r)
    cursor, changed = 0, 0
    for start, end, italic in filtered:
        if cursor < start:
            parent.insert(position, _new_text_run(run, run.text[cursor:start], run.italic))
            position += 1
        original_italic = run.italic
        parent.insert(position, _new_text_run(run, run.text[start:end], italic))
        position += 1
        changed += int(original_italic is not italic)
        cursor = end
    if cursor < len(run.text):
        parent.insert(position, _new_text_run(run, run.text[cursor:], run.italic))
    parent.remove(run._r)
    return changed


def _paragraph_text(paragraph):
    return "".join(paragraph._p.xpath(".//w:t/text() | .//m:t/text()"))


def _issue(issues, code, location, message, excerpt):
    key = (code, location, message)
    if not any((item["code"], item["location"], item["message"]) == key for item in issues):
        issues.append({"code": code, "location": location, "message": message, "excerpt": excerpt[:180]})


def _inspect_p_values(text, location, issues):
    matches = list(P_VALUE.finditer(text))
    for match in matches:
        value_text = match.group("value").replace("−", "-")
        try:
            value = Decimal(value_text.replace(".", "0.", 1) if value_text.startswith(".") else value_text)
        except InvalidOperation:
            continue
        operator = match.group("operator")
        decimals = len(value_text.partition(".")[2]) if "." in value_text else 0
        if value < 0 or value > 1:
            _issue(issues, "p_range", location, "p 值应在 0 到 1 之间；请核对原始分析输出。", match.group(0))
        if operator == "=" and value == 0:
            _issue(issues, "p_zero", location, "不要把 p 值报告为精确的 .000；若原始结果小于 .001，应报告为 p < .001。", match.group(0))
        if operator == "=" and decimals not in {2, 3}:
            _issue(issues, "p_precision", location, "精确 p 值通常报告到两位或三位小数；请根据原始结果决定如何舍入。", match.group(0))
        if operator in {"<", "<=", "≤"} and value != Decimal("0.001"):
            _issue(issues, "p_exact", location, "除小于 .001 等合理例外外，通常应报告精确 p 值；请核对原始输出。", match.group(0))
        if operator in {">", ">=", "≥"}:
            _issue(issues, "p_operator", location, "此 p 值使用了大于号；请确认这是预期的报告方式。", match.group(0))
    return len(matches)


def _inspect_test_reporting(text, location, issues):
    matches = list(TEST_STATISTIC.finditer(text))
    if not matches:
        return 0
    if not P_VALUE.search(text):
        _issue(issues, "test_missing_p", location, "识别到检验统计量，但未在同段找到 p 值；请确认报告是否完整。", text)
    if not EFFECT_SIZE.search(text):
        _issue(issues, "test_missing_effect", location, "未在同段识别到效应量；请按分析方法和学校或期刊要求确认是否需要补充。", text)
    if not CONFIDENCE_INTERVAL.search(text):
        _issue(issues, "test_missing_ci", location, "未在同段识别到置信区间；请按分析方法和学校或期刊要求确认是否需要补充。", text)
    for match in matches:
        if match.group("symbol") in {"t", "F", "χ²", "χ2"} and not match.group("degrees"):
            _issue(issues, "test_missing_df", location, "检验统计量后未识别到自由度；请核对报告是否完整。", match.group(0))
        if match.group("symbol") == "χ2":
            _issue(issues, "chi_superscript", location, "卡方符号中的 2 应使用上标。", match.group(0))
    return len(matches)


def _formula_like(text):
    matches = []
    occupied = [(match.start(), match.end()) for pattern in (P_VALUE, TEST_STATISTIC, SUMMARY_STATISTIC, GREEK_STATISTIC)
                for match in pattern.finditer(text)]
    for match in PLAIN_FORMULA.finditer(text):
        if any(match.start() < end and match.end() > start for start, end in occupied):
            continue
        # Avoid noisy matches such as "A = 5 participants". Plain-text
        # formulas may contain variables and familiar function names, but not
        # ordinary multi-letter prose words.
        words = re.findall(r"[A-Za-z]+", match.group("rhs"))
        if any(len(word) > 3 and word.lower() not in _MATH_FUNCTIONS for word in words):
            continue
        matches.append(match.group(0).strip())
    return matches


def _inspect_plain_script_notation(text, location, issues):
    for match in re.finditer(r"(?<![\w])(?:R2|η2|ηp2|ω2)(?=\s*(?:<=|>=|=|<|>|≤|≥))", text):
        _issue(issues, "stat_superscript", location,
               "统计符号中的 2 应使用上标；已保留原文，避免改变引用或公式结构。", match.group(0))


def format_statistics_and_equations(doc, roles=None, table_roles=None, profile="student", apply_safe_changes=True):
    """Format safe statistical presentation and return a conservative review report."""
    edits, issues, changed_locations = [], [], set()
    expressions = tests = symbols_formatted = formula_layouts = 0
    formula_candidates = []
    equation_locations, display_numbers = [], []

    for paragraph, location, role, editable in iter_document_paragraphs(doc, roles, table_roles):
        before = _paragraph_text(paragraph)
        if not before:
            continue
        if apply_safe_changes and editable and not paragraph._p.xpath(".//w:instrText | .//w:fldChar | .//w:fldSimple"):
            for text_node in paragraph._p.xpath(".//w:t[not(ancestor::w:txbxContent)]"):
                original = text_node.text or ""
                normalized = normalize_statistical_text(original)
                if normalized != original:
                    text_node.text = normalized
                    if normalized[:1].isspace() or normalized[-1:].isspace():
                        text_node.set(qn("xml:space"), "preserve")
                    edits.append({"location": location, "before": original, "after": normalized,
                                  "kind": "statistical_presentation"})
                    changed_locations.add(location)
            for run_node in list(paragraph._p.xpath(".//w:r[not(ancestor::m:oMath) and not(ancestor::w:txbxContent)]")):
                symbols_formatted += _format_symbols_in_run(Run(run_node, paragraph))
        text = _paragraph_text(paragraph)
        p_count = _inspect_p_values(text, location, issues)
        test_count = _inspect_test_reporting(text, location, issues)
        summary_count = sum(len(pattern.findall(text)) for pattern in (SUMMARY_STATISTIC, GREEK_STATISTIC))
        _inspect_plain_script_notation(text, location, issues)
        expressions += p_count + test_count + summary_count
        tests += test_count
        math_count = len(paragraph._p.xpath(".//m:oMath"))
        if not math_count:
            for formula in _formula_like(text):
                formula_candidates.append({"location": location, "text": formula[:180]})
        if math_count:
            displayed = bool(paragraph._p.xpath(".//m:oMathPara"))
            equation_locations.append({"location": location, "objects": math_count, "displayed": displayed})
            if displayed:
                number_match = re.search(r"\((\d+)\)\s*[.,;:]?\s*$", text)
                if number_match:
                    display_numbers.append((int(number_match.group(1)), location))
                else:
                    _issue(issues, "equation_number", location,
                           "识别到独立公式，但未确认到右侧括号编号；请结合正文引用和目标要求核对。", text)
                _issue(issues, "equation_punctuation", location,
                       "请确认独立公式的标点符合它在句子中的语法位置，并逐页核对编号位置。", text)
                if apply_safe_changes and editable and role == "equation" and location.startswith("p"):
                    pf = paragraph.paragraph_format
                    pf.left_indent = pf.right_indent = Inches(0)
                    pf.first_line_indent = Inches(0)
                    pf.space_before = pf.space_after = Pt(0)
                    pf.line_spacing = 2.0
                    changed_locations.add(location)
                    formula_layouts += 1

    if display_numbers:
        numbers = [number for number, _ in display_numbers]
        if len(numbers) != len(set(numbers)) or numbers != sorted(numbers):
            _issue(issues, "equation_sequence", ", ".join(location for _, location in display_numbers),
                   "公式编号存在重复或顺序异常；不会自动重编号，以免破坏正文引用。", ", ".join(map(str, numbers)))

    for item in formula_candidates:
        _issue(issues, "text_formula", item["location"],
               "识别到纯文本公式；已保留内容，请确认变量斜体、运算符、上下标、括号和标点。", item["text"])

    detected = expressions + len(equation_locations) + len(formula_candidates)
    status = "not_applicable" if not detected else "needs_review" if issues else "passed"
    return {
        "status": status,
        "profile": profile,
        "expressions_found": expressions,
        "tests_found": tests,
        "safe_text_edits": edits,
        "symbols_formatted": symbols_formatted,
        "changed_locations": sorted(changed_locations),
        "issues": issues,
        "equations": {
            "native_math_paragraphs": len(equation_locations),
            "displayed_equations": sum(item["displayed"] for item in equation_locations),
            "numbered_displayed_equations": len(display_numbers),
            "plain_text_formula_candidates": len(formula_candidates),
            "paragraphs_formatted": formula_layouts,
            "locations": equation_locations,
        },
        "data_values_changed": False,
        "limitation": (
            "Pattern-based formatting and review only. It does not recompute p values, verify analysis choices, "
            "round results, infer missing statistics, or rewrite native Word equations."
        ),
    }

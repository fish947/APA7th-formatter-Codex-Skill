"""Create APA-style Word tables and figures from structured results.

This module formats supplied values; it does not run statistical tests, infer
missing results, or change the meaning or precision of reported values.
"""
from __future__ import annotations

import argparse
import csv
import html
import json
import math
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


SCHEMA_VERSION = 1
SUPPORTED_CHARTS = {"bar", "line", "scatter"}
NUMBER_RE = re.compile(r"^\s*[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\s*$")
P_RE = re.compile(r"(?i)(?<![A-Za-z])p\s*([=<>≤≥])\s*(\.?\d+(?:\.\d+)?)")


class ResultsSpecError(ValueError):
    """Raised when the supplied result specification is unsafe or incomplete."""


def _text(value) -> str:
    if value is None:
        return ""
    return str(value)


def _integer(value, label: str, minimum: int = 1) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise ResultsSpecError(f"{label} 必须是整数。") from exc
    if parsed < minimum:
        raise ResultsSpecError(f"{label} 必须不小于 {minimum}。")
    return parsed


def _float(value, label: str) -> float:
    try:
        parsed = float(value)
    except (TypeError, ValueError) as exc:
        raise ResultsSpecError(f"{label} 必须是有限数字。") from exc
    if not math.isfinite(parsed):
        raise ResultsSpecError(f"{label} 必须是有限数字。")
    return parsed


def load_spec(path: Path) -> dict:
    """Load JSON while preserving the lexical precision of table numbers."""
    path = Path(path)
    try:
        value = json.loads(path.read_text(encoding="utf-8"), parse_float=str, parse_int=str)
    except (OSError, json.JSONDecodeError) as exc:
        raise ResultsSpecError(f"无法读取结果 JSON：{exc}") from exc
    if not isinstance(value, dict):
        raise ResultsSpecError("结果 JSON 的最外层必须是对象。")
    return value


def csv_table_spec(path: Path, *, title: str, number: int = 1, note: str = "") -> dict:
    path = Path(path)
    try:
        with path.open(newline="", encoding="utf-8-sig") as source:
            rows = list(csv.reader(source))
    except OSError as exc:
        raise ResultsSpecError(f"无法读取 CSV：{exc}") from exc
    if len(rows) < 2 or not rows[0]:
        raise ResultsSpecError("CSV 至少需要一行列名和一行数据。")
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise ResultsSpecError("CSV 每一行必须有相同的列数。")
    return {
        "schema_version": SCHEMA_VERSION,
        "tables": [{"number": number, "title": title, "columns": rows[0], "rows": rows[1:], "note": note}],
        "figures": [],
    }


def _validate_table(raw: dict, index: int) -> dict:
    if not isinstance(raw, dict):
        raise ResultsSpecError(f"tables[{index}] 必须是对象。")
    columns = raw.get("columns")
    rows = raw.get("rows")
    if not isinstance(columns, list) or not columns or not all(_text(x).strip() for x in columns):
        raise ResultsSpecError(f"tables[{index}].columns 必须是非空列名列表。")
    if not isinstance(rows, list) or not rows:
        raise ResultsSpecError(f"tables[{index}].rows 必须至少包含一行。")
    width = len(columns)
    if any(not isinstance(row, list) or len(row) != width for row in rows):
        raise ResultsSpecError(f"tables[{index}] 的每行必须正好包含 {width} 个值。")
    title = _text(raw.get("title")).strip()
    if not title:
        raise ResultsSpecError(f"tables[{index}].title 不能为空。")
    number = _integer(raw.get("number", index + 1), f"tables[{index}].number")
    numeric = raw.get("numeric_columns")
    if numeric is None:
        numeric = []
        for col in range(width):
            values = [_text(row[col]) for row in rows if _text(row[col]).strip()]
            if values and all(NUMBER_RE.match(value) for value in values):
                numeric.append(col)
    if not isinstance(numeric, list):
        raise ResultsSpecError(f"tables[{index}].numeric_columns 必须是列编号列表。")
    numeric = sorted({_integer(value, f"tables[{index}].numeric_columns", 0) for value in numeric})
    if any(value >= width for value in numeric):
        raise ResultsSpecError(f"tables[{index}].numeric_columns 超出列数。")
    return {
        "number": number,
        "title": title,
        "columns": [_text(value) for value in columns],
        "rows": [[_text(value) for value in row] for row in rows],
        "note": _text(raw.get("note")).strip(),
        "numeric_columns": numeric,
        "wide": bool(raw.get("wide", width > 7)),
    }


def _validate_figure(raw: dict, index: int) -> dict:
    if not isinstance(raw, dict):
        raise ResultsSpecError(f"figures[{index}] 必须是对象。")
    chart_type = _text(raw.get("type", "line")).lower()
    if chart_type not in SUPPORTED_CHARTS:
        raise ResultsSpecError(f"figures[{index}].type 仅支持 bar、line 或 scatter。")
    title = _text(raw.get("title")).strip()
    if not title:
        raise ResultsSpecError(f"figures[{index}].title 不能为空。")
    series = raw.get("series")
    if not isinstance(series, list) or not series:
        raise ResultsSpecError(f"figures[{index}].series 必须至少包含一组数据。")
    checked = []
    for series_index, item in enumerate(series):
        if not isinstance(item, dict):
            raise ResultsSpecError(f"figures[{index}].series[{series_index}] 必须是对象。")
        xs, ys = item.get("x"), item.get("y")
        if not isinstance(xs, list) or not isinstance(ys, list) or not xs or len(xs) != len(ys):
            raise ResultsSpecError(f"figures[{index}].series[{series_index}] 的 x、y 必须等长且非空。")
        numeric_x = chart_type == "scatter"
        x_values = [_float(x, f"figure {index + 1} x") for x in xs] if numeric_x else [_text(x) for x in xs]
        y_values = [_float(y, f"figure {index + 1} y") for y in ys]
        checked.append({"name": _text(item.get("name", f"Series {series_index + 1}")), "x": x_values, "y": y_values})
    return {
        "number": _integer(raw.get("number", index + 1), f"figures[{index}].number"),
        "title": title,
        "type": chart_type,
        "series": checked,
        "x_label": _text(raw.get("x_label")).strip(),
        "y_label": _text(raw.get("y_label")).strip(),
        "note": _text(raw.get("note")).strip(),
        "alt_text": _text(raw.get("alt_text")).strip() or title,
        "wide": bool(raw.get("wide", False)),
    }


def validate_spec(raw: dict) -> dict:
    version = _integer(raw.get("schema_version", SCHEMA_VERSION), "schema_version")
    if version != SCHEMA_VERSION:
        raise ResultsSpecError(f"不支持 schema_version {version}。")
    tables = raw.get("tables", [])
    figures = raw.get("figures", [])
    if not isinstance(tables, list) or not isinstance(figures, list):
        raise ResultsSpecError("tables 和 figures 必须是列表。")
    if not tables and not figures:
        raise ResultsSpecError("至少需要一个表格或一个图。")
    checked = {
        "schema_version": version,
        "document_title": _text(raw.get("document_title")).strip(),
        "tables": [_validate_table(item, index) for index, item in enumerate(tables)],
        "figures": [_validate_figure(item, index) for index, item in enumerate(figures)],
    }
    table_numbers = [item["number"] for item in checked["tables"]]
    figure_numbers = [item["number"] for item in checked["figures"]]
    if len(table_numbers) != len(set(table_numbers)):
        raise ResultsSpecError("表格编号不能重复。")
    if len(figure_numbers) != len(set(figure_numbers)):
        raise ResultsSpecError("图编号不能重复。")
    return checked


def _font(run, *, italic=None, bold=None, size=12, name="Times New Roman"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if italic is not None:
        run.italic = italic
    if bold is not None:
        run.bold = bold


def _body_paragraph(paragraph, *, alignment=WD_ALIGN_PARAGRAPH.LEFT, spacing=2.0):
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = spacing
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def _set_border(cell, edge: str, *, size="8", color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), size)
    element.set(qn("w:color"), color)


def _remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _decimal_tab(paragraph, text: str):
    # Decimal tabs inside narrow DOCX cells wrap differently in Word and
    # LibreOffice. Right alignment is stable and preserves the supplied text;
    # equal reported precision still produces a visually aligned decimal.
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text)
    _font(run, size=10)


def _caption(doc, label: str, number: int, title: str):
    number_p = doc.add_paragraph()
    _body_paragraph(number_p, spacing=1.0)
    _font(number_p.add_run(f"{label} {number}"), bold=True)
    title_p = doc.add_paragraph()
    _body_paragraph(title_p, spacing=1.0)
    _font(title_p.add_run(title), italic=True)


def _note(doc, text: str):
    if not text:
        return
    paragraph = doc.add_paragraph()
    _body_paragraph(paragraph, spacing=1.0)
    _font(paragraph.add_run("Note. "), italic=True, size=10)
    _font(paragraph.add_run(text), size=10)


def _configure_section(section, wide: bool):
    if wide:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11), Inches(8.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)


def _switch_section(doc, current_wide: bool, target_wide: bool, first: bool) -> bool:
    if first:
        _configure_section(doc.sections[0], target_wide)
    elif current_wide != target_wide:
        _configure_section(doc.add_section(WD_SECTION.NEW_PAGE), target_wide)
    else:
        doc.add_page_break()
    return target_wide


def add_apa_table(doc, spec: dict):
    _caption(doc, "Table", spec["number"], spec["title"])
    table = doc.add_table(rows=1, cols=len(spec["columns"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _remove_table_borders(table)
    header = table.rows[0]
    _repeat_header(header)
    for column, text in enumerate(spec["columns"]):
        cell = header.cells[column]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        _body_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
        _font(paragraph.add_run(text), italic=True, size=10)
        _set_border(cell, "top")
        _set_border(cell, "bottom")
    for row_values in spec["rows"]:
        row = table.add_row()
        for column, text in enumerate(row_values):
            cell = row.cells[column]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            _body_paragraph(paragraph, spacing=1.0)
            if column in spec["numeric_columns"]:
                _decimal_tab(paragraph, text)
            else:
                _font(paragraph.add_run(text), size=10)
    for cell in table.rows[-1].cells:
        _set_border(cell, "bottom")
    _note(doc, spec["note"])


def _chart_bounds(spec: dict):
    ys = [value for series in spec["series"] for value in series["y"]]
    low, high = min(ys), max(ys)
    if spec["type"] == "bar":
        low, high = min(0.0, low), max(0.0, high)
    if math.isclose(low, high):
        padding = abs(low) * 0.1 or 1.0
    else:
        padding = (high - low) * 0.08
    return low - (0 if spec["type"] == "bar" and low == 0 else padding), high + padding


def _font_file():
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    return next((path for path in candidates if Path(path).is_file()), None)


def _chart_geometry(spec: dict, width=1800, height=1100):
    left, right, top, bottom = 210, width - 100, 90, height - 190
    y_low, y_high = _chart_bounds(spec)
    if spec["type"] == "scatter":
        xs = [value for series in spec["series"] for value in series["x"]]
        x_low, x_high = min(xs), max(xs)
        if math.isclose(x_low, x_high):
            x_low, x_high = x_low - 1, x_high + 1
    else:
        labels = spec["series"][0]["x"]
        if any(series["x"] != labels for series in spec["series"]):
            raise ResultsSpecError("bar/line 图的所有系列必须使用相同的 x 标签。")
        x_low, x_high = 0.0, max(1.0, len(labels) - 1.0)
    def px(value):
        return left + (float(value) - x_low) / (x_high - x_low) * (right - left)
    def py(value):
        return bottom - (float(value) - y_low) / (y_high - y_low) * (bottom - top)
    return (left, right, top, bottom, y_low, y_high, px, py)


def _render_png(spec: dict, output: Path):
    width, height = 1800, 1100
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = _font_file()
    font = ImageFont.truetype(font_path, 34) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 28) if font_path else ImageFont.load_default()
    left, right, top, bottom, y_low, y_high, px, py = _chart_geometry(spec, width, height)
    draw.line((left, top, left, bottom), fill="black", width=4)
    draw.line((left, bottom, right, bottom), fill="black", width=4)
    for tick in range(6):
        value = y_low + (y_high - y_low) * tick / 5
        y = py(value)
        draw.line((left - 12, y, left, y), fill="black", width=3)
        draw.text((left - 25, y), f"{value:.2g}", font=small, fill="black", anchor="rm")
    palette = ["#111111", "#666666", "#999999", "#333333"]
    if spec["type"] == "bar":
        count, series_count = len(spec["series"][0]["x"]), len(spec["series"])
        group_width = (right - left) / max(count, 1) * 0.72
        bar_width = group_width / series_count
        zero = py(0)
        for series_index, series in enumerate(spec["series"]):
            for item_index, value in enumerate(series["y"]):
                center = left + (item_index + 0.5) * (right - left) / count
                x0 = center - group_width / 2 + series_index * bar_width
                x1 = x0 + bar_width * 0.9
                draw.rectangle((x0, min(zero, py(value)), x1, max(zero, py(value))), fill=palette[series_index % len(palette)], outline="black", width=2)
    else:
        for series_index, series in enumerate(spec["series"]):
            points = []
            for item_index, (x_value, y_value) in enumerate(zip(series["x"], series["y"])):
                x = px(x_value if spec["type"] == "scatter" else item_index)
                y = py(y_value)
                points.append((x, y))
            if spec["type"] == "line" and len(points) > 1:
                draw.line(points, fill=palette[series_index % len(palette)], width=5)
            for x, y in points:
                draw.ellipse((x - 9, y - 9, x + 9, y + 9), fill="white", outline=palette[series_index % len(palette)], width=5)
    if spec["type"] != "scatter":
        labels = spec["series"][0]["x"]
        count = len(labels)
        for index, label in enumerate(labels):
            x = left + (index + 0.5) * (right - left) / count if spec["type"] == "bar" else px(index)
            draw.text((x, bottom + 30), label, font=small, fill="black", anchor="ma")
    if spec["x_label"]:
        draw.text(((left + right) / 2, height - 45), spec["x_label"], font=font, fill="black", anchor="mm")
    if spec["y_label"]:
        label_layer = Image.new("RGBA", (500, 70), (255, 255, 255, 0))
        label_draw = ImageDraw.Draw(label_layer)
        label_draw.text((250, 35), spec["y_label"], font=font, fill="black", anchor="mm")
        label_layer = label_layer.rotate(90, expand=True)
        image.paste(label_layer, (25, int((height - label_layer.height) / 2)), label_layer)
    if len(spec["series"]) > 1 or spec["series"][0]["name"]:
        x = left
        for index, series in enumerate(spec["series"]):
            draw.line((x, 40, x + 55, 40), fill=palette[index % len(palette)], width=5)
            draw.text((x + 70, 40), series["name"], font=small, fill="black", anchor="lm")
            x += 70 + draw.textlength(series["name"], font=small) + 45
    image.save(output, format="PNG", dpi=(300, 300))


def _svg(spec: dict, width=900, height=550) -> str:
    left, right, top, bottom, y_low, y_high, px_big, py_big = _chart_geometry(spec, width * 2, height * 2)
    scale = 0.5
    left, right, top, bottom = [value * scale for value in (left, right, top, bottom)]
    px = lambda value: px_big(value) * scale
    py = lambda value: py_big(value) * scale
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="white"/>', f'<path d="M {left} {top} V {bottom} H {right}" fill="none" stroke="black" stroke-width="2"/>']
    for tick in range(6):
        value = y_low + (y_high - y_low) * tick / 5
        y = py(value)
        parts.append(f'<path d="M {left - 6} {y} H {left}" stroke="black" stroke-width="1.5"/>')
        parts.append(f'<text x="{left - 12}" y="{y + 5}" text-anchor="end" font-family="Arial" font-size="14">{html.escape(f"{value:.2g}")}</text>')
    palette = ["#111111", "#666666", "#999999", "#333333"]
    if spec["type"] == "bar":
        count, series_count = len(spec["series"][0]["x"]), len(spec["series"])
        group_width = (right - left) / max(count, 1) * 0.72
        bar_width = group_width / series_count
        zero = py(0)
        for series_index, series in enumerate(spec["series"]):
            for item_index, value in enumerate(series["y"]):
                center = left + (item_index + 0.5) * (right - left) / count
                x0 = center - group_width / 2 + series_index * bar_width
                y = py(value)
                parts.append(f'<rect x="{x0:.2f}" y="{min(zero, y):.2f}" width="{bar_width * .9:.2f}" height="{abs(zero-y):.2f}" fill="{palette[series_index % len(palette)]}" stroke="black"/>')
    else:
        for series_index, series in enumerate(spec["series"]):
            points = [(px(x if spec["type"] == "scatter" else i), py(y)) for i, (x, y) in enumerate(zip(series["x"], series["y"]))]
            if spec["type"] == "line" and len(points) > 1:
                d = " ".join(("M" if i == 0 else "L") + f" {x:.2f} {y:.2f}" for i, (x, y) in enumerate(points))
                parts.append(f'<path d="{d}" fill="none" stroke="{palette[series_index % len(palette)]}" stroke-width="2.5"/>')
            for x, y in points:
                parts.append(f'<circle cx="{x:.2f}" cy="{y:.2f}" r="4" fill="white" stroke="{palette[series_index % len(palette)]}" stroke-width="2.5"/>')
    if spec["type"] != "scatter":
        labels = spec["series"][0]["x"]
        for index, label in enumerate(labels):
            x = left + (index + .5) * (right - left) / len(labels) if spec["type"] == "bar" else px(index)
            parts.append(f'<text x="{x:.2f}" y="{bottom + 28}" text-anchor="middle" font-family="Arial" font-size="14">{html.escape(label)}</text>')
    if spec["x_label"]:
        parts.append(f'<text x="{(left+right)/2:.2f}" y="{height-18}" text-anchor="middle" font-family="Arial" font-size="16" font-weight="bold">{html.escape(spec["x_label"])}</text>')
    if spec["y_label"]:
        parts.append(f'<text x="22" y="{(top+bottom)/2:.2f}" text-anchor="middle" transform="rotate(-90 22 {(top+bottom)/2:.2f})" font-family="Arial" font-size="16" font-weight="bold">{html.escape(spec["y_label"])}</text>')
    parts.append("</svg>")
    return "".join(parts)


def add_apa_figure(doc, spec: dict, png_path: Path):
    _caption(doc, "Figure", spec["number"], spec["title"])
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.line_spacing = None
    shape = paragraph.add_run().add_picture(str(png_path), width=Inches(8.4 if spec["wide"] else 6.4))
    shape._inline.docPr.set("descr", spec["alt_text"])
    _note(doc, spec["note"])


def statistical_warnings(spec: dict) -> list[dict]:
    warnings = []
    locations = []
    for table in spec["tables"]:
        locations.append((f"Table {table['number']} title", table["title"]))
        locations.append((f"Table {table['number']} note", table["note"]))
        for row_index, row in enumerate(table["rows"], 1):
            for column_index, value in enumerate(row, 1):
                locations.append((f"Table {table['number']} R{row_index}C{column_index}", value))
    for figure in spec["figures"]:
        locations.extend(((f"Figure {figure['number']} title", figure["title"]), (f"Figure {figure['number']} note", figure["note"])))
    for location, value in locations:
        for match in P_RE.finditer(value):
            operator, raw = match.groups()
            try:
                number = float(raw)
            except ValueError:
                continue
            if operator == "=" and number == 0:
                warnings.append({"code": "p_zero", "location": location, "message": "p = .000 需要作者根据原始输出确认，工具没有改动该数值。"})
            if number < 0 or number > 1:
                warnings.append({"code": "p_range", "location": location, "message": "p 值不在 0 到 1 之间，需要作者确认。"})
    return warnings


def build_results_document(spec: dict, output: Path, *, export_svg: Path | None = None) -> tuple[Path, dict]:
    output = Path(output)
    if output.exists():
        raise FileExistsError(f"输出文件已存在：{output}")
    if output.suffix.lower() != ".docx":
        raise ResultsSpecError("输出文件必须使用 .docx 扩展名。")
    if export_svg is not None and Path(export_svg).exists():
        raise FileExistsError(f"SVG 输出文件夹已存在：{export_svg}")
    checked = validate_spec(spec)
    warnings = statistical_warnings(checked)
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if checked["document_title"]:
        paragraph = doc.add_paragraph()
        _body_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        _font(paragraph.add_run(checked["document_title"]), bold=True)
    items = [("table", item) for item in checked["tables"]] + [("figure", item) for item in checked["figures"]]
    current_wide = False
    first = True
    svg_manifest = []
    with tempfile.TemporaryDirectory(prefix="apa7_results_") as temporary:
        temporary = Path(temporary)
        for kind, item in items:
            current_wide = _switch_section(doc, current_wide, item["wide"], first)
            first = False
            if kind == "table":
                add_apa_table(doc, item)
            else:
                png = temporary / f"figure-{item['number']}.png"
                _render_png(item, png)
                add_apa_figure(doc, item, png)
                if export_svg is not None:
                    svg_manifest.append((f"figure-{item['number']}.svg", _svg(item)))
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)
    if export_svg is not None:
        export_svg = Path(export_svg)
        export_svg.mkdir(parents=True)
        items_manifest = []
        for filename, content in svg_manifest:
            (export_svg / filename).write_text(content, encoding="utf-8")
            items_manifest.append({"output": filename, "method": "reconstructed_from_supplied_data"})
        (export_svg / "manifest.json").write_text(json.dumps({"items": items_manifest}, indent=2), encoding="utf-8")
    report = {
        "tables_created": len(checked["tables"]),
        "figures_created": len(checked["figures"]),
        "values_recomputed": False,
        "statistical_warnings": warnings,
        "output": str(output),
    }
    if export_svg is not None:
        report["svg_output"] = str(export_svg)
    return output, report


def main(argv=None):
    parser = argparse.ArgumentParser(description="Generate APA-style Word tables and figures from supplied CSV/JSON results.")
    parser.add_argument("input", type=Path, help="JSON specification or CSV table")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--export-svg", type=Path, help="Optional new folder for data-derived SVG figures")
    parser.add_argument("--table-title", help="Required when input is CSV")
    parser.add_argument("--table-number", type=int, default=1)
    parser.add_argument("--note", default="")
    args = parser.parse_args(argv)
    try:
        if args.input.suffix.lower() == ".csv":
            if not args.table_title:
                parser.error("CSV 输入需要 --table-title。")
            spec = csv_table_spec(args.input, title=args.table_title, number=args.table_number, note=args.note)
        else:
            spec = load_spec(args.input)
        _, report = build_results_document(spec, args.output, export_svg=args.export_svg)
    except (ResultsSpecError, FileExistsError) as exc:
        parser.error(str(exc))
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Repeatable corpus benchmark for the APA 7 Word Formatter.

The benchmark never treats successful file creation as proof of quality. Each
case is checked for source preservation, a readable single output, saved-format
verification, expected warning codes, and (when requested) a fresh page review.
Private manuscripts are referenced in place and are never copied into the repo.
"""
from __future__ import annotations

import argparse
import base64
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.shared import Inches

import apa7_format as engine
import apa7_workflow as workflow


SCHEMA_VERSION = 1
SOURCE_REGISTRY_SCHEMA_VERSION = 1
CASE_ID = re.compile(r"[a-z0-9][a-z0-9_-]{0,63}")
HEX_MD5 = re.compile(r"[0-9a-f]{32}")
HEX_SHA256 = re.compile(r"[0-9a-f]{64}")
MACHINE_STATUSES = {"passed", "needs_review", "not_applicable"}
ISSUE_REPORTS = {
    "front_matter_check": "front_matter",
    "reference_quality_check": "reference_quality",
    "numbered_object_check": "numbered_object",
    "caption_object_check": "caption_object",
    "statistical_reporting": "statistics",
}
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def read_json(path):
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"JSON 顶层必须是对象：{path}")
    return value


def write_new_json(path, value):
    path = Path(path)
    if path.exists() or path.is_symlink():
        raise ValueError(f"不会覆盖已有文件：{path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as output:
        json.dump(value, output, ensure_ascii=False, indent=2)
        output.write("\n")
    return path


def atomic_json(path, value):
    path = Path(path)
    temporary = path.with_name(path.name + ".tmp")
    if temporary.exists() or temporary.is_symlink():
        raise ValueError(f"临时输出已存在：{temporary}")
    with temporary.open("x", encoding="utf-8") as output:
        json.dump(value, output, ensure_ascii=False, indent=2)
        output.write("\n")
    temporary.replace(path)


def write_new_text(path, text):
    path = Path(path)
    if path.exists() or path.is_symlink():
        raise ValueError(f"不会覆盖已有文件：{path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as output:
        output.write(text)
    return path


def atomic_text(path, text):
    path = Path(path)
    temporary = path.with_name(path.name + ".tmp")
    if temporary.exists() or temporary.is_symlink():
        raise ValueError(f"临时输出已存在：{temporary}")
    with temporary.open("x", encoding="utf-8") as output:
        output.write(text)
    temporary.replace(path)


def _string_list(value, label):
    if not isinstance(value, list) or any(not isinstance(item, str) or not item for item in value):
        raise ValueError(f"{label} 必须是非空字符串列表。")
    return value


def validate_manifest(manifest):
    if manifest.get("schema_version") != SCHEMA_VERSION:
        raise ValueError(f"schema_version 必须为 {SCHEMA_VERSION}。")
    cases = manifest.get("cases")
    if not isinstance(cases, list) or not cases:
        raise ValueError("cases 必须是非空列表。")
    seen = set()
    for case in cases:
        if not isinstance(case, dict):
            raise ValueError("每个测试案例必须是对象。")
        case_id = case.get("id")
        if not isinstance(case_id, str) or not CASE_ID.fullmatch(case_id):
            raise ValueError("案例 id 只能使用小写字母、数字、下划线和连字符。")
        if case_id in seen:
            raise ValueError(f"案例 id 重复：{case_id}")
        seen.add(case_id)
        if case.get("profile") not in {"student", "professional"}:
            raise ValueError(f"{case_id}: profile 必须是 student 或 professional。")
        for key in ("source", "config"):
            if not isinstance(case.get(key), str) or not case[key].strip():
                raise ValueError(f"{case_id}: 缺少 {key}。")
        if case["profile"] == "professional" and not str(case.get("running_head", "")).strip():
            raise ValueError(f"{case_id}: 专业论文基准需要明确 running_head。")
        if case.get("font", "Times New Roman") not in engine.FONTS:
            raise ValueError(f"{case_id}: 字体不在支持列表中。")
        if not isinstance(case.get("require_visual_review", True), bool):
            raise ValueError(f"{case_id}: require_visual_review 必须是布尔值。")
        _string_list(case.get("tags", []), f"{case_id}.tags")
        expect = case.get("expect", {})
        if not isinstance(expect, dict):
            raise ValueError(f"{case_id}.expect 必须是对象。")
        checks = expect.get("machine_checks", {"content_preservation": "passed"})
        if not isinstance(checks, dict) or any(
            not isinstance(key, str) or value not in MACHINE_STATUSES for key, value in checks.items()
        ):
            raise ValueError(f"{case_id}.expect.machine_checks 含无效状态。")
        for key in ("required_issue_codes", "forbidden_issue_codes"):
            _string_list(expect.get(key, []), f"{case_id}.expect.{key}")
        if "exact_issue_codes" in expect:
            _string_list(expect["exact_issue_codes"], f"{case_id}.expect.exact_issue_codes")
        maximum = expect.get("preflight_max_issues")
        if maximum is not None and (type(maximum) is not int or maximum < 0):
            raise ValueError(f"{case_id}.expect.preflight_max_issues 必须是非负整数。")
    return manifest


def validate_source_registry(registry):
    """Validate metadata for public, openly licensed real-document cases.

    The registry is committed; downloaded Word files are not. This keeps
    provenance and exact checksums reviewable without redistributing a corpus
    or silently changing a held-out test document.
    """
    if registry.get("schema_version") != SOURCE_REGISTRY_SCHEMA_VERSION:
        raise ValueError(f"source registry schema_version 必须为 {SOURCE_REGISTRY_SCHEMA_VERSION}。")
    sources = registry.get("sources")
    if not isinstance(sources, list) or not sources:
        raise ValueError("source registry 的 sources 必须是非空列表。")
    seen_ids, seen_names = set(), set()
    required = {
        "id", "title", "landing_page", "download_url", "license", "license_url",
        "local_filename", "size_bytes", "md5", "sha256", "intended_use",
    }
    for source in sources:
        if not isinstance(source, dict) or not required.issubset(source):
            raise ValueError("每个公开来源必须包含来源、许可、文件名、大小、哈希和用途。")
        source_id = source["id"]
        if not isinstance(source_id, str) or not CASE_ID.fullmatch(source_id) or source_id in seen_ids:
            raise ValueError(f"无效或重复的公开来源 id：{source_id!r}")
        seen_ids.add(source_id)
        filename = source["local_filename"]
        if not isinstance(filename, str) or Path(filename).name != filename or not filename.lower().endswith(".docx"):
            raise ValueError(f"{source_id}: local_filename 必须是单个 .docx 文件名。")
        if filename in seen_names:
            raise ValueError(f"公开来源文件名重复：{filename}")
        seen_names.add(filename)
        for key in ("title", "license", "intended_use"):
            if not isinstance(source[key], str) or not source[key].strip():
                raise ValueError(f"{source_id}: {key} 不能为空。")
        for key in ("landing_page", "download_url", "license_url"):
            if not isinstance(source[key], str) or not source[key].startswith("https://"):
                raise ValueError(f"{source_id}: {key} 必须是 https URL。")
        if type(source["size_bytes"]) is not int or source["size_bytes"] <= 0:
            raise ValueError(f"{source_id}: size_bytes 必须是正整数。")
        if not isinstance(source["md5"], str) or not HEX_MD5.fullmatch(source["md5"]):
            raise ValueError(f"{source_id}: md5 必须是 32 位小写十六进制。")
        if not isinstance(source["sha256"], str) or not HEX_SHA256.fullmatch(source["sha256"]):
            raise ValueError(f"{source_id}: sha256 必须是 64 位小写十六进制。")
    return registry


def verify_public_sources(registry_path, corpus_dir):
    registry_path = Path(registry_path).expanduser().resolve()
    corpus_dir = Path(corpus_dir).expanduser().resolve()
    registry = validate_source_registry(read_json(registry_path))
    results = []
    for source in registry["sources"]:
        path = corpus_dir / source["local_filename"]
        item = {"id": source["id"], "file": str(path), "status": "missing"}
        if path.is_file():
            data = path.read_bytes()
            actual = {
                "size_bytes": len(data),
                "md5": hashlib.md5(data).hexdigest(),  # nosec B324 - published integrity checksum
                "sha256": hashlib.sha256(data).hexdigest(),
            }
            expected = {key: source[key] for key in actual}
            item.update(actual=actual, expected=expected)
            item["status"] = "passed" if actual == expected else "mismatch"
        results.append(item)
    return {
        "registry": str(registry_path),
        "corpus_dir": str(corpus_dir),
        "status": "passed" if all(item["status"] == "passed" for item in results) else "failed",
        "sources": results,
    }


def corpus_inspection_markdown(report):
    summary = report["summary"]
    lines = [
        "# APA7 Public Corpus Inspection",
        "",
        f"Overall: **{report['status'].upper()}**",
        "",
        (
            f"Documents: {summary['documents']} | Inspected: {summary['inspected']} | "
            f"Failed: {summary['failed']} | Profile: {report['profile']}"
        ),
        "",
        (
            f"Paragraphs: {summary['paragraphs']} | Tables: {summary['tables']} | "
            f"Drawings: {summary['drawing_objects']} | Native charts: {summary['native_charts']} | "
            f"Native-equation paragraphs: {summary['native_math_paragraphs']} | "
            f"Statistical expressions: {summary['statistical_expressions']}"
        ),
        "",
        "| Source | Paragraphs | Tables | Drawings | Equations | Statistics | Preflight issues | Result |",
        "| --- | ---: | ---: | ---: | ---: | ---: | ---: | --- |",
    ]
    for item in report["documents"]:
        lines.append(
            f"| {item['id']} | {item.get('paragraphs', 0)} | {item.get('tables', 0)} | "
            f"{item.get('drawing_objects', 0)} | {item.get('native_math_paragraphs', 0)} | "
            f"{item.get('statistical_expressions', 0)} | {item.get('preflight_issue_count', 0)} | "
            f"{item['status']} |"
        )
    lines.extend([
        "",
        "This is a read-only ingestion and structure-inspection result. It is not proof that every document has completed APA formatting or page-by-page visual review.",
        "",
    ])
    return "\n".join(lines)


def inspect_public_sources(registry_path, corpus_dir, profile, output_dir):
    """Verify and inspect every registered DOCX without formatting or copying it."""
    if profile not in {"student", "professional"}:
        raise ValueError("profile 必须是 student 或 professional。")
    output_dir = Path(output_dir).expanduser().resolve()
    if output_dir.exists() or output_dir.is_symlink():
        raise ValueError("整库扫描输出目录必须是新目录，不能覆盖旧结果。")
    registry_path = Path(registry_path).expanduser().resolve()
    corpus_dir = Path(corpus_dir).expanduser().resolve()
    verification = verify_public_sources(registry_path, corpus_dir)
    if verification["status"] != "passed":
        bad = [item["id"] for item in verification["sources"] if item["status"] != "passed"]
        raise ValueError("公开来源校验失败，拒绝扫描：" + ", ".join(bad))
    registry = validate_source_registry(read_json(registry_path))
    totals = {
        "documents": len(registry["sources"]),
        "inspected": 0,
        "failed": 0,
        "paragraphs": 0,
        "tables": 0,
        "drawing_objects": 0,
        "native_charts": 0,
        "native_math_paragraphs": 0,
        "statistical_expressions": 0,
        "preflight_issues": 0,
    }
    documents = []
    for source in registry["sources"]:
        path = corpus_dir / source["local_filename"]
        before = engine.digest(path)
        item = {
            "id": source["id"],
            "title": source["title"],
            "local_filename": source["local_filename"],
            "source_sha256": before,
            "status": "failed",
        }
        try:
            draft = workflow.prepare(path, profile)
            if engine.digest(path) != before:
                raise RuntimeError("只读结构扫描意外改变了原稿。")
            review = draft.get("_review", {})
            objects = review.get("object_summary", {})
            preflight = review.get("preflight_summary", {})
            item.update(
                status="passed",
                paragraphs=len(review.get("paragraphs", [])),
                tables=objects.get("top_level_tables", 0),
                drawing_objects=objects.get("inline_drawings", 0) + objects.get("floating_drawings", 0),
                native_charts=objects.get("native_charts", 0),
                native_math_paragraphs=objects.get("native_math_paragraphs", 0),
                statistical_expressions=objects.get("statistical_expressions", 0),
                preflight_issue_count=preflight.get("issue_count", 0),
            )
            totals["inspected"] += 1
            for key in (
                "paragraphs", "tables", "drawing_objects", "native_charts",
                "native_math_paragraphs", "statistical_expressions",
            ):
                totals[key] += item[key]
            totals["preflight_issues"] += item["preflight_issue_count"]
        except Exception as exc:
            if path.is_file() and engine.digest(path) != before:
                item["source_changed"] = True
            item["error"] = f"{type(exc).__name__}: {exc}"
            totals["failed"] += 1
        documents.append(item)
    report = {
        "schema_version": SOURCE_REGISTRY_SCHEMA_VERSION,
        "formatter_version": engine.VERSION,
        "created_at": utc_now(),
        "profile": profile,
        "registry": str(registry_path),
        "registry_sha256": engine.digest(registry_path),
        "corpus_dir": str(corpus_dir),
        "source_verification": verification["status"],
        "status": "passed" if totals["failed"] == 0 else "failed",
        "summary": totals,
        "documents": documents,
        "limitation": (
            "Read-only ingestion and structure inspection only; this does not certify completed APA formatting "
            "or page-by-page visual review."
        ),
    }
    write_new_json(output_dir / "corpus-inspection.json", report)
    write_new_text(output_dir / "corpus-inspection.md", corpus_inspection_markdown(report))
    return report


def resolve_case_paths(manifest_path, case):
    base = Path(manifest_path).resolve().parent

    def resolve(value):
        path = Path(value).expanduser()
        return path.resolve() if path.is_absolute() else (base / path).resolve()

    return resolve(case["source"]), resolve(case["config"])


def collect_issue_codes(report):
    codes = set()
    for report_key, prefix in ISSUE_REPORTS.items():
        value = report.get(report_key) or {}
        for issue in value.get("issues", []):
            if isinstance(issue, dict) and issue.get("code"):
                codes.add(f"{prefix}:{issue['code']}")
    citation = report.get("citation_reference_check") or {}
    if citation.get("unmatched_citations"):
        codes.add("citation_reference:unmatched_citation")
    if citation.get("uncited_references"):
        codes.add("citation_reference:uncited_reference")
    if citation.get("unparsed_reference_paragraphs"):
        codes.add("citation_reference:unparsed_reference")
    links = report.get("reference_link_check") or {}
    if links.get("unlinked_urls"):
        codes.add("reference_link:unlinked_url")
    if links.get("bare_dois"):
        codes.add("reference_link:bare_doi")
    if (report.get("visuals") or {}).get("export_error"):
        codes.add("visual_export:error")
    return sorted(codes)


def evaluate_expectations(case, report, hard_gates):
    failures = [label for label, passed in hard_gates.items() if not passed]
    expect = case.get("expect", {})
    actual_checks = {item["id"]: item["status"] for item in report.get("machine_checks", [])}
    expected_checks = expect.get("machine_checks", {"content_preservation": "passed"})
    for check_id, expected in expected_checks.items():
        actual = actual_checks.get(check_id)
        if actual != expected:
            failures.append(f"machine_check {check_id}: expected {expected}, got {actual or 'missing'}")
    actual_codes = collect_issue_codes(report)
    required = set(expect.get("required_issue_codes", []))
    forbidden = set(expect.get("forbidden_issue_codes", []))
    missing = sorted(required - set(actual_codes))
    unexpected = sorted(forbidden & set(actual_codes))
    if missing:
        failures.append("missing required issue codes: " + ", ".join(missing))
    if unexpected:
        failures.append("found forbidden issue codes: " + ", ".join(unexpected))
    if "exact_issue_codes" in expect:
        exact = sorted(set(expect["exact_issue_codes"]))
        if actual_codes != exact:
            failures.append(f"issue codes differ: expected {exact}, got {actual_codes}")
    maximum = expect.get("preflight_max_issues")
    issue_count = (report.get("preflight_summary") or {}).get("issue_count")
    if maximum is not None and (type(issue_count) is not int or issue_count > maximum):
        failures.append(f"preflight issue count {issue_count!r} exceeds {maximum}")
    return failures, actual_codes, actual_checks


def review_template(document_hash, pages):
    return {
        "document_sha256": document_hash,
        "pages": [
            {"page": page["page"], "status": "pending", "notes": ""}
            for page in pages
        ],
    }


def run_case(manifest_path, case, run_dir, renderer=None):
    case_id = case["id"]
    case_dir = run_dir / "cases" / case_id
    case_dir.mkdir(parents=True)
    source, config_path = resolve_case_paths(manifest_path, case)
    source_before = engine.digest(source) if source.is_file() else None
    output = case_dir / f"{case_id}_APA7.docx"
    result = {
        "id": case_id,
        "profile": case["profile"],
        "tags": case.get("tags", []),
        "require_visual_review": case.get("require_visual_review", True),
        "engine_status": "failed",
        "expectation_status": "failed",
        "visual_status": "not_run",
        "status": "failed",
        "failures": [],
    }
    try:
        if not source.is_file() or source.suffix.lower() != ".docx":
            raise ValueError("案例源文件必须是现有 .docx。")
        if not config_path.is_file():
            raise ValueError("案例缺少已审核的结构配置。")
        out, report = workflow.apply_reviewed(
            source,
            read_json(config_path),
            profile=case["profile"],
            output=output,
            font=case.get("font", "Times New Roman"),
            running_head=case.get("running_head", ""),
            export_visuals=False,
            add_styles=False,
        )
        source_unchanged = engine.digest(source) == source_before
        output_exists = out.is_file()
        output_openable = False
        if output_exists:
            Document(out)
            output_openable = True
        one_output = len(list(case_dir.glob("*.docx"))) == 1
        content_check = next(
            (item for item in report.get("machine_checks", []) if item.get("id") == "content_preservation"),
            {},
        )
        hard_gates = {
            "source_unchanged": source_unchanged,
            "single_output": one_output,
            "output_openable": output_openable,
            "formatter_preservation_passed": str(report.get("preservation", "")).startswith("passed"),
            "content_preservation_check_passed": content_check.get("status") == "passed",
        }
        failures, issue_codes, machine_checks = evaluate_expectations(case, report, hard_gates)
        write_new_json(case_dir / "formatter-report.json", report)
        result.update(
            engine_status="passed" if all(hard_gates.values()) else "failed",
            expectation_status="passed" if not failures else "failed",
            hard_gates=hard_gates,
            issue_codes=issue_codes,
            machine_checks=machine_checks,
            preflight_issue_count=(report.get("preflight_summary") or {}).get("issue_count"),
            source_sha256=source_before,
            output_sha256=engine.digest(out),
            output=str(out.relative_to(run_dir)),
            failures=failures,
        )
        if failures:
            result["status"] = "failed"
        elif result["require_visual_review"]:
            if renderer is None:
                result["visual_status"] = "not_rendered"
            else:
                render_dir = case_dir / "render"
                render = workflow.render_document(out, renderer, render_dir)
                write_new_json(
                    case_dir / "page-review.template.json",
                    review_template(render["document_sha256"], render["pages"]),
                )
                result["visual_status"] = "awaiting_review"
                result["render_dir"] = str(render_dir.relative_to(run_dir))
                result["page_count"] = len(render["pages"])
            result["status"] = "awaiting_visual_review"
        else:
            result["visual_status"] = "not_required"
            result["status"] = "passed"
    except Exception as exc:
        if source_before and source.is_file() and engine.digest(source) != source_before:
            result["failures"].append("source changed during failed run")
        result["failures"].append(f"{type(exc).__name__}: {exc}")
    write_new_json(case_dir / "result.json", result)
    return result


def quality_summary(results):
    cases = results.get("cases", [])
    failed = sum(item.get("status") == "failed" for item in cases)
    pending = sum(item.get("status") == "awaiting_visual_review" for item in cases)
    passed = sum(item.get("status") == "passed" for item in cases)
    visual_required = sum(bool(item.get("require_visual_review")) for item in cases)
    visual_passed = sum(item.get("visual_status") == "passed" for item in cases)
    visual_issues = sum(item.get("visual_status") == "issue" for item in cases)
    overall = "failed" if failed or visual_issues else "needs_visual_review" if pending else "passed"
    return {
        "status": overall,
        "cases": len(cases),
        "passed": passed,
        "failed": failed,
        "awaiting_visual_review": pending,
        "visual_required": visual_required,
        "visual_passed": visual_passed,
        "visual_issues": visual_issues,
    }


def summary_markdown(results):
    summary = quality_summary(results)
    labels = {"passed": "PASS", "failed": "FAIL", "needs_visual_review": "NEEDS VISUAL REVIEW"}
    lines = [
        "# APA7 Benchmark Summary",
        "",
        f"Overall: **{labels[summary['status']]}**",
        "",
        (
            f"Cases: {summary['cases']} | Passed: {summary['passed']} | Failed: {summary['failed']} | "
            f"Awaiting visual review: {summary['awaiting_visual_review']}"
        ),
        "",
        "| Case | Engine | Expectations | Visual | Result |",
        "| --- | --- | --- | --- | --- |",
    ]
    for case in results.get("cases", []):
        lines.append(
            f"| {case['id']} | {case.get('engine_status', 'unknown')} | "
            f"{case.get('expectation_status', 'unknown')} | {case.get('visual_status', 'unknown')} | "
            f"{case.get('status', 'unknown')} |"
        )
    failures = [case for case in results.get("cases", []) if case.get("failures")]
    if failures:
        lines.extend(["", "## Failures", ""])
        for case in failures:
            lines.append(f"- **{case['id']}**: " + "; ".join(case["failures"]))
    lines.extend(
        [
            "",
            "A suite is good only when every engine gate and expectation passes and every required visual review is recorded as passed.",
            "",
        ]
    )
    return "\n".join(lines)


def save_run_summary(run_dir, results, *, creating=False):
    results["summary"] = quality_summary(results)
    results["updated_at"] = utc_now()
    result_path = run_dir / "benchmark-results.json"
    summary_path = run_dir / "benchmark-summary.md"
    if creating:
        write_new_json(result_path, results)
        write_new_text(summary_path, summary_markdown(results))
    else:
        atomic_json(result_path, results)
        atomic_text(summary_path, summary_markdown(results))


def run_suite(manifest_path, output_dir, renderer=None):
    manifest_path = Path(manifest_path).expanduser().resolve()
    manifest = validate_manifest(read_json(manifest_path))
    output_dir = Path(output_dir).expanduser().resolve()
    if output_dir.exists() or output_dir.is_symlink():
        raise ValueError("基准输出目录必须是新目录，不能覆盖旧结果。")
    output_dir.mkdir(parents=True)
    renderer_path = Path(renderer).expanduser().resolve() if renderer else None
    if renderer_path is not None and not renderer_path.is_file():
        raise ValueError("找不到指定的 render_docx.py。")
    results = {
        "schema_version": SCHEMA_VERSION,
        "suite": manifest.get("suite", manifest_path.stem),
        "formatter_version": engine.VERSION,
        "created_at": utc_now(),
        "manifest_sha256": engine.digest(manifest_path),
        "cases": [],
    }
    for case in manifest["cases"]:
        results["cases"].append(run_case(manifest_path, case, output_dir, renderer_path))
    save_run_summary(output_dir, results, creating=True)
    return results


def record_case_qa(run_dir, case_id, review_path):
    run_dir = Path(run_dir).expanduser().resolve()
    results_path = run_dir / "benchmark-results.json"
    results = read_json(results_path)
    matches = [item for item in results.get("cases", []) if item.get("id") == case_id]
    if len(matches) != 1:
        raise ValueError(f"基准结果中找不到唯一案例：{case_id}")
    case = matches[0]
    if case.get("status") == "failed":
        raise ValueError("该案例的引擎或预期检查已失败，不能记录为视觉通过。")
    if not case.get("render_dir") or not case.get("output"):
        raise ValueError("该案例没有可用的渲染页面，请用 --renderer 重新运行。")
    case_dir = run_dir / "cases" / case_id
    qa_path = case_dir / "qa.json"
    qa = workflow.record_qa(
        run_dir / case["output"],
        run_dir / case["render_dir"],
        read_json(review_path),
    )
    write_new_json(qa_path, qa)
    case["visual_status"] = "passed" if qa["status"] == "VISUAL_REVIEW_RECORDED" else "issue"
    case["status"] = "passed" if case["visual_status"] == "passed" else "failed"
    if case["visual_status"] == "issue":
        case.setdefault("failures", []).append("逐页视觉复核发现需要修正的页面。")
    case["qa"] = str(qa_path.relative_to(run_dir))
    save_run_summary(run_dir, results)
    return qa, results["summary"]


def _paragraph_adder(doc, roles):
    def add(text, role, style=None, *, page_break=False):
        paragraph = doc.add_paragraph(text, style)
        if page_break:
            paragraph.paragraph_format.page_break_before = True
        roles[str(len(doc.paragraphs))] = role
        return paragraph
    return add


def _compliance(present, needs_review=()):
    present = set(present)
    needs_review = set(needs_review)
    reasons = {
        "target_requirements": "Synthetic benchmark uses the selected APA profile without an external template.",
        "title_page": "Synthetic title-page elements are explicitly identified for the selected profile.",
        "abstract_keywords": "No abstract or keywords are included in this focused fixture.",
        "headings": "The fixture declares the intended heading hierarchy explicitly.",
        "citations_references": "Citation and reference expectations are defined by the fixture.",
        "statistics": "Statistical expressions are included to exercise reporting checks.",
        "equations": "No independent equation is included in this fixture.",
        "tables": "Table purpose and header rows are explicitly defined by the fixture.",
        "figures": "Figure purpose, caption, and callout are explicitly defined by the fixture.",
        "appendices": "No appendix is included in this focused fixture.",
    }
    result = {}
    for key in workflow.COMPLIANCE_AREAS:
        if key in {"target_requirements", "title_page"} or key in present:
            status = "needs_review" if key in needs_review else "pass"
        else:
            status = "not_applicable"
        result[key] = {"status": status, "reason": reasons[key]}
    return result


def _reviewed_config(source, profile, roles, table_roles, header_rows, title_page, compliance, *, uncertain_tables=()):
    config = workflow.prepare(source, profile)
    config.update(
        review_status="reviewed",
        profile=profile,
        roles=roles,
        table_roles=table_roles,
        table_header_rows=header_rows,
        title_page=title_page,
        run_in_headings={},
        replace_headers=False,
        compliance_review=compliance,
        unresolved=[],
    )
    uncertain_tables = {str(item) for item in uncertain_tables}
    config["decision_notes"] = {
        "paragraphs": {
            key: {
                "confidence": "high",
                "reason": f"Synthetic fixture explicitly defines paragraph {key} as {role}.",
            }
            for key, role in roles.items()
        },
        "tables": {
            key: {
                "confidence": "uncertain" if key in uncertain_tables else "high",
                "reason": (
                    "Synthetic merged layout is intentionally preserved for safety."
                    if key in uncertain_tables
                    else "Synthetic fixture explicitly defines this as a simple data table with one header row."
                ),
            }
            for key in table_roles
        },
    }
    return config


def _student_core(folder):
    source = folder / "student_core.docx"
    doc, roles = Document(), {}
    add = _paragraph_adder(doc, roles)
    add("Effects of Practice on Response Time", "title", "Title")
    add("Alex Example", "title_meta")
    add("Department of Psychology, Example University", "title_meta")
    add("PSY 301: Research Methods", "title_meta")
    add("Dr. Taylor Example", "title_meta")
    add("September 18, 2026", "title_meta")
    add("Effects of Practice on Response Time", "title", "Title", page_break=True)
    add("Example (2024) predicted faster responses after practice.", "body")
    add("Method", "heading1", "Heading 1")
    add("Participants", "heading2", "Heading 2")
    add("Forty synthetic participants completed the task.", "body")
    add("Results", "heading1", "Heading 1")
    add("Table 1 summarizes the conditions. The effect was t(38) = 2.41, p = .021, d = 0.76, 95% CI [0.12, 1.40].", "body")
    add("Table 1", "caption_number")
    add("Synthetic Response-Time Summary", "caption_title")
    table = doc.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, [("Condition", "M", "SD"), ("Practice", "420", "35"), ("Control", "455", "41")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    add("Note. Values are entirely synthetic.", "note")
    add("References", "section", "Heading 1", page_break=True)
    add("Example, A. (2024). Synthetic practice effects. Example Journal. https://example.org/article", "reference")
    doc.save(source)
    config = _reviewed_config(
        source, "student", roles, {"1": "data"}, {"1": 1}, [1, 6],
        _compliance({"headings", "citations_references", "statistics", "tables"}),
    )
    return source, config, {
        "id": "student_core",
        "profile": "student",
        "source": source.name,
        "config": "student_core.structure.json",
        "font": "Times New Roman",
        "tags": ["student", "headings", "statistics", "table", "references", "hyperlink"],
        "require_visual_review": False,
        "expect": {
            "machine_checks": {"content_preservation": "passed"},
            "exact_issue_codes": [],
            "preflight_max_issues": 0,
        },
    }


def _professional_warnings(folder, image):
    source = folder / "professional_warnings.docx"
    doc, roles = Document(), {}
    add = _paragraph_adder(doc, roles)
    add("A Synthetic Professional Manuscript", "title", "Title")
    add("Alex Example", "title_meta")
    add("Example Research Institute", "title_meta")
    add("A Synthetic Professional Manuscript", "title", "Title", page_break=True)
    add("Figure 2 presents the synthetic trend.", "body")
    add("Results", "heading1", "Heading 1")
    add("The model was significant, F(1, 28) = 4.20, p = .000.", "body")
    add("Figure 2", "caption_number")
    add("Synthetic Trend", "caption_title")
    picture = doc.add_paragraph()
    picture.add_run().add_picture(str(image), width=Inches(2.0))
    roles[str(len(doc.paragraphs))] = "preserve"
    add("Note. The image is a deterministic synthetic fixture.", "note")
    doc.save(source)
    config = _reviewed_config(
        source, "professional", roles, {}, {}, [1, 3],
        _compliance({"headings", "statistics", "figures"}, {"statistics"}),
    )
    expected = [
        "numbered_object:sequence_gap",
        "statistics:p_zero",
        "statistics:test_missing_ci",
        "statistics:test_missing_effect",
    ]
    return source, config, {
        "id": "professional_warnings",
        "profile": "professional",
        "running_head": "SYNTHETIC PROFESSIONAL MANUSCRIPT",
        "source": source.name,
        "config": "professional_warnings.structure.json",
        "font": "Times New Roman",
        "tags": ["professional", "statistics", "figure", "expected-warning"],
        "require_visual_review": False,
        "expect": {
            "machine_checks": {"content_preservation": "passed"},
            "exact_issue_codes": expected,
        },
    }


def _complex_preserve(folder):
    source = folder / "complex_preserve.docx"
    doc, roles = Document(), {}
    add = _paragraph_adder(doc, roles)
    add("Preservation of a Complex Layout", "title", "Title")
    add("Alex Example", "title_meta")
    add("Department of Psychology, Example University", "title_meta")
    add("PSY 301: Research Methods", "title_meta")
    add("Dr. Taylor Example", "title_meta")
    add("September 18, 2026", "title_meta")
    add("Preservation of a Complex Layout", "title", "Title", page_break=True)
    add("The following merged table is intentionally treated as an unsupported layout object.", "body")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(1, 0).text = "Label"
    table.cell(1, 1).text = "Value"
    doc.save(source)
    config = _reviewed_config(
        source, "student", roles, {"1": "preserve"}, {}, [1, 6],
        _compliance({"headings", "tables"}, {"tables"}), uncertain_tables={1},
    )
    return source, config, {
        "id": "complex_preserve",
        "profile": "student",
        "source": source.name,
        "config": "complex_preserve.structure.json",
        "font": "Times New Roman",
        "tags": ["student", "merged-table", "preservation"],
        "require_visual_review": False,
        "expect": {
            "machine_checks": {"content_preservation": "passed"},
            "exact_issue_codes": [],
            "preflight_max_issues": 0,
        },
    }


def generate_synthetic_suite(output_dir):
    output_dir = Path(output_dir).expanduser().resolve()
    if output_dir.exists() or output_dir.is_symlink():
        raise ValueError("合成基准目录必须是新目录。")
    output_dir.mkdir(parents=True)
    image = output_dir / "synthetic.png"
    image.write_bytes(TINY_PNG)
    cases = []
    for builder in (_student_core, lambda folder: _professional_warnings(folder, image), _complex_preserve):
        source, config, case = builder(output_dir)
        write_new_json(output_dir / case["config"], config)
        cases.append(case)
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "suite": "synthetic-regression",
        "description": "Generated fictional documents; safe to recreate for deterministic regression testing.",
        "cases": cases,
    }
    write_new_json(output_dir / "cases.json", manifest)
    return output_dir / "cases.json"


def prepare_private_case(source, case_id, profile, output_dir, running_head=""):
    if not CASE_ID.fullmatch(case_id):
        raise ValueError("案例 id 只能使用小写字母、数字、下划线和连字符。")
    source = Path(source).expanduser().resolve()
    if not source.is_file() or source.suffix.lower() != ".docx":
        raise ValueError("需要现有 .docx 原稿。")
    if profile == "professional" and not running_head.strip():
        raise ValueError("专业论文测试案例需要作者确认的 running head。")
    output_dir = Path(output_dir).expanduser().resolve()
    if output_dir.exists() or output_dir.is_symlink():
        raise ValueError("私有案例目录必须是新目录。")
    # Complete the read-only inspection before creating anything so an
    # unsupported or damaged manuscript cannot leave a misleading case folder.
    draft = workflow.prepare(source, profile)
    output_dir.mkdir(parents=True)
    config_path = output_dir / f"{case_id}.structure.json"
    workflow.write_new_json(config_path, draft)
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "suite": case_id,
        "description": "Private case. The source is referenced in place and is not copied.",
        "cases": [{
            "id": case_id,
            "profile": profile,
            "running_head": running_head,
            "source": str(source),
            "config": config_path.name,
            "font": "Times New Roman",
            "tags": ["private", "real-manuscript"],
            "require_visual_review": True,
            "expect": {
                "machine_checks": {"content_preservation": "passed"},
                "required_issue_codes": [],
                "forbidden_issue_codes": [],
            },
        }],
    }
    manifest_path = write_new_json(output_dir / "cases.private.json", manifest)
    return manifest_path, config_path


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    generate = commands.add_parser("generate", help="生成可重复的虚构 Word 基准套件")
    generate.add_argument("output_dir", type=Path)
    prepare = commands.add_parser("prepare-case", help="为私有真实论文建立不复制原文的案例草稿")
    prepare.add_argument("source", type=Path)
    prepare.add_argument("--id", required=True)
    prepare.add_argument("--profile", choices=("student", "professional"), required=True)
    prepare.add_argument("--running-head", default="")
    prepare.add_argument("--output-dir", type=Path, required=True)
    run = commands.add_parser("run", help="运行一个新基准批次")
    run.add_argument("manifest", type=Path)
    run.add_argument("--output-dir", type=Path, required=True)
    run.add_argument("--renderer", type=Path)
    qa = commands.add_parser("record-qa", help="记录某个案例的完整逐页视觉复核")
    qa.add_argument("run_dir", type=Path)
    qa.add_argument("case_id")
    qa.add_argument("review", type=Path)
    summary = commands.add_parser("summarize", help="重新计算并显示基准汇总")
    summary.add_argument("run_dir", type=Path)
    verify = commands.add_parser("verify-sources", help="核对本地公开测试文档与来源登记的大小和哈希")
    verify.add_argument("registry", type=Path)
    verify.add_argument("--corpus-dir", type=Path, required=True)
    inspect = commands.add_parser("inspect-sources", help="校验并只读扫描全部公开 Word 测试文档")
    inspect.add_argument("registry", type=Path)
    inspect.add_argument("--corpus-dir", type=Path, required=True)
    inspect.add_argument("--profile", choices=("student", "professional"), required=True)
    inspect.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        if args.command == "generate":
            path = generate_synthetic_suite(args.output_dir)
            print(path)
            return 0
        if args.command == "prepare-case":
            manifest, config = prepare_private_case(
                args.source, args.id, args.profile, args.output_dir, args.running_head
            )
            print(json.dumps({"manifest": str(manifest), "config_to_review": str(config)}, ensure_ascii=False, indent=2))
            return 0
        if args.command == "run":
            results = run_suite(args.manifest, args.output_dir, args.renderer)
            print(summary_markdown(results))
            return 1 if results["summary"]["status"] == "failed" else 0
        if args.command == "record-qa":
            qa_result, run_summary = record_case_qa(args.run_dir, args.case_id, args.review)
            print(json.dumps({"qa": qa_result, "summary": run_summary}, ensure_ascii=False, indent=2))
            return 1 if run_summary["status"] == "failed" else 0
        if args.command == "verify-sources":
            result = verify_public_sources(args.registry, args.corpus_dir)
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 0 if result["status"] == "passed" else 1
        if args.command == "inspect-sources":
            result = inspect_public_sources(args.registry, args.corpus_dir, args.profile, args.output_dir)
            print(corpus_inspection_markdown(result))
            return 0 if result["status"] == "passed" else 1
        run_dir = args.run_dir.expanduser().resolve()
        results = read_json(run_dir / "benchmark-results.json")
        save_run_summary(run_dir, results)
        print(summary_markdown(results))
        return 1 if results["summary"]["status"] == "failed" else 0
    except Exception as exc:
        print("未完成：" + str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
